import tempfile
import unittest
import zipfile
from collections import OrderedDict
from pathlib import Path

from docx import Document
from lxml import etree

from docx_footnotes import (
    CONTENT_TYPE_FOOTNOTES,
    CT_NS,
    PKGREL_NS,
    REL_TYPE_FOOTNOTES,
    W_NS,
    FootnotePatchError,
    inspect_docx_footnotes,
    patch_docx_footnotes,
)


W = "{%s}" % W_NS
REL = "{%s}" % PKGREL_NS
CT = "{%s}" % CT_NS


class DocxFootnoteTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self) -> None:
        self.temp.cleanup()

    def _source(self, name: str = "source.docx") -> Path:
        path = self.root / name
        document = Document()
        first = document.add_paragraph()
        run = first.add_run("Before [[FN:alpha]] between [[FN:beta]] after")
        run.bold = True
        second = document.add_paragraph()
        second.add_run("Split [[FN:")
        split = second.add_run("split-id")
        split.italic = True
        second.add_run("]] tail")
        document.save(path)
        return path

    @staticmethod
    def _xml(archive: zipfile.ZipFile, name: str) -> etree._Element:
        return etree.fromstring(archive.read(name))

    def test_materializes_true_footnotes_and_preserves_surrounding_runs(self) -> None:
        source = self._source()
        output = self.root / "patched.docx"
        notes = OrderedDict(
            [
                ("beta", "Second definition."),
                ("alpha", "First definition."),
                ("split-id", "A definition with two lines.\nContinuation."),
            ]
        )

        result = patch_docx_footnotes(source, output, notes)

        self.assertEqual(
            dict(result.stable_to_word_id),
            {"beta": 1, "alpha": 2, "split-id": 3},
        )
        self.assertTrue(result.inventory.valid, result.inventory.problems)
        self.assertEqual(result.inventory.reference_ids, (2, 1, 3))
        self.assertEqual(result.inventory.definition_ids, (1, 2, 3))
        self.assertEqual(
            result.inventory.separators,
            ((-1, "separator"), (0, "continuationSeparator")),
        )

        reopened = Document(output)
        self.assertEqual(reopened.paragraphs[0].text, "Before  between  after")
        self.assertEqual(reopened.paragraphs[1].text, "Split  tail")

        with zipfile.ZipFile(output) as archive:
            document = self._xml(archive, "word/document.xml")
            self.assertFalse(
                any(
                    "[[FN:" in (text.text or "")
                    for text in document.findall(".//" + W + "t")
                )
            )
            first_paragraph = document.find(".//" + W + "p")
            sequence = []
            for run in first_paragraph.findall(W + "r"):
                text = "".join(node.text or "" for node in run.findall(W + "t"))
                reference = run.find(W + "footnoteReference")
                if text:
                    sequence.append(("text", text))
                    self.assertIsNotNone(run.find(W + "rPr/" + W + "b"))
                elif reference is not None:
                    sequence.append(("ref", int(reference.get(W + "id"))))
            self.assertEqual(
                sequence,
                [
                    ("text", "Before "),
                    ("ref", 2),
                    ("text", " between "),
                    ("ref", 1),
                    ("text", " after"),
                ],
            )

            footnotes = self._xml(archive, "word/footnotes.xml")
            separator = footnotes.find(W + "footnote[@" + W + "id='-1']")
            continuation = footnotes.find(W + "footnote[@" + W + "id='0']")
            self.assertEqual(separator.get(W + "type"), "separator")
            self.assertIsNotNone(separator.find(".//" + W + "separator"))
            self.assertEqual(
                continuation.get(W + "type"), "continuationSeparator"
            )
            self.assertIsNotNone(
                continuation.find(".//" + W + "continuationSeparator")
            )
            multiline = footnotes.find(W + "footnote[@" + W + "id='3']")
            self.assertEqual(len(multiline.findall(W + "p")), 2)

            relationships = self._xml(
                archive, "word/_rels/document.xml.rels"
            )
            note_relationships = [
                relationship
                for relationship in relationships.findall(REL + "Relationship")
                if relationship.get("Type") == REL_TYPE_FOOTNOTES
            ]
            self.assertEqual(len(note_relationships), 1)
            self.assertEqual(note_relationships[0].get("Target"), "footnotes.xml")

            content_types = self._xml(archive, "[Content_Types].xml")
            overrides = [
                override
                for override in content_types.findall(CT + "Override")
                if override.get("PartName") == "/word/footnotes.xml"
            ]
            self.assertEqual(len(overrides), 1)
            self.assertEqual(
                overrides[0].get("ContentType"), CONTENT_TYPE_FOOTNOTES
            )

    def test_same_path_patch_is_atomic_and_auditable(self) -> None:
        source = self._source()
        patch_docx_footnotes(
            source,
            source,
            [
                ("alpha", "Alpha."),
                ("beta", "Beta."),
                ("split-id", "Split."),
            ],
        )
        inventory = inspect_docx_footnotes(source)
        self.assertTrue(inventory.valid, inventory.problems)
        self.assertEqual(inventory.reference_ids, (1, 2, 3))
        with zipfile.ZipFile(source) as archive:
            self.assertIsNone(archive.testzip())

    def test_marker_contract_failure_does_not_replace_existing_output(self) -> None:
        source = self._source()
        output = self.root / "existing.docx"
        sentinel = b"existing-output-must-survive"
        output.write_bytes(sentinel)

        with self.assertRaisesRegex(
            FootnotePatchError, "each footnote needs exactly one marker"
        ):
            patch_docx_footnotes(
                source,
                output,
                [
                    ("alpha", "Alpha."),
                    ("beta", "Beta."),
                    ("split-id", "Split."),
                    ("missing", "Missing."),
                ],
            )

        self.assertEqual(output.read_bytes(), sentinel)

    def test_rejects_unknown_markers_and_duplicate_stable_ids(self) -> None:
        source = self._source()
        with self.assertRaisesRegex(
            FootnotePatchError, "markers have no footnote definitions"
        ):
            patch_docx_footnotes(source, self.root / "unknown.docx", [("alpha", "A")])

        with self.assertRaisesRegex(FootnotePatchError, "duplicate footnote stable ID"):
            patch_docx_footnotes(
                source,
                self.root / "duplicate.docx",
                [("alpha", "A"), ("alpha", "B")],
            )


if __name__ == "__main__":
    unittest.main()
