from __future__ import annotations

import io
import json
import tempfile
import unittest
from pathlib import Path
from typing import Sequence
from unittest.mock import patch

from book_pipeline import write_knowledge_base
from knowledge_base_cli import main
from rag_knowledge_base import metadata_sidecar_path_for, vector_index_path_for


ROWS = [
    {
        "id": "a" * 40,
        "title": "量子理论",
        "chapter_id": "chapter-1",
        "chapter_order": 1,
        "content": "量子叠加描述微观系统。",
    },
    {
        "id": "b" * 40,
        "title": "历史方法",
        "chapter_id": "chapter-2",
        "chapter_order": 2,
        "content": "历史研究依赖档案互证。",
    },
]


class FakeZhipuProvider:
    provider_name = "zhipu"
    model = "embedding-3"

    def __init__(self, **kwargs: object) -> None:
        self.kwargs = kwargs

    @staticmethod
    def _vector(text: str) -> list[float]:
        return [float(text.count("量子")), float(text.count("历史")), 1.0]

    def embed_documents(self, texts: Sequence[str]) -> Sequence[Sequence[float]]:
        return [self._vector(text) for text in texts]

    def embed_query(self, text: str) -> Sequence[float]:
        return self._vector(text)


class MissingKeyProvider(FakeZhipuProvider):
    def embed_query(self, text: str) -> Sequence[float]:
        raise RuntimeError("missing key")


def _json(stdout: io.StringIO) -> dict[str, object]:
    return json.loads(stdout.getvalue())


class KnowledgeBaseCliTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tempdir = tempfile.TemporaryDirectory()
        self.addCleanup(self.tempdir.cleanup)
        self.root = Path(self.tempdir.name)
        self.output = self.root / "book"
        self.kb = self.output / "knowledge_base.jsonl"
        write_knowledge_base(self.kb, ROWS)

    def test_register_builds_zhipu_embedding_index_without_network(self) -> None:
        stdout = io.StringIO()

        with patch("knowledge_base_cli.ZhipuEmbeddingProvider", FakeZhipuProvider):
            code = main(["register", str(self.output), "--dimensions", "2048"], stdout)

        payload = _json(stdout)
        self.assertEqual(code, 0)
        self.assertTrue(vector_index_path_for(self.kb).is_file())
        self.assertEqual(payload["embedding"]["status"], "ready")
        self.assertEqual(payload["registered"]["provider"], "zhipu")

    def test_retrieve_uses_semantic_when_requested_and_index_ready(self) -> None:
        with patch("knowledge_base_cli.ZhipuEmbeddingProvider", FakeZhipuProvider):
            self.assertEqual(main(["register", str(self.output)], io.StringIO()), 0)
            stdout = io.StringIO()
            code = main(
                ["retrieve", str(self.output), "量子问题", "--semantic", "--top-k", "1"],
                stdout,
            )

        payload = _json(stdout)
        self.assertEqual(code, 0)
        self.assertTrue(payload["semantic_used"])
        self.assertEqual(payload["retrieval_mode"], "semantic")
        self.assertEqual(payload["hits"][0]["id"], "a" * 40)

    def test_retrieve_falls_back_to_lexical_when_semantic_is_not_ready(self) -> None:
        stdout = io.StringIO()

        code = main(
            ["retrieve", str(self.output), "历史档案", "--semantic", "--top-k", "1"],
            stdout,
        )

        payload = _json(stdout)
        self.assertEqual(code, 0)
        self.assertFalse(payload["semantic_used"])
        self.assertEqual(payload["retrieval_mode"], "lexical")
        self.assertEqual(payload["hits"][0]["id"], "b" * 40)

    def test_retrieve_falls_back_to_lexical_when_query_provider_fails(self) -> None:
        with patch("knowledge_base_cli.ZhipuEmbeddingProvider", FakeZhipuProvider):
            self.assertEqual(main(["register", str(self.output)], io.StringIO()), 0)
        stdout = io.StringIO()

        with patch("knowledge_base_cli.ZhipuEmbeddingProvider", MissingKeyProvider):
            code = main(
                ["retrieve", str(self.output), "历史档案", "--semantic", "--top-k", "1"],
                stdout,
            )

        payload = _json(stdout)
        self.assertEqual(code, 0)
        self.assertFalse(payload["semantic_used"])
        self.assertEqual(payload["retrieval_mode"], "lexical")
        self.assertIn("missing key", payload["semantic_error"])

    def test_status_reports_current_manifest_and_embedding_state(self) -> None:
        stdout = io.StringIO()

        code = main(["status", str(self.output)], stdout)

        payload = _json(stdout)
        self.assertEqual(code, 0)
        self.assertTrue(payload["current"])
        self.assertEqual(payload["chunk_count"], 2)
        self.assertEqual(payload["embedding"]["status"], "awaiting_provider")
        self.assertFalse(payload["metadata"]["exists"])
        self.assertEqual(payload["metadata"]["coverage"], 0.0)

    def test_status_reports_metadata_coverage(self) -> None:
        metadata_sidecar_path_for(self.kb).write_text(
            json.dumps(
                {
                    "id": "a" * 40,
                    "book_id": "book-japan",
                    "book_title": "日本的思想",
                },
                ensure_ascii=False,
            )
            + "\n",
            encoding="utf-8",
        )
        stdout = io.StringIO()

        code = main(["status", str(self.output)], stdout)

        payload = _json(stdout)
        self.assertEqual(code, 0)
        self.assertTrue(payload["metadata"]["exists"])
        self.assertEqual(payload["metadata"]["row_count"], 1)
        self.assertEqual(payload["metadata"]["coverage"], 0.5)

    def test_retrieve_auto_routes_explicit_book_mention(self) -> None:
        sidecar = metadata_sidecar_path_for(self.kb)
        sidecar.write_text(
            "\n".join(
                [
                    json.dumps(
                        {
                            "id": "a" * 40,
                            "book_id": "book-japan",
                            "book_title": "日本的思想",
                            "author": "丸山真男",
                            "language": "zh",
                        },
                        ensure_ascii=False,
                    ),
                    json.dumps(
                        {
                            "id": "b" * 40,
                            "book_id": "book-history",
                            "book_title": "历史方法论",
                            "author": "作者乙",
                            "language": "zh",
                        },
                        ensure_ascii=False,
                    ),
                ]
            )
            + "\n",
            encoding="utf-8",
        )
        stdout = io.StringIO()

        code = main(
            [
                "retrieve",
                str(self.output),
                "丸山真男如何理解日本思想中的量子理论",
                "--mode",
                "lexical",
                "--top-k",
                "2",
            ],
            stdout,
        )

        payload = _json(stdout)
        self.assertEqual(code, 0)
        self.assertEqual(payload["routing"]["inferred_books"], ["日本的思想"])
        self.assertEqual(payload["routing"]["matched_authors"], ["丸山真男"])
        self.assertEqual([hit["id"] for hit in payload["hits"]], ["a" * 40])

    def test_derive_docx_creates_artifact_directory_and_jsonl(self) -> None:
        from docx import Document

        docx_path = self.root / "source.docx"
        document = Document()
        document.add_heading("第一章", level=1)
        document.add_paragraph("第一章正文。")
        document.add_heading("第二章", level=1)
        document.add_paragraph("第二章正文。")
        document.save(docx_path)
        stdout = io.StringIO()

        code = main(["derive-docx", str(docx_path)], stdout)

        payload = _json(stdout)
        derived = self.root / "source" / "knowledge_base.jsonl"
        self.assertEqual(code, 0)
        self.assertTrue(derived.is_file())
        self.assertEqual(payload["chunk_count"], 2)
        self.assertEqual(payload["derived_from"], str(docx_path.resolve()))

    def test_derive_docx_fails_without_heading_contract(self) -> None:
        from docx import Document

        docx_path = self.root / "flat.docx"
        document = Document()
        document.add_paragraph("没有标题结构。")
        document.save(docx_path)
        stdout = io.StringIO()

        code = main(["derive-docx", str(docx_path)], stdout)

        payload = _json(stdout)
        self.assertEqual(code, 1)
        self.assertEqual(payload["error"], "RuntimeError")
        self.assertIn("Heading 1", payload["message"])


if __name__ == "__main__":
    unittest.main()
