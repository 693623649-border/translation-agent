import json
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document

from docx_footnotes import patch_docx_footnotes
from docx_semantic_migration import (
    DocxSemanticMigrationError,
    apply_semantic_migration,
    extract_docx_semantic_markdown,
    write_semantic_migration,
)
from publication_semantics import parse_markdown_footnotes


class DocxSemanticMigrationTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self) -> None:
        self.temp.cleanup()

    @property
    def manifest(self):
        return [
            {
                "id": "chapter-one",
                "display_title": "第一章",
                "filename": "001_第一章.md",
                "pdf_page": 10,
                "end_pdf_page": 19,
            },
            {
                "id": "chapter-two",
                "display_title": "第二章",
                "filename": "002_第二章.md",
                "pdf_page": 20,
                "end_pdf_page": 29,
            },
        ]

    def _accepted_docx(self) -> Path:
        source = self.root / "source.docx"
        document = Document()
        document.add_paragraph("书名", style="Title")
        document.add_heading("第一章", level=1)
        document.add_paragraph("甲[[FN:alpha]]乙[[FN:beta]]。")
        document.add_heading("小节", level=2)
        document.add_paragraph("引文内容", style="Quote")
        document.add_paragraph("编号项目", style="List Number")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "名称"
        table.cell(0, 1).text = "说明"
        table.cell(1, 0).text = "表项"
        table.cell(1, 1).text = "表注[[FN:table-note]]"
        document.add_heading("第二章", level=1)
        document.add_paragraph("末章[[FN:omega]]。")
        document.save(source)

        output = self.root / "accepted.docx"
        # Definition order deliberately differs from body occurrence order.
        # This yields body Word IDs (2, 1, 3, 4), which must not exchange notes.
        patch_docx_footnotes(
            source,
            output,
            [
                ("beta", "乙注。"),
                ("alpha", "甲注。"),
                ("table-note", "表格注。"),
                ("omega", "末注。"),
            ],
        )
        return output

    @staticmethod
    def _rewrite_member(path: Path, member: str, transform) -> None:
        replacement = path.with_suffix(".rewritten.docx")
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
            replacement, "w"
        ) as target:
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename == member:
                    data = transform(data)
                target.writestr(info, data)
        replacement.replace(path)

    def test_extracts_document_order_markdown_and_exact_true_footnotes(self) -> None:
        migration = extract_docx_semantic_markdown(
            self._accepted_docx(), self.manifest
        )

        self.assertEqual(len(migration.chapters), 2)
        first = migration.chapters[0]
        self.assertIn("# 第一章", first.markdown)
        self.assertIn("## 小节", first.markdown)
        self.assertIn("> 引文内容", first.markdown)
        self.assertIn("1. 编号项目", first.markdown)
        self.assertIn("| 名称 | 说明 |", first.markdown)
        self.assertIn("| 表项 | 表注[^accepted-docx-fn-000003] |", first.markdown)
        self.assertIn(
            "甲[^accepted-docx-fn-000001]乙[^accepted-docx-fn-000002]。",
            first.markdown,
        )
        self.assertEqual(
            [item.word_id for item in first.footnotes], [2, 1, 3]
        )
        self.assertEqual(
            [item.text for item in first.footnotes],
            ["甲注。", "乙注。", "表格注。"],
        )
        self.assertEqual(first.footnotes[0].context_before, "甲")
        self.assertTrue(first.footnotes[0].context_after.startswith("乙"))
        self.assertEqual(
            [item.occurrence for item in migration.chapters[1].footnotes], [4]
        )

        for chapter in migration.chapters:
            inventory = parse_markdown_footnotes(chapter.markdown)
            self.assertTrue(inventory.valid)
            self.assertEqual(
                len(inventory.references), len(inventory.definitions)
            )
        audit = migration.audit_dict()
        self.assertEqual(audit["status"], "passed")
        self.assertFalse(audit["release_blocked"])
        self.assertEqual(audit["generated_by"], "docx_semantic_migration")
        self.assertEqual(audit["mode"], "accepted-docx-semantic-extraction")
        self.assertEqual(audit["summary"]["footnote_count"], 4)
        self.assertEqual(
            audit["migration"]["reference_order"], "document-occurrence"
        )
        self.assertEqual(
            audit["migration"]["word_id_non_monotonic_positions"],
            [
                {
                    "occurrence": 2,
                    "word_id": 1,
                    "previous_occurrence": 1,
                    "previous_word_id": 2,
                    "chapter_id": "chapter-one",
                    "block_index": 2,
                    "stable_id": "accepted-docx-fn-000002",
                }
            ],
        )
        self.assertTrue(
            all(chapter["reviewed_override"] for chapter in audit["chapters"])
        )

    def test_writes_verifier_compatible_chapters_audit_and_manifest_fields(self) -> None:
        migration = extract_docx_semantic_markdown(
            self._accepted_docx(), self.manifest
        )
        chapter_dir = self.root / "chapters"
        audit_path = self.root / "audit" / "semantic-reconstruction.json"
        write_semantic_migration(migration, chapter_dir, audit_path=audit_path)

        self.assertEqual(
            (chapter_dir / "001_第一章.md").read_text(encoding="utf-8"),
            migration.chapters[0].markdown,
        )
        audit = json.loads(audit_path.read_text(encoding="utf-8"))
        self.assertFalse(audit["summary"]["release_blocked"])
        updated = migration.updated_manifest(self.manifest)
        self.assertEqual(
            [item["semantic_footnote_count"] for item in updated], [3, 1]
        )
        self.assertTrue(
            all(item["semantic_issue_count"] == 0 for item in updated)
        )
        self.assertTrue(all(item["reviewed_override"] for item in updated))

    def test_applies_complete_reviewed_bundle_with_external_recovery_archive(self) -> None:
        accepted = self._accepted_docx()
        output = self.root / "publication"
        (output / "chapters").mkdir(parents=True)
        (output / "audit").mkdir()
        (output / "chapters.json").write_text(
            json.dumps(self.manifest, ensure_ascii=False),
            encoding="utf-8",
        )
        for item in self.manifest:
            (output / "chapters" / item["filename"]).write_text(
                "# legacy\n",
                encoding="utf-8",
            )
        (output / "toc.json").write_text("{}\n", encoding="utf-8")
        published_docx = output / "accepted.docx"
        published_docx.write_bytes(accepted.read_bytes())
        migration = extract_docx_semantic_markdown(
            published_docx,
            self.manifest,
        )
        backup = self.root / "recovery" / "before.zip"

        result = apply_semantic_migration(
            migration,
            output,
            backup_path=backup,
        )

        self.assertEqual(result["chapter_count"], 2)
        self.assertEqual(result["footnote_count"], 4)
        self.assertTrue(backup.is_file())
        with zipfile.ZipFile(backup) as archive:
            names = set(archive.namelist())
            self.assertIn("backup-manifest.json", names)
            self.assertIn("chapters.json", names)
            self.assertIn("chapters/001_第一章.md", names)
            self.assertIn("accepted.docx", names)
            backup_manifest = json.loads(
                archive.read("backup-manifest.json").decode("utf-8")
            )
            self.assertFalse(backup_manifest["reviewed_chapters_existed"])
        updated = json.loads(
            (output / "chapters.json").read_text(encoding="utf-8")
        )
        self.assertTrue(all(item["reviewed_override"] for item in updated))
        for chapter in migration.chapters:
            canonical = output / "chapters" / chapter.filename
            reviewed = output / "reviewed_chapters" / f"{chapter.chapter_id}.md"
            self.assertEqual(canonical.read_bytes(), reviewed.read_bytes())
        semantic = json.loads(
            (output / "audit" / "semantic-reconstruction.json").read_text(
                encoding="utf-8"
            )
        )
        migration_audit = json.loads(
            (output / "audit" / "docx-semantic-migration.json").read_text(
                encoding="utf-8"
            )
        )
        self.assertEqual(semantic, migration_audit)

    def test_apply_rejects_backup_inside_publication_output(self) -> None:
        accepted = self._accepted_docx()
        output = self.root / "publication"
        (output / "chapters").mkdir(parents=True)
        (output / "chapters.json").write_text(
            json.dumps(self.manifest, ensure_ascii=False),
            encoding="utf-8",
        )
        for item in self.manifest:
            (output / "chapters" / item["filename"]).write_text(
                "# legacy\n",
                encoding="utf-8",
            )
        published_docx = output / "accepted.docx"
        published_docx.write_bytes(accepted.read_bytes())
        migration = extract_docx_semantic_markdown(
            published_docx,
            self.manifest,
        )

        with self.assertRaisesRegex(
            DocxSemanticMigrationError,
            "backup must be outside",
        ):
            apply_semantic_migration(
                migration,
                output,
                backup_path=output / "audit" / "backup.zip",
            )

    def test_rejects_heading_manifest_mismatch(self) -> None:
        manifest = [dict(item) for item in self.manifest]
        manifest[1]["display_title"] = "错误标题"
        with self.assertRaisesRegex(
            DocxSemanticMigrationError, "Heading 1/manifest mismatch"
        ):
            extract_docx_semantic_markdown(self._accepted_docx(), manifest)

    def test_rejects_duplicate_or_orphan_word_footnotes(self) -> None:
        duplicate = self._accepted_docx()

        def add_duplicate(data: bytes) -> bytes:
            needle = b'<w:footnoteReference w:id="2"/>'
            self.assertIn(needle, data)
            return data.replace(needle, needle + needle, 1)

        self._rewrite_member(duplicate, "word/document.xml", add_duplicate)
        with self.assertRaisesRegex(
            DocxSemanticMigrationError, "invalid Word footnote package"
        ):
            extract_docx_semantic_markdown(duplicate, self.manifest)

        orphan = self._accepted_docx()

        def remove_reference(data: bytes) -> bytes:
            needle = b'<w:footnoteReference w:id="2"/>'
            self.assertIn(needle, data)
            return data.replace(needle, b"", 1)

        self._rewrite_member(orphan, "word/document.xml", remove_reference)
        with self.assertRaisesRegex(
            DocxSemanticMigrationError, "invalid Word footnote package"
        ):
            extract_docx_semantic_markdown(orphan, self.manifest)


if __name__ == "__main__":
    unittest.main()
