import json
import hashlib
import io
import os
import re
import sys
import tempfile
import threading
import unittest
import zipfile
import xml.etree.ElementTree as ET
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest.mock import patch

import fitz
from PIL import Image
from docx.enum.text import WD_ALIGN_PARAGRAPH

from book_pipeline import (
    ChatTranslator,
    CodingPlanVisionOCR,
    DeepSeekClient,
    GlmClient,
    McpStdioClient,
    PageRecord,
    TocEntry,
    _exclusive_stage_lock,
    annotate_printed_page_markers,
    apply_page_mapping,
    build_parser,
    build_bookmarked_pdf,
    build_docx,
    build_epub,
    build_translation_client,
    compile_chapters,
    clean_ocr_text,
    clean_translation_text,
    detect_language,
    infer_page_offset,
    import_existing_ocr,
    main,
    markdown_inline_to_plain_text,
    normalize_target_script,
    ocr_pdf,
    output_status,
    parse_page_spec,
    remove_duplicate_title,
    resolve_api_key,
    resolve_translation_api_key,
    resolve_worker_counts,
    save_page_record,
    strip_publication_metadata,
    strip_reviewed_publication_metadata,
    trim_before_next_title,
    translate_non_chinese_pages,
    write_json,
)


class UtilityTests(unittest.TestCase):
    def test_model_stage_lock_decorator_preserves_keyword_call_api(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            translate_non_chinese_pages(
                records=[],
                output_dir=output,
                translator=object(),
                target_language="简体中文",
                force=False,
            )
            self.assertTrue((output / ".stage_locks" / "translate.lock").is_file())

    def test_chapter_selector_is_rejected_outside_verify_phase(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaises(SystemExit):
                main(
                    [
                        "-o",
                        directory,
                        "--phase",
                        "status",
                        "--chapter-id",
                        "1",
                    ]
                )

    def test_english_translation_prompt_uses_source_language_grammar(self) -> None:
        class RecordingClient:
            def __init__(self) -> None:
                self.prompts: list[str] = []

            def chat_text(
                self,
                prompt: str,
                *,
                system: str,
                max_tokens: int = 16384,
            ) -> str:
                self.prompts.append(prompt)
                return f"译文{len(self.prompts)}。"

        client = RecordingClient()
        translator = ChatTranslator(client)

        translated = translator.translate(
            "# 审定标题\n\nEnglish body.\n\n## 御宅世界影像\n\nMore English body.",
            source_language="en",
            target_language="简体中文",
        )

        self.assertEqual(
            translated,
            "# 审定标题\n\n译文1。\n\n## 御宅世界影像\n\n译文2。",
        )
        self.assertEqual(len(client.prompts), 2)
        for prompt in client.prompts:
            self.assertIn("源语言标签：en", prompt)
            self.assertIn("依据源语言的语法和上下文", prompt)
            self.assertNotIn("日语语法", prompt)
            self.assertIn("Markdown 标题已由程序保护", prompt)
            self.assertNotIn("审定标题", prompt)
            self.assertNotIn("御宅世界影像", prompt)

    def test_translation_prompt_requires_every_numbered_footnote_definition(self) -> None:
        class RecordingClient:
            def __init__(self) -> None:
                self.prompts: list[str] = []

            def chat_text(
                self,
                prompt: str,
                *,
                system: str,
                max_tokens: int = 16384,
            ) -> str:
                self.prompts.append(prompt)
                source = prompt.split("\n原文：\n", 1)[1]
                label = source.split(maxsplit=1)[0]
                return f"{label} 第{label}条脚注全文。\n第{label}条续行。"

        client = RecordingClient()
        translated = ChatTranslator(client).translate(
            "1 First footnote definition in full.\nFirst continuation.\n"
            "2 Second footnote definition in full.\nSecond continuation.\n"
            "3 Third footnote definition in full.\nThird continuation.\n"
            "4 Fourth footnote definition in full.\nFourth continuation.",
            source_language="en",
            target_language="简体中文",
        )

        self.assertEqual(len(client.prompts), 4)
        sources = [prompt.split("\n原文：\n", 1)[1] for prompt in client.prompts]
        self.assertEqual(
            sources,
            [
                "1 First footnote definition in full.\nFirst continuation.",
                "2 Second footnote definition in full.\nSecond continuation.",
                "3 Third footnote definition in full.\nThird continuation.",
                "4 Fourth footnote definition in full.\nFourth continuation.",
            ],
        )
        self.assertEqual(
            translated,
            "1 第1条脚注全文。\n第1条续行。\n\n"
            "2 第2条脚注全文。\n第2条续行。\n\n"
            "3 第3条脚注全文。\n第3条续行。\n\n"
            "4 第4条脚注全文。\n第4条续行。",
        )
        for index, prompt in enumerate(client.prompts, start=1):
            self.assertIn("每一个行首脚注编号及其对应定义全文", prompt)
            self.assertIn("严禁合并、跳号、截断、只保留编号或省略出处", prompt)
            self.assertIn(f"必须逐项原样保留）：{index}", prompt)

    def test_dotted_numbers_are_not_footnotes_or_strict_labels(self) -> None:
        class RecordingClient:
            def __init__(self) -> None:
                self.prompts: list[str] = []

            def chat_text(
                self,
                prompt: str,
                *,
                system: str,
                max_tokens: int = 16384,
            ) -> str:
                self.prompts.append(prompt)
                return "这是包含年代和列表编号的译文。"

        client = RecordingClient()
        translated = ChatTranslator(client).translate(
            "1853. A year in prose.\n240. A dotted list item.",
            source_language="en",
            target_language="简体中文",
        )

        self.assertEqual(translated, "这是包含年代和列表编号的译文。")
        self.assertEqual(len(client.prompts), 1)
        self.assertNotIn("本分块检测到的行首编号", client.prompts[0])

    def test_parenthesized_number_is_strict_but_year_is_not(self) -> None:
        class RecordingClient:
            def __init__(self) -> None:
                self.prompts: list[str] = []

            def chat_text(
                self,
                prompt: str,
                *,
                system: str,
                max_tokens: int = 16384,
            ) -> str:
                self.prompts.append(prompt)
                return "(7) 保留的编号条目。\n年代出处已翻译。"

        client = RecordingClient()
        ChatTranslator(client).translate(
            "(7) A parenthesized numbered item.\n(1853) A bibliographic year.",
            source_language="en",
            target_language="简体中文",
        )

        self.assertEqual(len(client.prompts), 2)
        self.assertIn("必须逐项原样保留）：7。", client.prompts[0])
        self.assertNotIn("本分块检测到的行首编号", client.prompts[1])

    def test_numbered_item_isolated_but_indented_volume_is_not_strict(self) -> None:
        class RecordingClient:
            def __init__(self) -> None:
                self.prompts: list[str] = []

            def chat_text(
                self,
                prompt: str,
                *,
                system: str,
                max_tokens: int = 16384,
            ) -> str:
                self.prompts.append(prompt)
                source = prompt.split("\n原文：\n", 1)[1]
                if source.startswith("(2) "):
                    return "(2) 第二项。"
                return "译文。"

        client = RecordingClient()
        ChatTranslator(client).translate(
            "Introductory prose.\n"
            "(2) Second enumerated body item.\n\n"
            "2016. The International Encyclopedia (Volume\n"
            "     2) (Malden: Wiley-Blackwell)",
            source_language="en",
            target_language="简体中文",
        )

        self.assertEqual(len(client.prompts), 2)
        self.assertIn("必须逐项原样保留）：2", client.prompts[1])
        self.assertNotIn("必须逐项原样保留）：2、2", client.prompts[1])

    def test_translation_rejects_missing_numbered_definition(self) -> None:
        class DroppingClient:
            def chat_text(
                self,
                prompt: str,
                *,
                system: str,
                max_tokens: int = 16384,
            ) -> str:
                return "1 只返回了第一条脚注。"

        with self.assertRaisesRegex(RuntimeError, "leading labels: \\['2'\\]"):
            ChatTranslator(DroppingClient()).translate(
                "1 First footnote.\n2 Second footnote.",
                source_language="en",
                target_language="简体中文",
            )

    def test_translation_accepts_table_artifact_labels(self) -> None:
        class DroppingClient:
            def chat_text(
                self,
                prompt: str,
                *,
                system: str,
                max_tokens: int = 16384,
            ) -> str:
                return "全部正文翻译完成。"

        # Two-column chronology rows: the ``|`` artifact marks a table page,
        # so stray leading numbers are OCR noise, not omitted footnotes.
        translator = ChatTranslator(DroppingClient())
        translator.translate(
            "2 出狱被允许的陀思妥耶夫斯基的前途 | 〇 托尔斯泰诞生。\n"
            "3 流刑第二年 | 〇 拿破仑政变。",
            source_language="ja",
            target_language="简体中文",
        )

    def test_translation_accepts_zero_and_header_labels(self) -> None:
        class DroppingClient:
            def chat_text(
                self,
                prompt: str,
                *,
                system: str,
                max_tokens: int = 16384,
            ) -> str:
                return "正文译文。\n\n塞米巴拉金斯克。"

        # "0" is a misread circle marker and "4" here is the chapter number
        # in the running head; neither is a footnote definition.
        translator = ChatTranslator(DroppingClient())
        translator.translate(
            "4 塞米巴拉金斯克\n\n正文文字。\n0 对自由的向往。",
            source_language="ja",
            target_language="简体中文",
        )

    def test_translation_rejects_japanese_only_output(self) -> None:
        class EchoClient:
            def chat_text(
                self,
                prompt: str,
                *,
                system: str,
                max_tokens: int = 16384,
            ) -> str:
                return "文芸時評、大衆時評をはじめいろいろな時評が、元来輿論の代表者として責めるべきところを。"

        with self.assertRaisesRegex(RuntimeError, "retained Japanese text"):
            ChatTranslator(EchoClient()).translate(
                "文芸時評、大衆時評をはじめいろいろな時評が、元来輿論の代表者として責めるべきところを。",
                source_language="ja",
                target_language="简体中文",
            )

    def test_translation_accepts_kana_name_glosses(self) -> None:
        class GlossClient:
            def chat_text(
                self,
                prompt: str,
                *,
                system: str,
                max_tokens: int = 16384,
            ) -> str:
                return "在文学上给予基里尔（キイ）最强烈刺激的朋友，是父亲的朋友。"

        translator = ChatTranslator(GlossClient())
        result = translator.translate(
            "文学上キリルに最も強い刺激を与へた友は父の友であつた。",
            source_language="ja",
            target_language="简体中文",
        )
        self.assertIn("基里尔（キイ）", result)

    def test_status_accepts_multiple_ocr_model_prefixes(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            save_page_record(
                output,
                PageRecord(
                    pdf_page=1,
                    text="竖排",
                    ocr_model="coding-plan/glm-4.6v-vision-mcp/vertical-v2",
                ),
            )
            save_page_record(
                output,
                PageRecord(
                    pdf_page=2,
                    text="横排",
                    ocr_model="coding-plan/glm-4.6v-vision-mcp/horizontal-v2",
                ),
            )

            status = output_status(
                output,
                expected_ocr_model_prefix=(
                    "coding-plan/glm-4.6v-vision-mcp/vertical-v2,"
                    "coding-plan/glm-4.6v-vision-mcp/horizontal-v2"
                ),
            )

        self.assertEqual(status["ocr_pages_profile_fresh"], 2)

    def test_status_distinguishes_exact_ocr_identity_from_explicit_prefix(self) -> None:
        models = (
            "tesseract/chi_sim/psm-3",
            "tesseract/chi_sim/psm-30",
            "coding-plan/glm-4.6v-vision-mcp/horizontal-v2",
            "coding-plan/glm-4.6v-vision-mcp/horizontal-v20",
        )
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            for pdf_page, model in enumerate(models, start=1):
                save_page_record(
                    output,
                    PageRecord(
                        pdf_page=pdf_page,
                        text="正文",
                        ocr_model=model,
                    ),
                )

            tesseract_exact = output_status(
                output,
                expected_ocr_model_exact="tesseract/chi_sim/psm-3",
            )
            tesseract_prefix = output_status(
                output,
                expected_ocr_model_prefix="tesseract/chi_sim/psm-3",
            )
            glm_exact = output_status(
                output,
                expected_ocr_model_exact=(
                    "coding-plan/glm-4.6v-vision-mcp/horizontal-v2"
                ),
            )
            glm_prefix = output_status(
                output,
                expected_ocr_model_prefix=(
                    "coding-plan/glm-4.6v-vision-mcp/horizontal-v2"
                ),
            )

        self.assertEqual(tesseract_exact["ocr_pages_profile_fresh"], 1)
        self.assertEqual(tesseract_prefix["ocr_pages_profile_fresh"], 2)
        self.assertEqual(glm_exact["ocr_pages_profile_fresh"], 1)
        self.assertEqual(glm_prefix["ocr_pages_profile_fresh"], 2)

    def test_status_only_accepts_fresh_full_publication_report(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            audit = output / "audit"
            chapters = output / "chapters"
            audit.mkdir()
            chapters.mkdir()
            report_path = audit / "release-report.json"
            chapter_path = chapters / "001_test.md"
            chapter_path.write_text("# 测试\n\n正文。\n", encoding="utf-8")
            report_path.write_text(
                json.dumps({"mode": "chapters", "status": "passed"}),
                encoding="utf-8",
            )

            status = output_status(output)
            self.assertFalse(status["verification_ready"])
            self.assertIsNone(status["verification_status"])

            report_path.write_text(
                json.dumps(
                    {
                        "mode": "full",
                        "status": "passed",
                        "summary": {"chapter_count": 1},
                    }
                ),
                encoding="utf-8",
            )
            os.utime(report_path, ns=(1_000_000_000, 1_000_000_000))
            os.utime(chapter_path, ns=(2_000_000_000, 2_000_000_000))
            stale = output_status(output)
            self.assertTrue(stale["verification_ready"])
            self.assertTrue(stale["verification_stale"])

            os.utime(report_path, ns=(3_000_000_000, 3_000_000_000))
            fresh = output_status(output)
            self.assertEqual(fresh["verification_status"], "passed")
            self.assertFalse(fresh["verification_stale"])

    def test_status_accepts_fresh_word_publication_report(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            audit = output / "audit"
            chapters = output / "chapters"
            audit.mkdir()
            chapters.mkdir()
            report_path = audit / "word-release-report.json"
            chapter_path = chapters / "001_test.md"
            chapter_path.write_text("# 测试\n\n正文。\n", encoding="utf-8")
            report_path.write_text(
                json.dumps(
                    {
                        "mode": "full",
                        "publication_profile": "word",
                        "status": "passed",
                        "release_ready": True,
                        "summary": {"chapter_count": 1},
                    }
                ),
                encoding="utf-8",
            )
            os.utime(chapter_path, ns=(1_000_000_000, 1_000_000_000))
            os.utime(report_path, ns=(2_000_000_000, 2_000_000_000))

            status = output_status(output)

            self.assertTrue(status["verification_ready"])
            self.assertEqual(status["verification_profile"], "word")
            self.assertEqual(status["verification_status"], "passed")
            self.assertTrue(status["verification_release_ready"])
            self.assertEqual(status["verification_report"], str(report_path))

    def test_stage_lock_rejects_duplicate_process_for_same_output(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            with _exclusive_stage_lock(output, "ocr"):
                with self.assertRaisesRegex(RuntimeError, "Another ocr process"):
                    with _exclusive_stage_lock(output, "ocr"):
                        self.fail("duplicate stage lock should not be acquired")

    def test_closed_vision_backend_cannot_recreate_mcp_process(self) -> None:
        backend = object.__new__(CodingPlanVisionOCR)
        backend.closed = threading.Event()
        backend.closed.set()
        with self.assertRaisesRegex(RuntimeError, "backend is closed"):
            backend._client()

    def test_vertical_mcp_prompt_requires_right_to_left_column_order(self) -> None:
        client = object.__new__(McpStdioClient)
        client.reading_direction = "vertical"
        client.tool = {
            "inputSchema": {
                "properties": {
                    "image_source": {"type": "string"},
                    "prompt": {"type": "string"},
                },
                "required": ["image_source", "prompt"],
            }
        }
        arguments = client._tool_arguments(Path("page.jpg"))
        self.assertIn("最右侧文字列", arguments["prompt"])
        self.assertIn("逐列向左", arguments["prompt"])
        self.assertIn("先完整读完上方版块", arguments["prompt"])

    def test_shared_page_trim_removes_review_publication_header(self) -> None:
        text = (
            "本篇结尾。\n\n（参考文献）\n奈须蘑菇《DDD》\n"
            "TYPE-MOON 相关近作评论\nCross Review\n"
            "哈莫尼亚\n《Fate/Grand Order 终局特异点冠位时间神殿所罗门》\n（游戏）\n正文。"
        )
        trimmed = trim_before_next_title(
            text,
            "哈莫尼亚 《Fate/Grand Order 终局特异点冠位时间神殿所罗门》（游戏）",
        )
        self.assertEqual(trimmed, "本篇结尾。\n\n（参考文献）\n奈须蘑菇《DDD》")

    def test_write_json_is_safe_for_concurrent_writers(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "page.json"
            with ThreadPoolExecutor(max_workers=8) as executor:
                list(executor.map(lambda value: write_json(path, {"value": value}), range(64)))
            payload = json.loads(path.read_text(encoding="utf-8"))
            self.assertIn(payload["value"], range(64))
            self.assertEqual(list(path.parent.glob(f".{path.name}.*.tmp")), [])

    def test_normalize_target_script_to_simplified_chinese(self) -> None:
        self.assertEqual(
            normalize_target_script("傳統藝術與歷史", "简体中文"),
            "传统艺术与历史",
        )
        self.assertEqual(
            normalize_target_script("傳統藝術與歷史", "繁体中文"),
            "傳統藝術與歷史",
        )
        self.assertEqual(
            normalize_target_script("作者望著天空，阅读其所著名著，成就顯著。", "简体中文"),
            "作者望着天空，阅读其所著名著，成就显著。",
        )

    def test_glm_chat_disables_thinking(self) -> None:
        class RecordingClient(GlmClient):
            def __init__(self) -> None:
                super().__init__(api_key="test-key")
                self.requests: list[tuple[str, dict]] = []

            def _post(self, endpoint: str, payload: dict) -> dict:
                self.requests.append((endpoint, payload))
                return {"choices": [{"message": {"content": '{"entries": []}'}}]}

        client = RecordingClient()
        self.assertEqual(client.chat_text("text", system="system"), '{"entries": []}')
        self.assertEqual(client.chat_json("json", system="system"), {"entries": []})
        self.assertEqual(len(client.requests), 2)
        for endpoint, payload in client.requests:
            self.assertEqual(endpoint, "chat/completions")
            self.assertEqual(payload["thinking"], {"type": "disabled"})

    def test_deepseek_chat_disables_thinking(self) -> None:
        class RecordingClient(DeepSeekClient):
            def __init__(self) -> None:
                super().__init__(api_key="test-key", text_model="deepseek-v4-flash")
                self.requests: list[tuple[str, dict]] = []

            def _post(self, endpoint: str, payload: dict) -> dict:
                self.requests.append((endpoint, payload))
                return {"choices": [{"message": {"content": '{"entries": []}'}}]}

        client = RecordingClient()
        self.assertEqual(client.chat_text("text", system="system"), '{"entries": []}')
        self.assertEqual(client.chat_json("json", system="system"), {"entries": []})
        self.assertEqual(len(client.requests), 2)
        for endpoint, payload in client.requests:
            self.assertEqual(endpoint, "chat/completions")
            self.assertEqual(payload["model"], "deepseek-v4-flash")
            self.assertEqual(payload["thinking"], {"type": "disabled"})

    def test_standard_glm_ocr_content_filter_uses_segmented_fallback(self) -> None:
        class FilterThenReadGlm(GlmClient):
            def __init__(self) -> None:
                super().__init__(
                    api_key="test-key",
                    reading_direction="horizontal",
                )

            def _ocr_image_once(self, image_path: Path) -> tuple[str, str]:
                if "_segment_" not in image_path.stem:
                    raise RuntimeError(
                        'GLM request failed: GLM HTTP 400: {"contentFilter":[],'
                        '"error":{"code":"1301","message":"potentially unsafe content"}}'
                    )
                match = re.search(r"_segment_(\d+)_", image_path.stem)
                assert match is not None
                values = {"1": "（无可见文字）", "2": "彼得", "3": "堡正文。", "4": "结尾。"}
                return values[match.group(1)], "glm-test"

        with tempfile.TemporaryDirectory() as directory:
            image_path = Path(directory) / "page.jpg"
            Image.new("RGB", (240, 800), "white").save(image_path)
            text, request_id = FilterThenReadGlm().ocr_image(image_path)

            self.assertEqual(text, "彼得堡正文。\n\n结尾。")
            self.assertTrue(request_id.startswith("glm-segmented-"))
            self.assertEqual(list(image_path.parent.glob("*_segment_*.jpg")), [])

    def test_provider_keys_are_isolated(self) -> None:
        parser = build_parser()
        with patch.dict(
            "os.environ",
            {"GLM_CODING_API_KEY": "glm-sentinel", "DEEPSEEK_API_KEY": "deepseek-sentinel"},
            clear=True,
        ):
            args = parser.parse_args(["sample.pdf"])
            self.assertEqual(resolve_api_key(args), "glm-sentinel")
            self.assertEqual(resolve_translation_api_key(args), "deepseek-sentinel")

            args = parser.parse_args(
                ["sample.pdf", "--api-key", "glm-cli", "--translation-api-key", "deepseek-cli"]
            )
            self.assertEqual(resolve_api_key(args), "glm-cli")
            self.assertEqual(resolve_translation_api_key(args), "deepseek-cli")

        with patch.dict("os.environ", {"GLM_CODING_API_KEY": "glm-only"}, clear=True):
            args = parser.parse_args(["sample.pdf"])
            self.assertEqual(resolve_api_key(args), "glm-only")
            self.assertEqual(resolve_translation_api_key(args), "")

    def test_explicit_ocr_backend_overrides_profile_adapter(self) -> None:
        class RecordingGlmOCR:
            instances: list["RecordingGlmOCR"] = []

            def __init__(self, **kwargs: object) -> None:
                self.ocr_model = str(kwargs["ocr_model"])
                self.kwargs = kwargs
                self.instances.append(self)

            def ocr_image(self, image_path: Path) -> tuple[str, str]:
                return "识别正文内容。", "glm-override-test"

            def close(self) -> None:
                return

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            pdf_path = root / "one-page.pdf"
            document = fitz.open()
            page = document.new_page(width=200, height=200)
            page.insert_text((72, 72), "source")
            document.save(pdf_path)
            document.close()
            config_path = root / "profiles.toml"
            config_path.write_text(
                "\n".join(
                    (
                        "schema_version = 1",
                        "[profiles.vision]",
                        'adapter = "coding-plan-mcp"',
                        'provider = "glm"',
                        'model = "glm-4.6v"',
                        'credential_env = "PROFILE_CODING_KEY"',
                        "[pipeline]",
                        'ocr_profile = "vision"',
                    )
                )
                + "\n",
                encoding="utf-8",
            )
            stderr = io.StringIO()

            with (
                patch.dict("os.environ", {"GLM_OCR_API_KEY": "standard-key"}, clear=True),
                patch("book_pipeline.GlmClient", RecordingGlmOCR),
                patch("book_pipeline.load_env_file", lambda path: None),
                patch("sys.stderr", stderr),
            ):
                result = main(
                    [
                        str(pdf_path),
                        "-o",
                        str(root / "out"),
                        "--phase",
                        "ocr",
                        "--config",
                        str(config_path),
                        "--ocr-backend",
                        "glm-ocr",
                        "--ocr-model",
                        "glm-ocr",
                        "--ocr-concurrency",
                        "1",
                    ]
                )

            self.assertEqual(result, 0)
            self.assertEqual(len(RecordingGlmOCR.instances), 1)
            self.assertEqual(RecordingGlmOCR.instances[0].ocr_model, "glm-ocr")
            self.assertIn("--ocr-backend=glm-ocr overrides OCR profile", stderr.getvalue())

    def test_worker_count_precedence_and_legacy_compatibility(self) -> None:
        parser = build_parser()
        with patch.dict("os.environ", {}, clear=True):
            self.assertEqual(resolve_worker_counts(parser.parse_args(["sample.pdf"])), (4, 16))
            self.assertEqual(
                resolve_worker_counts(parser.parse_args(["sample.pdf", "--concurrency", "7"])),
                (7, 7),
            )
            self.assertEqual(
                resolve_worker_counts(
                    parser.parse_args(
                        ["sample.pdf", "--concurrency", "7", "--ocr-concurrency", "2"]
                    )
                ),
                (2, 7),
            )
            self.assertEqual(
                resolve_worker_counts(
                    parser.parse_args(
                        [
                            "sample.pdf",
                            "--concurrency",
                            "7",
                            "--translation-concurrency",
                            "9",
                        ]
                    )
                ),
                (7, 9),
            )

    def test_build_translation_client_routes_deepseek(self) -> None:
        parser = build_parser()
        with patch.dict("os.environ", {"DEEPSEEK_API_KEY": "deepseek-sentinel"}, clear=True):
            args = parser.parse_args(["sample.pdf"])
            client = build_translation_client(args, glm_api_base="https://glm.invalid")
        self.assertIsInstance(client, DeepSeekClient)
        assert isinstance(client, DeepSeekClient)
        self.assertEqual(client.api_base, "https://api.deepseek.com")
        self.assertEqual(client.text_model, "deepseek-v4-flash")
        self.assertNotIn("deepseek-sentinel", repr(client.__dict__).replace(client.api_key, ""))

    def test_api_timeout_cli(self) -> None:
        args = build_parser().parse_args(["sample.pdf", "--api-timeout", "37"])
        self.assertEqual(args.api_timeout, 37)

    def test_parse_page_spec(self) -> None:
        self.assertEqual(parse_page_spec("2-4, 7,4"), [2, 3, 4, 7])
        with self.assertRaises(ValueError):
            parse_page_spec("4-2")

    def test_clean_ocr_text(self) -> None:
        self.assertEqual(clean_ocr_text("```\n# 标题\n正文\n```"), "# 标题\n正文")
        self.assertEqual(
            clean_ocr_text("```markdown\n# 标题\n正文\n```\n\n**Quality Notes**\n清晰"),
            "# 标题\n正文",
        )
        self.assertEqual(clean_ocr_text("# 标题\n正文"), "# 标题\n正文")
        self.assertEqual(clean_ocr_text("plaintext\n日本語本文"), "日本語本文")
        self.assertEqual(
            clean_ocr_text(
                "Extracted Text\n\n正文\n\nContent Type: 书页文字\nLanguage/Format: Markdown"
            ),
            "正文",
        )
        self.assertEqual(
            clean_ocr_text("### Extracted Text\n\n正文\n\n### Content Type\n书页文字"),
            "正文",
        )
        self.assertEqual(
            clean_ocr_text(
                "Extracted Text\n\nContent Type: 书页文字\nQuality Notes: 截图内容为空白。"
            ),
            "",
        )
        self.assertEqual(clean_ocr_text("```\n```\n\n正文"), "正文")
        self.assertEqual(clean_ocr_text("```\n14\n\n正文"), "14\n\n正文")
        self.assertEqual(clean_ocr_text("14\n\n正文\n```"), "14\n\n正文")
        self.assertEqual(
            clean_ocr_text(
                "```\n```\n\n**Content Type** 空白书页\n\n**Quality Notes** 截图内容为空白。"
            ),
            "",
        )
        self.assertEqual(
            clean_ocr_text("正文\n\n```\n[无法辨认]\n```"),
            "正文",
        )
        self.assertEqual(
            clean_ocr_text(
                "**Extracted Text**\n\n正文\n\n**Content Type**  \n书籍\n\n**Quality Notes**  \n清晰"
            ),
            "正文",
        )

    def test_clean_translation_text_removes_model_notes(self) -> None:
        source = """修正说明：OCR 有破损，以下进行了整理。

第一段译文。

第二段译文。

***
*译者注：这不是原书内容。*
*[注1] 模型自行添加的解释。*
"""
        self.assertEqual(clean_translation_text(source), "第一段译文。\n\n第二段译文。")

    def test_clean_translation_text_removes_inline_transition_line(self) -> None:
        source = "第一段。\n\n在段能乐堂，翻译如下。\n\n第二段。"
        self.assertEqual(clean_translation_text(source), "第一段。\n\n第二段。")

    def test_publication_metadata_is_removed(self) -> None:
        source = """# 第一章

<!-- source-pdf: source.pdf -->
<!-- pdf-pages: 4-5 -->

<span epub:type="pagebreak" id="pdf-page-4" title="4"></span>
<!-- PDF_PAGE: 4 -->

第一页正文。

1

<span epub:type="pagebreak" id="pdf-page-5" title="5"></span>
<!-- PDF_PAGE: 5 -->

第二页正文。

2
"""
        result = strip_publication_metadata(source)
        self.assertEqual(result, "# 第一章\n\n第一页正文。\n\n第二页正文。\n")

    def test_publication_metadata_is_removed_when_only_pdf_page_comment_exists(self) -> None:
        source = """# 第一章

<!-- PDF_PAGE: 13 -->

13

后文。
"""
        self.assertEqual(
            strip_publication_metadata(source),
            "# 第一章\n\n后文。\n",
        )

    def test_publication_metadata_keeps_small_numbers_without_page_boundary(self) -> None:
        source = """# 章节

119

后文。
"""
        self.assertEqual(
            strip_publication_metadata(source),
            "# 章节\n\n119\n\n后文。\n",
        )

    def test_reader_output_removes_trailing_split_printed_page(self) -> None:
        # Column-aware OCR split printed page 40 into two lines; both are a
        # page footer and must be discarded, unlike mid-text numbers.
        source = """# 章节

正文。

4
0
"""
        self.assertEqual(
            strip_publication_metadata(source),
            "# 章节\n\n正文。\n",
        )

    def test_reader_output_removes_source_page_but_keeps_sections_and_years(self) -> None:
        source = """# 章节

正文。

119

后文。

2

1921

年表结束。
"""
        self.assertEqual(
            strip_publication_metadata(source),
            "# 章节\n\n正文。\n\n119\n\n后文。\n\n2\n\n1921\n\n年表结束。\n",
        )

    def test_contents_chapter_removes_navigation_page_numbers_only(self) -> None:
        source = """# 目录

第一章

6

第二章

42
"""
        self.assertEqual(
            strip_publication_metadata(source, chapter_title="目录"),
            "# 目录\n\n第一章\n\n第二章\n",
        )
        self.assertIn(
            "\n\n6\n",
            strip_publication_metadata(source, chapter_title="数据表"),
        )

    def test_known_two_page_spread_numbers_are_removed_from_publication(self) -> None:
        source = "正文。\n\n6\n\n后文。\n\n7\n\n42\n\n数据。"
        marked = annotate_printed_page_markers(source, [6, 7])
        self.assertIn('id="printed-page-6"', marked)
        self.assertEqual(
            strip_publication_metadata(marked),
            "正文。\n\n后文。\n\n42\n\n数据。\n",
        )

    def test_decorated_printed_page_numbers_are_removed_without_losing_body(self) -> None:
        source = "上一页未完\n\n27 ●\n\n●30出血的代价。\n\n© 330肃。\n\n⊙ 458闹。\n\n1968年正文。"
        marked = annotate_printed_page_markers(source, [27, 30, 330, 458])
        self.assertEqual(marked.count('epub:type="pagebreak"'), 4)
        self.assertNotIn("27 ●", marked)
        self.assertNotIn("●30", marked)
        self.assertNotIn("© 330", marked)
        self.assertNotIn("⊙ 458", marked)
        self.assertIn("出血的代价。", marked)
        self.assertIn("肃。", marked)
        self.assertIn("闹。", marked)
        self.assertIn("1968年正文。", marked)

    def test_decorated_page_marker_allows_cross_page_sentence_join(self) -> None:
        marked = annotate_printed_page_markers("付\n●30", [30])
        source = f"""# 章节

{marked}

<span epub:type="pagebreak" id="pdf-page-42" title="42"></span>
<!-- PDF_PAGE: 42 -->

出血的代价。
"""
        self.assertEqual(
            strip_publication_metadata(source),
            "# 章节\n\n付出血的代价。\n",
        )

    def test_reader_removes_numeric_header_only_at_page_boundary(self) -> None:
        source = """# 章节

正文中的数据。

119

<span epub:type="pagebreak" id="pdf-page-120" title="120"></span>
<!-- PDF_PAGE: 120 -->

120

下一页正文。
"""
        self.assertEqual(
            strip_publication_metadata(source),
            "# 章节\n\n正文中的数据。\n\n下一页正文。\n",
        )

    def test_reader_has_only_one_h1_per_chapter(self) -> None:
        source = "# 章节\n\n正文。\n\n# OCR 子标题\n\n后文。\n"
        self.assertEqual(
            strip_publication_metadata(source),
            "# 章节\n\n正文。\n\n## OCR 子标题\n\n后文。\n",
        )

    def test_reader_discards_bare_ocr_hash_heading(self) -> None:
        source = "# 章节\n\n正文。\n\n#\n\n后文。\n"
        self.assertEqual(
            strip_publication_metadata(source),
            "# 章节\n\n正文。\n\n后文。\n",
        )

    def test_reader_discards_hash_prefixed_printed_page_heading(self) -> None:
        source = "# 章节\n\n正文。\n\n#225\n\n后文。\n"
        self.assertEqual(
            strip_publication_metadata(source),
            "# 章节\n\n正文。\n\n后文。\n",
        )

    def test_reader_demotes_setext_and_html_h1(self) -> None:
        source = "# 章节\n\n子标题\n====\n\n<h1>另一子标题</h1>\n"
        self.assertEqual(
            strip_publication_metadata(source),
            "# 章节\n\n子标题\n---\n\n<h2>另一子标题</h2>\n",
        )

    def test_reader_keeps_year_and_only_discards_one_leading_page_number(self) -> None:
        source = """# 年表

1930

<span epub:type="pagebreak" id="pdf-page-12" title="12"></span>
<!-- PDF_PAGE: 12 -->

— 12 —

42

后文。
"""
        self.assertEqual(
            strip_publication_metadata(source),
            "# 年表\n\n1930\n\n42\n\n后文。\n",
        )

    def test_publication_running_title_is_removed(self) -> None:
        source = """# 消费／消灭之理——《魔法使之夜》与〈普遍经济〉

正文。

87 消费／消灭之理——《魔法使之夜》与〈普通经济〉｜会良ひめる

后文。
"""
        result = strip_publication_metadata(
            source,
            publication_title="消费／消灭之理——《魔法使之夜》与〈普遍经济〉",
        )
        self.assertEqual(
            result,
            "# 消费／消灭之理——《魔法使之夜》与〈普遍经济〉\n\n正文。\n\n后文。\n",
        )
        noisy = """# 消费／消灭之理——《魔法使之夜》与〈普遍经济〉

正文。

消费/消灭的理——《魔法使之夜》与〈普通经济〉[原文存疑] 會良ひめろ[原文存疑]
"""
        self.assertEqual(
            strip_publication_metadata(noisy, publication_title="消费／消灭之理——《魔法使之夜》与〈普遍经济〉"),
            "# 消费／消灭之理——《魔法使之夜》与〈普遍经济〉\n\n正文。\n",
        )

    def test_publication_title_inside_body_sentence_is_preserved(self) -> None:
        source = """# 第一章 历史想象力

中产阶级的孩子们

老左派衰落后，中产阶级的孩子们重新拿出了革命旗帜，这句话属于正文。
"""
        self.assertEqual(
            strip_publication_metadata(
                source,
                publication_title="中产阶级的孩子们：60年代与文化领导权",
            ),
            "# 第一章 历史想象力\n\n老左派衰落后，中产阶级的孩子们重新拿出了革命旗帜，这句话属于正文。\n",
        )

    def test_vertical_publication_running_title_is_removed(self) -> None:
        source = """# 主要参考书目

上一条书目。

<span epub:type="pagebreak" id="pdf-page-487" title="487"></span>
<!-- PDF_PAGE: 487 -->

中
产
阶
级
的
孩
子
们
M.E.Sharpe, Inc., 1983.
"""
        self.assertEqual(
            strip_publication_metadata(
                source,
                publication_title="中产阶级的孩子们：60年代与文化领导权",
            ),
            "# 主要参考书目\n\n上一条书目。\n\nM.E.Sharpe, Inc., 1983.\n",
        )

    def test_chapter_running_title_and_page_number_are_removed(self) -> None:
        source = """# 第一章 歌德的《浮士德》：发展的悲剧

正文。

第一章 歌德的《浮士德》：发展的悲剧 47

后文。
"""
        self.assertEqual(
            strip_publication_metadata(
                source,
                publication_title="一切坚固的东西都烟消云散了",
                chapter_title="第一章 歌德的《浮士德》：发展的悲剧",
            ),
            "# 第一章 歌德的《浮士德》：发展的悲剧\n\n正文。\n\n后文。\n",
        )

    def test_body_line_starting_with_book_title_is_kept(self) -> None:
        source = """# 二

包法利夫人回答道：

“汪洋一片，无边无涯。”
"""
        self.assertEqual(
            strip_publication_metadata(
                source,
                publication_title="包法利夫人",
                chapter_title="二",
            ),
            "# 二\n\n包法利夫人回答道：\n\n“汪洋一片，无边无涯。”\n",
        )

    def test_docx_link_list_keeps_inter_link_spaces(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            chapter_dir = output / "chapters"
            chapter_dir.mkdir(parents=True)
            filename = "001_目录.md"
            (chapter_dir / filename).write_text(
                "# 目录\n\n- [第一部](a.xhtml) [一](b.xhtml) [二](c.xhtml)\n",
                encoding="utf-8",
            )
            manifest = [
                {
                    "sequence": 1,
                    "id": "chapter",
                    "display_title": "目录",
                    "filename": filename,
                    "reviewed_override": False,
                }
            ]
            docx_path = output / "toc.docx"
            build_docx(docx_path, chapter_dir, manifest, book_title="书")

            from docx import Document

            document = Document(docx_path)
            paragraph = next(
                paragraph
                for paragraph in document.paragraphs
                if "第一部" in paragraph.text
            )
            self.assertEqual(paragraph.text, "第一部 一 二")

    def test_epub_subsection_heading_repeating_chapter_title_is_kept(self) -> None:
        source = """# 第一章 欲望机器

## 1. 欲望机器

正文。

## 2. 无器官身体

后文。
"""
        self.assertEqual(
            strip_publication_metadata(
                source,
                publication_title="反俄狄浦斯",
                chapter_title="第一章 欲望机器",
            ),
            "# 第一章 欲望机器\n\n"
            "## 1. 欲望机器\n\n"
            "正文。\n\n"
            "## 2. 无器官身体\n\n"
            "后文。\n",
        )

    def test_markdown_formatted_short_running_title_is_removed(self) -> None:
        source = """# 讲故事的人 论尼古拉·列斯科夫

正文。

# 启迪

后文。
"""
        self.assertEqual(
            strip_publication_metadata(
                source,
                publication_title="启迪 本雅明文选",
                chapter_title="讲故事的人 论尼古拉·列斯科夫",
            ),
            "# 讲故事的人 论尼古拉·列斯科夫\n\n正文。\n\n后文。\n",
        )

    def test_short_running_title_is_removed_before_cross_page_join(self) -> None:
        source = """# 企鹅版前言：宽广开放的理解方式

《地下室手

<span epub:type="pagebreak" id="pdf-page-14" title="14"></span>
<!-- PDF_PAGE: 14 -->

企鹅版前言 5

记》中继续讨论。
"""
        self.assertEqual(
            strip_publication_metadata(
                source,
                publication_title="一切坚固的东西都烟消云散了：现代性体验",
                chapter_title="企鹅版前言：宽广开放的理解方式",
            ),
            "# 企鹅版前言：宽广开放的理解方式\n\n《地下室手记》中继续讨论。\n",
        )
        self.assertEqual(
            strip_publication_metadata(
                "# 现代性研究译丛·总序\n\n正文。\n\n总序 3\n\n后文。",
                publication_title="一切坚固的东西都烟消云散了：现代性体验",
                chapter_title="现代性研究译丛·总序",
            ),
            "# 现代性研究译丛·总序\n\n正文。\n\n后文。\n",
        )

    def test_publication_joins_only_cross_page_sentence_fragments(self) -> None:
        source = """# 章节

这一页结束于作为对

25

<span epub:type="pagebreak" id="pdf-page-26" title="26"></span>
<!-- PDF_PAGE: 26 -->

象展开的讨论。

完整句子。

26

<span epub:type="pagebreak" id="pdf-page-27" title="27"></span>
<!-- PDF_PAGE: 27 -->

下一段另起。
"""
        self.assertEqual(
            strip_publication_metadata(source),
            "# 章节\n\n这一页结束于作为对象展开的讨论。\n\n完整句子。\n\n下一段另起。\n",
        )

    def test_publication_does_not_join_signed_note_to_next_page(self) -> None:
        source = """# 第一章

① Wells，美国电影演员。——译者

<span epub:type="pagebreak" id="pdf-page-60" title="60"></span>
<!-- PDF_PAGE: 60 -->

处于我的绝对控制之下！
"""
        self.assertEqual(
            strip_publication_metadata(source),
            "# 第一章\n\n① Wells，美国电影演员。——译者\n\n处于我的绝对控制之下！\n",
        )

    def test_publication_does_not_join_chronology_year_at_page_boundary(self) -> None:
        source = """# 年表

前一年条目 GS VI 185—187

<span epub:type="pagebreak" id="pdf-page-290" title="290"></span>
<!-- PDF_PAGE: 290 -->

启迪

1930

文学批评纲要
"""
        self.assertEqual(
            strip_publication_metadata(source, publication_title="启迪 本雅明文选"),
            "# 年表\n\n前一年条目 GS VI 185—187\n\n1930\n\n文学批评纲要\n",
        )

    def test_docx_inline_markup_becomes_plain_text(self) -> None:
        self.assertEqual(
            markdown_inline_to_plain_text("术语<sup>[原文存疑]</sup>与[链接](https://example.test)"),
            "术语[原文存疑]与链接",
        )

    def test_remove_multiline_duplicate_title(self) -> None:
        text = "第一章\n\n陀思妥耶夫斯基的复调\n小说和评论著述对它的阐释\n\n正文"
        title = "第一章 陀思妥耶夫斯基的复调小说和评论著述对它的阐释"
        self.assertEqual(remove_duplicate_title(text, title), "正文")

    def test_remove_multiline_duplicate_title_with_margin_page_number(self) -> None:
        text = "37 第一章 歌德的《浮士德》：\n发展的悲剧\n\n正文"
        title = "第一章 歌德的《浮士德》：发展的悲剧"
        self.assertEqual(remove_duplicate_title(text, title), "正文")

    def test_remove_duplicate_title_ignores_running_header_markup_and_particles(self) -> None:
        text = "小林秀雄全集第三卷\n\n**小说の问题 I**\n\n正文"
        self.assertEqual(remove_duplicate_title(text, "小说的问题 I"), "正文")

    def test_remove_duplicate_title_handles_reordered_descriptor(self) -> None:
        text = "84\n\n第八章\n\n批判细田守《无尽的斯嘉丽》\n\n正文"
        self.assertEqual(
            remove_duplicate_title(text, "细田守《无尽的斯嘉丽》批判"),
            "正文",
        )

    def test_parallel_translation_checkpoints(self) -> None:
        class FakeTranslator:
            def translate(self, text: str, *, source_language: str, target_language: str) -> str:
                return f"{target_language}:{text}"

        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            records = [
                PageRecord(1, "日本語の本文です。", language="ja"),
                PageRecord(2, "中文正文。", language="zh"),
                PageRecord(3, "別の日本語です。", language="ja"),
            ]
            translate_non_chinese_pages(
                records,
                output,
                FakeTranslator(),
                target_language="简体中文",
                force=False,
                concurrency=2,
                request_delay=0,
            )
            self.assertTrue((output / "pages/page_0001.json").exists())
            self.assertFalse((output / "pages/page_0002.json").exists())
            self.assertTrue(records[0].translated_text.startswith("简体中文:"))
            self.assertTrue(records[2].translated_text.startswith("简体中文:"))
            self.assertTrue(records[0].translation_is_fresh)
            self.assertTrue(records[2].translation_is_fresh)

    def test_translation_reruns_when_ocr_text_changed(self) -> None:
        class CountingTranslator:
            def __init__(self) -> None:
                self.calls = 0

            def translate(self, text: str, *, source_language: str, target_language: str) -> str:
                self.calls += 1
                return f"译文：{text}"

        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            record = PageRecord(1, "最初の本文", language="ja")
            translator = CountingTranslator()
            translate_non_chinese_pages(
                [record],
                output,
                translator,
                target_language="简体中文",
                force=False,
                concurrency=1,
            )
            self.assertEqual(translator.calls, 1)
            record.text = "修正後の本文"
            save_page_record(output, record)
            translate_non_chinese_pages(
                [record],
                output,
                translator,
                target_language="简体中文",
                force=False,
                concurrency=1,
            )
            self.assertEqual(translator.calls, 2)
            self.assertTrue(record.translation_is_fresh)

    def test_translation_reruns_when_model_changed(self) -> None:
        class CountingTranslator:
            def __init__(self) -> None:
                self.calls = 0

            def translate(self, text: str, *, source_language: str, target_language: str) -> str:
                self.calls += 1
                return f"译文：{text}"

        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            record = PageRecord(1, "日本語本文", language="ja")
            translator = CountingTranslator()
            translate_non_chinese_pages(
                [record],
                output,
                translator,
                target_language="简体中文",
                force=False,
                concurrency=1,
                translation_provider="deepseek",
                translation_model="deepseek-v4-flash",
            )
            translate_non_chinese_pages(
                [record],
                output,
                translator,
                target_language="简体中文",
                force=False,
                concurrency=1,
                translation_provider="deepseek",
                translation_model="deepseek-v4-pro",
            )
            self.assertEqual(translator.calls, 2)
            self.assertEqual(record.translation_provider, "deepseek")
            self.assertEqual(record.translation_model, "deepseek-v4-pro")
            self.assertEqual(record.translation_target_language, "简体中文")

    def test_import_does_not_mark_unverified_translation_fresh(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "source"
            output = root / "output"
            (source / "pages").mkdir(parents=True)
            (source / "pages/page_0001.json").write_text(
                json.dumps(
                    {
                        "pdf_page": 1,
                        "text": "更新后的OCR",
                        "language": "ja",
                        "translated_text": "旧译文",
                        "translation_source_sha256": "does-not-match",
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            self.assertEqual(import_existing_ocr(source, output), 1)
            record = PageRecord(
                **json.loads((output / "pages/page_0001.json").read_text(encoding="utf-8"))
            )
            self.assertFalse(record.translation_is_fresh)
            self.assertEqual(record.compile_text, "更新后的OCR")

    def test_translation_workers_actually_overlap(self) -> None:
        barrier = threading.Barrier(2)

        class BarrierTranslator:
            def translate(self, text: str, *, source_language: str, target_language: str) -> str:
                barrier.wait(timeout=2)
                return f"{target_language}:{text}"

        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            records = [
                PageRecord(1, "日本語一。", language="ja"),
                PageRecord(2, "日本語二。", language="ja"),
            ]
            translate_non_chinese_pages(
                records,
                output,
                BarrierTranslator(),
                target_language="简体中文",
                force=False,
                concurrency=2,
            )
            self.assertTrue(all(record.translated_text for record in records))

    def test_detect_language(self) -> None:
        self.assertEqual(detect_language("这是一本中文书籍，包含足够多的中文文字用于判断。"), "zh")
        self.assertEqual(detect_language("これは日本語の文章です。日本語として判定されます。"), "ja")
        self.assertEqual(detect_language("This is a sufficiently long English sentence for language detection."), "en")

    def test_coding_plan_vision_mcp_protocol(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            server = root / "fake_mcp.py"
            server.write_text(
                """
import json
import sys
for line in sys.stdin:
    message = json.loads(line)
    if "id" not in message:
        continue
    method = message.get("method")
    if method == "initialize":
        result = {"protocolVersion": "2024-11-05", "capabilities": {"tools": {}}, "serverInfo": {"name": "fake", "version": "1"}}
    elif method == "tools/list":
        result = {"tools": [{"name": "extract_text_from_screenshot", "inputSchema": {"type": "object", "properties": {"image_path": {"type": "string"}}, "required": ["image_path"]}}]}
    elif method == "tools/call":
        result = {"content": [{"type": "text", "text": "# 识别标题\\n\\n识别正文"}]}
    else:
        result = {}
    print(json.dumps({"jsonrpc": "2.0", "id": message["id"], "result": result}), flush=True)
""".strip()
                + "\n",
                encoding="utf-8",
            )
            image = root / "page.jpg"
            Image.new("RGB", (100, 160), "white").save(image)
            backend = CodingPlanVisionOCR(
                api_key="test-key",
                command=f"{sys.executable} -u {server}",
                reading_direction="vertical",
            )
            try:
                text, request_id = backend.ocr_image(image)
            finally:
                backend.close()
            self.assertEqual(text, "# 识别标题\n\n识别正文")
            self.assertTrue(request_id.startswith("mcp-"))
            self.assertEqual(
                backend.ocr_model,
                "coding-plan/glm-4.6v-vision-mcp/vertical-v2",
            )

    def test_coding_plan_content_filter_uses_segmented_fallback(self) -> None:
        class FilterThenReadClient:
            def extract_text(self, image_path: Path) -> tuple[str, str]:
                if "_segment_" not in image_path.stem:
                    raise RuntimeError(
                        'HTTP 400: {"contentFilter":[],"error":{"code":"1301",'
                        '"message":"potentially unsafe content"}}'
                    )
                match = re.search(r"_segment_(\d+)_", image_path.stem)
                assert match is not None
                values = {"1": "（无可见文字）", "2": "彼得", "3": "堡正文。", "4": "结尾。"}
                return values[match.group(1)], "mcp-test"

        with tempfile.TemporaryDirectory() as directory:
            image_path = Path(directory) / "page.jpg"
            Image.new("RGB", (240, 800), "white").save(image_path)
            backend = object.__new__(CodingPlanVisionOCR)
            with patch.object(backend, "_client", return_value=FilterThenReadClient()):
                text, request_id = backend.ocr_image(image_path)
            self.assertEqual(text, "彼得堡正文。\n\n结尾。")
            self.assertTrue(request_id.startswith("mcp-segmented-"))
            self.assertEqual(list(image_path.parent.glob("*_segment_*.jpg")), [])

    def test_coding_plan_timeout_splits_dense_page_instead_of_retrying_whole_page(self) -> None:
        class TimeoutThenReadClient:
            def __init__(self) -> None:
                self.whole_page_calls = 0

            def extract_text(self, image_path: Path) -> tuple[str, str]:
                if "_segment_" not in image_path.stem:
                    self.whole_page_calls += 1
                    raise RuntimeError("Vision MCP request timed out after 120 seconds.")
                match = re.search(r"_segment_(\d+)_", image_path.stem)
                assert match is not None
                return {"1": "右页正文。", "2": "左页正文。"}[match.group(1)], "mcp-test"

            def close(self) -> None:
                return

        with tempfile.TemporaryDirectory() as directory:
            image_path = Path(directory) / "dense-spread.jpg"
            Image.new("RGB", (400, 600), "white").save(image_path)
            backend = object.__new__(CodingPlanVisionOCR)
            backend.reading_direction = "vertical"
            client = TimeoutThenReadClient()
            with (
                patch.dict(
                    os.environ,
                    {
                        "CODING_PLAN_SPREAD_SEGMENTS": "2",
                        "CODING_PLAN_VERTICAL_PAGE_ROWS": "1",
                        "CODING_PLAN_VERTICAL_PAGE_COLUMNS": "1",
                    },
                ),
                patch.object(backend, "_client", return_value=client),
            ):
                text, request_id = backend.ocr_image(image_path)
            self.assertEqual(client.whole_page_calls, 1)
            self.assertEqual(text, "右页正文。\n\n左页正文。")
            self.assertTrue(request_id.startswith("mcp-segmented-"))
            self.assertEqual(list(image_path.parent.glob("*_segment_*.jpg")), [])

    def test_vertical_two_page_spread_is_split_before_first_model_call(self) -> None:
        class ReadSegmentClient:
            def __init__(self) -> None:
                self.whole_page_calls = 0

            def extract_text(self, image_path: Path) -> tuple[str, str]:
                if "_segment_" not in image_path.stem:
                    self.whole_page_calls += 1
                    raise AssertionError("the whole spread must not be submitted")
                match = re.search(r"_segment_(\d+)_", image_path.stem)
                assert match is not None
                return {"1": "右页。", "2": "左页。"}[match.group(1)], "mcp-test"

        with tempfile.TemporaryDirectory() as directory:
            image_path = Path(directory) / "spread.jpg"
            Image.new("RGB", (800, 400), "white").save(image_path)
            backend = object.__new__(CodingPlanVisionOCR)
            backend.reading_direction = "vertical"
            client = ReadSegmentClient()
            with (
                patch.dict(
                    os.environ,
                    {
                        "CODING_PLAN_SPREAD_SEGMENTS": "2",
                        "CODING_PLAN_VERTICAL_PAGE_ROWS": "1",
                        "CODING_PLAN_VERTICAL_PAGE_COLUMNS": "1",
                    },
                ),
                patch.object(backend, "_client", return_value=client),
            ):
                text, request_id = backend.ocr_image(image_path)
            self.assertEqual(client.whole_page_calls, 0)
            self.assertEqual(text, "右页。\n\n左页。")
            self.assertEqual(text.physical_page_texts, ("右页。", "左页。"))
            self.assertTrue(request_id.startswith("mcp-segmented-"))

    def test_segment_rate_limit_does_not_restart_completed_segments(self) -> None:
        class RateLimitedSegmentClient:
            def __init__(self) -> None:
                self.calls = {"1": 0, "2": 0}

            def extract_text(self, image_path: Path) -> tuple[str, str]:
                match = re.search(r"_segment_(\d+)_", image_path.stem)
                assert match is not None
                segment = match.group(1)
                self.calls[segment] += 1
                if segment == "2" and self.calls[segment] == 1:
                    raise RuntimeError("HTTP 429: Rate limit reached for requests")
                return {"1": "右页。", "2": "左页。"}[segment], "mcp-test"

        with tempfile.TemporaryDirectory() as directory:
            image_path = Path(directory) / "rate-limited-spread.jpg"
            Image.new("RGB", (800, 400), "white").save(image_path)
            backend = object.__new__(CodingPlanVisionOCR)
            backend.reading_direction = "vertical"
            client = RateLimitedSegmentClient()
            with (
                patch.dict(
                    os.environ,
                    {
                        "CODING_PLAN_SEGMENT_ATTEMPTS": "2",
                        "CODING_PLAN_SEGMENT_RATE_LIMIT_DELAY": "0",
                        "CODING_PLAN_SEGMENT_RATE_LIMIT_MAX_DELAY": "0",
                        "CODING_PLAN_SEGMENT_RATE_LIMIT_JITTER": "0",
                        "CODING_PLAN_VERTICAL_PAGE_ROWS": "1",
                        "CODING_PLAN_VERTICAL_PAGE_COLUMNS": "1",
                    },
                ),
                patch.object(backend, "_client", return_value=client),
            ):
                text, request_id = backend._ocr_segmented(
                    image_path,
                    client,
                    segments=2,
                )
            self.assertEqual(text, "右页。\n\n左页。")
            self.assertTrue(request_id.startswith("mcp-segmented-"))
            self.assertEqual(client.calls, {"1": 1, "2": 2})

    def test_vertical_segment_fallback_reads_right_to_left(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            image_path = Path(directory) / "vertical.jpg"
            image = Image.new("L", (240, 400), 0)
            image.paste(255, (116, 0, 240, 400))
            image.save(image_path)
            backend = object.__new__(CodingPlanVisionOCR)
            backend.reading_direction = "vertical"

            def read_band(path: Path, *, depth: int) -> tuple[str, str]:
                with Image.open(path).convert("L") as band:
                    average = sum(band.get_flattened_data()) / (band.width * band.height)
                    return ("右。" if average > 127 else "左。"), "test"

            with (
                patch.dict(
                    os.environ,
                    {
                        "CODING_PLAN_VERTICAL_PAGE_ROWS": "1",
                        "CODING_PLAN_VERTICAL_PAGE_COLUMNS": "1",
                    },
                ),
                patch.object(backend, "_ocr_filtered_band", side_effect=read_band),
            ):
                text, request_id = backend._ocr_segmented(
                    image_path,
                    object(),
                    segments=2,
                )
            self.assertEqual(text, "右。\n\n左。")
            self.assertTrue(request_id.startswith("mcp-segmented-"))
            self.assertEqual(list(image_path.parent.glob("*_segment_*.jpg")), [])

    def test_vertical_two_page_grid_reads_rows_then_columns_in_page_order(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            image_path = Path(directory) / "vertical-grid.jpg"
            image = Image.new("RGB", (800, 400), "white")
            colors = {
                "左页上左。": (255, 0, 0),
                "左页上右。": (0, 255, 0),
                "左页下左。": (0, 0, 255),
                "左页下右。": (255, 255, 0),
                "右页上左。": (255, 0, 255),
                "右页上右。": (0, 255, 255),
                "右页下左。": (128, 0, 0),
                "右页下右。": (0, 128, 0),
            }
            image.paste(colors["左页上左。"], (0, 0, 200, 200))
            image.paste(colors["左页上右。"], (200, 0, 400, 200))
            image.paste(colors["左页下左。"], (0, 200, 200, 400))
            image.paste(colors["左页下右。"], (200, 200, 400, 400))
            image.paste(colors["右页上左。"], (400, 0, 600, 200))
            image.paste(colors["右页上右。"], (600, 0, 800, 200))
            image.paste(colors["右页下左。"], (400, 200, 600, 400))
            image.paste(colors["右页下右。"], (600, 200, 800, 400))
            image.save(image_path, quality=100, subsampling=0)
            backend = object.__new__(CodingPlanVisionOCR)
            backend.reading_direction = "vertical"

            def read_grid_cell(path: Path, *, depth: int) -> tuple[str, str]:
                with Image.open(path).convert("RGB") as cell:
                    self.assertEqual(cell.size, (200, 200))
                    color = cell.getpixel((100, 100))
                label = min(
                    colors,
                    key=lambda candidate: sum(
                        abs(actual - expected)
                        for actual, expected in zip(color, colors[candidate])
                    ),
                )
                return label, "test"

            with (
                patch.dict(
                    os.environ,
                    {
                        "CODING_PLAN_VERTICAL_PAGE_ROWS": "2",
                        "CODING_PLAN_VERTICAL_PAGE_COLUMNS": "2",
                    },
                ),
                patch.object(
                    backend,
                    "_blank_row_cuts",
                    side_effect=lambda page, count: [page.height // 2],
                ),
                patch.object(
                    backend,
                    "_blank_column_cuts",
                    side_effect=lambda page, count: [page.width // 2],
                ),
                patch.object(backend, "_ocr_filtered_band", side_effect=read_grid_cell),
            ):
                text, request_id = backend._ocr_segmented(
                    image_path,
                    object(),
                    segments=2,
                )

            self.assertEqual(
                text,
                "\n\n".join(
                    (
                        "右页上右。",
                        "右页上左。",
                        "右页下右。",
                        "右页下左。",
                        "左页上右。",
                        "左页上左。",
                        "左页下右。",
                        "左页下左。",
                    )
                ),
            )
            self.assertTrue(request_id.startswith("mcp-segmented-"))
            self.assertEqual(list(image_path.parent.glob("*_segment_*.jpg")), [])

    def test_vertical_portrait_filtered_fallback_reads_top_then_bottom(self) -> None:
        class TimeoutThenReadRegionClient:
            def extract_text(self, image_path: Path) -> tuple[str, str]:
                if "_filtered_" not in image_path.stem:
                    raise RuntimeError("Vision MCP request timed out after 120 seconds.")
                with Image.open(image_path).convert("RGB") as region:
                    self.test_case.assertEqual(region.size, (120, 100))
                    red, _green, blue = region.getpixel((60, 50))
                    return ("上段。" if red > blue else "下段。"), "mcp-test"

            def close(self) -> None:
                return

        with tempfile.TemporaryDirectory() as directory:
            image_path = Path(directory) / "portrait-page.jpg"
            image = Image.new("RGB", (120, 200), "red")
            image.paste("blue", (0, 100, 120, 200))
            image.save(image_path, quality=100, subsampling=0)
            backend = object.__new__(CodingPlanVisionOCR)
            backend.reading_direction = "vertical"
            backend.local = threading.local()
            client = TimeoutThenReadRegionClient()
            client.test_case = self

            with (
                patch.object(backend, "_client", return_value=client),
                patch.object(backend, "_blank_row_cuts", return_value=[100]) as row_cuts,
                patch.object(
                    backend,
                    "_blank_column_cuts",
                    side_effect=AssertionError("portrait page must not be split into columns"),
                ),
            ):
                text, request_id = backend._ocr_filtered_band(image_path, depth=0)

            self.assertEqual(text, "上段。\n\n下段。")
            self.assertTrue(request_id.startswith("mcp-filtered-"))
            row_cuts.assert_called_once()
            self.assertEqual(list(image_path.parent.glob("*_filtered_*.jpg")), [])

    def test_short_physical_page_output_uses_ordered_fallback(self) -> None:
        class ShortThenReadRegionClient:
            def extract_text(self, image_path: Path) -> tuple[str, str]:
                if "_filtered_" not in image_path.stem:
                    return "短。", "mcp-short"
                with Image.open(image_path).convert("RGB") as region:
                    red, _green, blue = region.getpixel((60, 50))
                    return ("上段正文。" if red > blue else "下段正文。"), "mcp-test"

            def close(self) -> None:
                return

        with tempfile.TemporaryDirectory() as directory:
            image_path = Path(directory) / "short-portrait-page.jpg"
            image = Image.new("RGB", (120, 200), "red")
            image.paste("blue", (0, 100, 120, 200))
            image.save(image_path, quality=100, subsampling=0)
            backend = object.__new__(CodingPlanVisionOCR)
            backend.reading_direction = "vertical"
            backend.local = threading.local()
            client = ShortThenReadRegionClient()

            with (
                patch.dict(os.environ, {"CODING_PLAN_MIN_OCR_CHARS": "10"}),
                patch.object(backend, "_client", return_value=client),
                patch.object(backend, "_blank_row_cuts", return_value=[100]),
            ):
                text, request_id = backend._ocr_filtered_band(image_path, depth=0)

            self.assertEqual(text, "上段正文。\n\n下段正文。")
            self.assertTrue(request_id.startswith("mcp-filtered-"))
            self.assertEqual(list(image_path.parent.glob("*_filtered_*.jpg")), [])

    def test_vertical_narrow_filtered_fallback_reads_right_then_left(self) -> None:
        class TimeoutThenReadRegionClient:
            def extract_text(self, image_path: Path) -> tuple[str, str]:
                if "_filtered_" not in image_path.stem:
                    raise RuntimeError("Vision MCP request timed out after 120 seconds.")
                with Image.open(image_path).convert("RGB") as region:
                    self.test_case.assertEqual(region.size, (30, 240))
                    red, _green, blue = region.getpixel((15, 120))
                    return ("右列。" if red > blue else "左列。"), "mcp-test"

            def close(self) -> None:
                return

        with tempfile.TemporaryDirectory() as directory:
            image_path = Path(directory) / "narrow-column.jpg"
            image = Image.new("RGB", (60, 240), "blue")
            image.paste("red", (30, 0, 60, 240))
            image.save(image_path, quality=100, subsampling=0)
            backend = object.__new__(CodingPlanVisionOCR)
            backend.reading_direction = "vertical"
            backend.local = threading.local()
            client = TimeoutThenReadRegionClient()
            client.test_case = self

            with (
                patch.object(backend, "_client", return_value=client),
                patch.object(backend, "_blank_column_cuts", return_value=[30]) as column_cuts,
                patch.object(
                    backend,
                    "_blank_row_cuts",
                    side_effect=AssertionError("narrow vertical band must not be split into rows"),
                ),
            ):
                text, request_id = backend._ocr_filtered_band(image_path, depth=0)

            self.assertEqual(text, "右列。\n\n左列。")
            self.assertTrue(request_id.startswith("mcp-filtered-"))
            column_cuts.assert_called_once()
            self.assertEqual(list(image_path.parent.glob("*_filtered_*.jpg")), [])

    def test_vertical_grid_cell_fallback_keeps_right_to_left_columns(self) -> None:
        class TimeoutThenReadRegionClient:
            def extract_text(self, image_path: Path) -> tuple[str, str]:
                if "_filtered_" not in image_path.stem:
                    raise RuntimeError("Vision MCP request timed out after 120 seconds.")
                with Image.open(image_path).convert("RGB") as region:
                    red, _green, blue = region.getpixel((30, 100))
                    return ("右列。" if red > blue else "左列。"), "mcp-test"

            def close(self) -> None:
                return

        with tempfile.TemporaryDirectory() as directory:
            image_path = Path(directory) / "page_segment_2_grid.jpg"
            image = Image.new("RGB", (120, 200), "blue")
            image.paste("red", (60, 0, 120, 200))
            image.save(image_path, quality=100, subsampling=0)
            backend = object.__new__(CodingPlanVisionOCR)
            backend.reading_direction = "vertical"
            backend.local = threading.local()
            client = TimeoutThenReadRegionClient()

            with (
                patch.dict(os.environ, {"CODING_PLAN_VERTICAL_PAGE_ROWS": "2"}),
                patch.object(backend, "_client", return_value=client),
                patch.object(backend, "_blank_column_cuts", return_value=[60]),
                patch.object(
                    backend,
                    "_blank_row_cuts",
                    side_effect=AssertionError("grid row must retain vertical columns"),
                ),
            ):
                text, request_id = backend._ocr_filtered_band(image_path, depth=0)

            self.assertEqual(text, "右列。\n\n左列。")
            self.assertTrue(request_id.startswith("mcp-filtered-"))
            self.assertEqual(list(image_path.parent.glob("*_filtered_*.jpg")), [])

    def test_ocr_pdf_marks_visually_blank_unreadable_page_as_blank_checkpoint(self) -> None:
        class UnreadableBlankBackend:
            ocr_model = "glm-ocr"

            def ocr_image(self, image_path: Path) -> tuple[str, str]:
                return "The image is too blurry to read.", "blank-test"

            def close(self) -> None:
                return

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            pdf_path = root / "blank.pdf"
            document = fitz.open()
            document.new_page(width=200, height=200)
            document.save(pdf_path)
            document.close()

            records = ocr_pdf(
                pdf_path,
                root / "out",
                UnreadableBlankBackend(),
                start_page=1,
                end_page=1,
                concurrency=1,
                dpi=72,
                max_image_side=400,
                jpeg_quality=90,
                keep_page_images=False,
                force=True,
            )

            self.assertEqual(records[0].text, "[空白页]")
            self.assertEqual(records[0].ocr_model, "manual/visually-confirmed-blank")
            self.assertIn("visual_blank=true", records[0].notes)


class MappingAndCompilationTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.pdf_path = self.root / "sample.pdf"
        document = fitz.open()
        for page_number in range(1, 13):
            page = document.new_page()
            page.insert_text((72, 72), f"PDF page {page_number}")
        document.save(self.pdf_path)
        document.close()

    def tearDown(self) -> None:
        self.temp.cleanup()

    @staticmethod
    def sample_entries() -> list[TocEntry]:
        return [
            TocEntry("toc-1", "第一章", "起点", 1, "chapter", 1),
            TocEntry("toc-2", "第一节", "细节", 2, "section", 3),
            TocEntry("toc-2b", "第二节", "深入", 2, "section", 5),
            TocEntry("toc-3", "第二章", "终点", 1, "chapter", 6),
        ]

    @staticmethod
    def sample_records() -> list[PageRecord]:
        records = [PageRecord(page, f"这是 PDF 第 {page} 页的正文内容。") for page in range(1, 13)]
        records[3] = PageRecord(4, "# 第一章 起点\n第一章正文")
        records[5] = PageRecord(6, "## 第一节 细节\n分节正文")
        records[7] = PageRecord(8, "## 第二节 深入\n分节正文")
        records[8] = PageRecord(9, "# 第二章 终点\n第二章正文")
        return records

    def test_chapter_compile_removes_publication_running_title(self) -> None:
        entries = [
            TocEntry("chapter-1", "第一章", "历史想象力", 1, "chapter", 1, pdf_page=1),
        ]
        records = [
            PageRecord(1, "中产阶级的孩子们\n\n正文第一页。"),
            PageRecord(
                2,
                "中产阶级的孩子们\n\n正文提到中产阶级的孩子们参与了运动。",
            ),
        ]
        output = self.root / "publication-title-chapters"
        manifest, rows = compile_chapters(
            self.pdf_path,
            output,
            records,
            {"page_offset": 0, "entries": [entry.__dict__ for entry in entries]},
            granularity="chapter",
            publication_title="中产阶级的孩子们：60年代与文化领导权",
        )
        markdown = (output / "chapters" / manifest[0]["filename"]).read_text(
            encoding="utf-8"
        )
        self.assertNotIn("\n中产阶级的孩子们\n", markdown)
        self.assertIn("正文提到中产阶级的孩子们参与了运动。", markdown)
        self.assertTrue(rows)
        self.assertNotIn("\n中产阶级的孩子们\n", rows[0]["content"])
        self.assertIn("正文提到中产阶级的孩子们参与了运动。", rows[0]["content"])
        audit = json.loads(
            (output / "audit" / "semantic-reconstruction.json").read_text(
                encoding="utf-8"
            )
        )
        self.assertEqual(
            audit["chapters"][0]["markdown_sha256"],
            hashlib.sha256(
                (output / "chapters" / manifest[0]["filename"]).read_bytes()
            ).hexdigest(),
        )

    def test_infer_page_offset(self) -> None:
        offset, evidence = infer_page_offset(self.sample_entries(), self.sample_records(), toc_end=2)
        self.assertEqual(offset, 3)
        self.assertGreaterEqual(len(evidence), 3)

    def test_manual_offset_is_authoritative(self) -> None:
        payload = {
            "toc_pdf_pages": [2],
            "entries": [entry.__dict__ for entry in self.sample_entries()],
        }
        records = self.sample_records()
        records[3] = PageRecord(4, "无关内容")
        records[4] = PageRecord(5, "# 第一章 起点\n错误的标题匹配位置")
        mapped = apply_page_mapping(payload, records, page_offset=3)
        first = mapped["entries"][0]
        self.assertEqual(first["pdf_page"], 4)

    def test_title_evidence_handles_piecewise_page_offsets(self) -> None:
        entries = [
            TocEntry("preface", "", "前言", 1, "frontmatter", 1),
            TocEntry("chapter", "第一章", "起点", 1, "chapter", 6),
            TocEntry("ending", "", "结语", 1, "other", 10),
        ]
        records = [PageRecord(page, f"普通正文 {page}") for page in range(1, 13)]
        records[3] = PageRecord(4, "前言\n正文")
        records[7] = PageRecord(8, "第一章 起点\n正文")
        records[11] = PageRecord(12, "结语\n正文")
        payload = {"toc_pdf_pages": [2], "entries": [entry.__dict__ for entry in entries]}
        mapped = apply_page_mapping(payload, records, page_offset=None)
        self.assertEqual(mapped["page_offset"], 2)
        self.assertEqual([item["pdf_page"] for item in mapped["entries"]], [4, 8, 12])

    def test_frontmatter_before_toc_maps_by_title_and_stops_before_nested_chapter(self) -> None:
        entries = [
            TocEntry("ack", "", "致谢", 1, "frontmatter", 1),
            TocEntry("intro", "", "导言：批判的停顿：没有反对派的社会", 1, "frontmatter", 1),
            TocEntry("part", "", "单向度的社会", 1, "part", None),
            TocEntry("chapter", "第一章", "控制的新形式", 2, "chapter", 3),
        ]
        records = [PageRecord(page, f"普通正文 {page}") for page in range(1, 24)]
        records[4] = PageRecord(5, "致谢\n感谢正文")
        records[5] = PageRecord(6, "导言\n批判的停顿：没有\n反对派的社会\n导言正文")
        records[18] = PageRecord(
            19,
            "目录\n致谢\n001\n导言\n批判的停顿：没有反对派的社会\n第一章\n003\n控制的新形式",
        )
        records[19] = PageRecord(20, "单向度的社会")
        records[21] = PageRecord(22, "第一章\n控制的新形式\n正文第一页")
        payload = {"toc_pdf_pages": [19], "entries": [entry.__dict__ for entry in entries]}

        mapped = apply_page_mapping(payload, records, page_offset=None)
        self.assertEqual(mapped["page_offset"], 19)
        self.assertEqual(
            [item["pdf_page"] for item in mapped["entries"]],
            [5, 6, None, 22],
        )
        manifest, rows = compile_chapters(
            self.pdf_path,
            self.root / "frontmatter-before-toc",
            records,
            mapped,
            granularity="chapter",
        )

        self.assertEqual(
            [(item["display_title"], item["pdf_page"], item["end_pdf_page"]) for item in manifest],
            [
                ("致谢", 5, 5),
                ("导言：批判的停顿：没有反对派的社会", 6, 18),
                ("第一章 控制的新形式", 22, 23),
            ],
        )
        introduction = (
            self.root
            / "frontmatter-before-toc"
            / "chapters"
            / manifest[1]["filename"]
        ).read_text(encoding="utf-8")
        self.assertNotIn("目录", introduction)
        self.assertTrue(all(row["content"].strip() for row in rows))

    def test_mapping_rejects_early_title_mentions_and_excludes_all_toc_pages(self) -> None:
        entries = [
            TocEntry("intro", "", "引言：作为文化批判的艺术", 1, "frontmatter", 1),
            TocEntry("chapter-1", "一", "德国艺术家小说引言", 1, "chapter", 105),
            TocEntry("chapter-2", "二", "文化的肯定性质", 1, "chapter", 121),
            TocEntry("chapter-3", "三", "单向度社会中的艺术", 1, "chapter", 165),
            TocEntry("chapter-4", "四", "作为艺术品的社会", 1, "chapter", 181),
        ]
        records = [PageRecord(page, f"普通正文 {page}") for page in range(1, 201)]
        records[8] = PageRecord(9, "目录\n引言：作为文化批判的艺术\n001")
        records[9] = PageRecord(10, "目录（续）")
        records[10] = PageRecord(11, "引言：作为文化批判的艺术\n正文")
        records[70] = PageRecord(71, "单向度社会中的艺术\n正文中的早期回顾")
        records[114] = PageRecord(115, "一 德国艺术家小说引言\n正文")
        records[130] = PageRecord(131, "二 文化的肯定性质\n正文")
        records[174] = PageRecord(175, "三 单向度社会中的艺术\n正文首页")
        records[194] = PageRecord(195, "四 作为艺术品的社会\n正文首页")
        payload = {
            "toc_pdf_pages": [9, 10],
            "entries": [entry.__dict__ for entry in entries],
        }

        mapped = apply_page_mapping(payload, records, page_offset=None)

        self.assertEqual(mapped["page_offset"], 10)
        self.assertEqual(
            [item["pdf_page"] for item in mapped["entries"]],
            [11, 115, 131, 175, 195],
        )
        intro_evidence = next(
            item for item in mapped["offset_evidence"] if item["entry_id"] == "intro"
        )
        self.assertEqual(intro_evidence["pdf_page"], 11)

    def test_two_printed_pages_per_pdf_page_are_inferred(self) -> None:
        entries = [
            TocEntry("preface", "", "まえがき", 1, "frontmatter", None),
            TocEntry("chapter-1", "第一章", "新海誠論", 1, "chapter", 6),
            TocEntry("chapter-2", "第二章", "庵野秀明論", 1, "chapter", 17),
            TocEntry("chapter-3", "第三章", "細田守論", 1, "chapter", 26),
        ]
        records = [PageRecord(page, f"普通正文 {page}") for page in range(1, 16)]
        records[3] = PageRecord(4, "# 前書き\n本文")
        records[4] = PageRecord(5, "# 第一章 新海誠論\n本文")
        records[9] = PageRecord(10, "# 第二章 庵野秀明論\n本文")
        records[14] = PageRecord(15, "# 第三章 細田守論\n本文")
        payload = {"toc_pdf_pages": [3], "entries": [entry.__dict__ for entry in entries]}

        mapped = apply_page_mapping(
            payload,
            records,
            page_offset=None,
            source_page_count=47,
        )

        self.assertEqual(mapped["printed_pages_per_pdf_page"], 2)
        self.assertEqual(mapped["page_offset"], 2)
        self.assertEqual(
            [item["pdf_page"] for item in mapped["entries"]],
            [4, 5, 10, 15],
        )

    def test_two_page_spread_chapter_ranges_overlap(self) -> None:
        entries = [
            TocEntry("chapter-1", "第一章", "起点", 1, "chapter", 6, pdf_page=5),
            TocEntry("chapter-2", "第二章", "终点", 1, "chapter", 17, pdf_page=10),
        ]
        payload = {
            "page_offset": 2,
            "printed_pages_per_pdf_page": 2,
            "entries": [entry.__dict__ for entry in entries],
        }
        manifest, _rows = compile_chapters(
            self.pdf_path,
            self.root / "spread-chapters",
            self.sample_records(),
            payload,
            granularity="chapter",
        )
        self.assertEqual(
            [(item["pdf_page"], item["end_pdf_page"]) for item in manifest],
            [(5, 10), (10, 12)],
        )
        self.assertEqual(manifest[0]["boundary_mode"], "closed-overlap")

    def test_two_page_spread_even_boundary_does_not_overlap(self) -> None:
        entries = [
            TocEntry("chapter-1", "第一章", "起点", 1, "chapter", 6, pdf_page=5),
            TocEntry("chapter-2", "第二章", "终点", 1, "chapter", 16, pdf_page=10),
        ]
        payload = {
            "page_offset": 2,
            "printed_pages_per_pdf_page": 2,
            "entries": [entry.__dict__ for entry in entries],
        }
        manifest, _rows = compile_chapters(
            self.pdf_path,
            self.root / "even-spread-chapters",
            self.sample_records(),
            payload,
            granularity="chapter",
        )
        self.assertEqual(
            [(item["pdf_page"], item["end_pdf_page"]) for item in manifest],
            [(5, 9), (10, 12)],
        )
        self.assertEqual(manifest[0]["boundary_mode"], "non-overlap")

    def test_shared_spread_boundary_is_trimmed_for_both_chapters(self) -> None:
        entries = [
            TocEntry("chapter-1", "第一章", "起点", 1, "chapter", 6, pdf_page=5),
            TocEntry("chapter-2", "第二章", "终点", 1, "chapter", 17, pdf_page=10),
        ]
        payload = {
            "page_offset": 2,
            "printed_pages_per_pdf_page": 2,
            "entries": [entry.__dict__ for entry in entries],
        }
        records = self.sample_records()
        records[9] = PageRecord(
            10,
            "16\n上一章结尾\n\n17\n第二章\n终点\n下一章正文",
        )
        manifest, _rows = compile_chapters(
            self.pdf_path,
            self.root / "trimmed-spread-chapters",
            records,
            payload,
            granularity="chapter",
        )
        first = (self.root / "trimmed-spread-chapters" / "chapters" / manifest[0]["filename"]).read_text(
            encoding="utf-8"
        )
        second = (self.root / "trimmed-spread-chapters" / "chapters" / manifest[1]["filename"]).read_text(
            encoding="utf-8"
        )
        self.assertIn("上一章结尾", first)
        self.assertNotIn("下一章正文", first)
        self.assertIn("下一章正文", second)
        self.assertNotIn("上一章结尾", second)

    def test_ocr_checkpoints_and_page_markdown(self) -> None:
        class FakeOCR:
            ocr_model = "fake-ocr"

            def ocr_image(self, image_path: Path) -> tuple[str, str]:
                self.assert_image(image_path)
                return "# 页标题\n\n逐页正文", "request-test"

            @staticmethod
            def assert_image(image_path: Path) -> None:
                if not image_path.exists() or image_path.stat().st_size == 0:
                    raise AssertionError("Rendered page image is missing")

            def close(self) -> None:
                return

        output = self.root / "ocr-output"
        records = ocr_pdf(
            self.pdf_path,
            output,
            FakeOCR(),
            start_page=1,
            end_page=2,
            concurrency=2,
            dpi=96,
            max_image_side=1000,
            jpeg_quality=80,
            keep_page_images=False,
            force=False,
        )
        self.assertEqual([record.pdf_page for record in records], [1, 2])
        self.assertTrue((output / "pages" / "page_0001.json").exists())
        self.assertEqual((output / "pages" / "page_0001.md").read_text(encoding="utf-8"), "# 页标题\n\n逐页正文\n")

    def test_ocr_cache_model_prefix_replaces_fallback(self) -> None:
        class FakeOCR:
            def __init__(self, model: str) -> None:
                self.ocr_model = model

            def ocr_image(self, image_path: Path) -> tuple[str, str]:
                return f"由 {self.ocr_model} 识别", "request"

            def close(self) -> None:
                return

        output = self.root / "model-prefix-output"
        ocr_pdf(
            self.pdf_path,
            output,
            FakeOCR("tesseract/test"),
            start_page=1,
            end_page=1,
            concurrency=1,
            dpi=72,
            max_image_side=800,
            jpeg_quality=80,
            keep_page_images=False,
            force=False,
        )
        records = ocr_pdf(
            self.pdf_path,
            output,
            FakeOCR("coding-plan/test"),
            start_page=1,
            end_page=1,
            concurrency=1,
            dpi=72,
            max_image_side=800,
            jpeg_quality=80,
            keep_page_images=False,
            force=False,
            cache_model_prefix="coding-plan/",
        )
        self.assertEqual(records[0].ocr_model, "coding-plan/test")

    def test_ocr_exact_cache_identity_rejects_similar_prefix(self) -> None:
        class FakeOCR:
            def __init__(self, model: str) -> None:
                self.ocr_model = model
                self.calls = 0

            def ocr_image(self, _image_path: Path) -> tuple[str, str]:
                self.calls += 1
                return f"由 {self.ocr_model} 识别", "request"

            def close(self) -> None:
                return

        output = self.root / "model-exact-output"
        first = FakeOCR("tesseract/chi_sim/psm-30")
        ocr_pdf(
            self.pdf_path,
            output,
            first,
            start_page=1,
            end_page=1,
            concurrency=1,
            dpi=72,
            max_image_side=800,
            jpeg_quality=80,
            keep_page_images=False,
            force=False,
        )
        second = FakeOCR("tesseract/chi_sim/psm-3")
        records = ocr_pdf(
            self.pdf_path,
            output,
            second,
            start_page=1,
            end_page=1,
            concurrency=1,
            dpi=72,
            max_image_side=800,
            jpeg_quality=80,
            keep_page_images=False,
            force=False,
            cache_model_exact="tesseract/chi_sim/psm-3",
        )
        self.assertEqual(first.calls, 1)
        self.assertEqual(second.calls, 1)
        self.assertEqual(records[0].ocr_model, "tesseract/chi_sim/psm-3")

    def test_chapter_and_section_boundaries(self) -> None:
        payload = {
            "toc_pdf_pages": [2],
            "entries": [entry.__dict__ for entry in self.sample_entries()],
        }
        mapped = apply_page_mapping(payload, self.sample_records(), page_offset=3)

        chapter_output = self.root / "chapter-output"
        chapter_manifest, rows = compile_chapters(
            self.pdf_path,
            chapter_output,
            self.sample_records(),
            mapped,
            granularity="chapter",
        )
        self.assertEqual([(item["pdf_page"], item["end_pdf_page"]) for item in chapter_manifest], [(4, 8), (9, 12)])
        self.assertTrue(all(item["reviewed_override"] is False for item in chapter_manifest))
        self.assertTrue(rows)
        first_markdown = (chapter_output / "chapters" / chapter_manifest[0]["filename"]).read_text(encoding="utf-8")
        self.assertTrue(first_markdown.startswith("# 第一章 起点\n"))
        self.assertEqual(first_markdown.count("# 第一章 起点"), 1)
        self.assertNotIn("source-pdf", first_markdown)
        self.assertNotIn("PDF_PAGE", first_markdown)
        self.assertNotIn("pagebreak", first_markdown)

        section_output = self.root / "section-output"
        section_manifest, _ = compile_chapters(
            self.pdf_path,
            section_output,
            self.sample_records(),
            mapped,
            granularity="section",
        )
        self.assertEqual(
            [(item["pdf_page"], item["end_pdf_page"]) for item in section_manifest],
            [(6, 8), (8, 8)],
        )
        self.assertEqual(section_manifest[0]["boundary_mode"], "closed-overlap")

    def test_chapter_compilation_keeps_top_level_book_matter(self) -> None:
        entries = [
            TocEntry("preface", "", "中译本前言", 1, "frontmatter", 1, pdf_page=1),
            TocEntry("chapter", "第一章", "正文", 1, "chapter", 4, pdf_page=4),
            TocEntry("section", "第一节", "细节", 2, "section", 6, pdf_page=6),
            TocEntry("ending", "", "结语", 1, "other", 10, pdf_page=10),
        ]
        payload = {"page_offset": 0, "entries": [entry.__dict__ for entry in entries]}
        manifest, _ = compile_chapters(
            self.pdf_path,
            self.root / "complete-book",
            self.sample_records(),
            payload,
            granularity="chapter",
        )
        self.assertEqual(
            [item["display_title"] for item in manifest],
            ["中译本前言", "第一章 正文", "结语"],
        )
        self.assertEqual(
            [(item["pdf_page"], item["end_pdf_page"]) for item in manifest],
            [(1, 3), (4, 9), (10, 12)],
        )

    def test_compile_translation_gate_and_null_frontmatter_printed_page(self) -> None:
        record = PageRecord(
            1,
            "日本語本文",
            language="ja",
            translated_text="中文译文",
        )
        payload = {
            "page_offset": 10,
            "entries": [
                TocEntry(
                    "preface",
                    "",
                    "前言",
                    1,
                    "frontmatter",
                    None,
                    pdf_page=1,
                ).__dict__
            ],
        }
        with self.assertRaisesRegex(ValueError, "Missing or stale translation"):
            compile_chapters(
                self.pdf_path,
                self.root / "stale-translation",
                [record],
                payload,
                granularity="chapter",
                require_translation=True,
            )
        record.translation_source_sha256 = record.text_sha256
        record.translation_provider = "deepseek"
        record.translation_model = "deepseek-v4-pro"
        record.translation_target_language = "简体中文"
        _, rows = compile_chapters(
            self.pdf_path,
            self.root / "fresh-translation",
            [record],
            payload,
            granularity="chapter",
            require_translation=True,
        )
        self.assertNotIn("printed_page", rows[0])
        self.assertNotIn("source_pdf", rows[0])
        self.assertNotIn("pdf_page_start", rows[0])

        # The gate means "every page that needs translation", not every
        # nonblank page. A Chinese source page must compile without a no-op
        # Chinese-to-Chinese translation.
        chinese_record = PageRecord(1, "这是中文正文，已经可以直接发布。", language="zh")
        compile_chapters(
            self.pdf_path,
            self.root / "chinese-source-no-translation",
            [chinese_record],
            payload,
            granularity="chapter",
            require_translation=True,
        )

    def test_reviewed_chapter_override_replaces_output_and_knowledge_rows(self) -> None:
        output = self.root / "reviewed-override"
        reviewed_dir = output / "reviewed_chapters"
        reviewed_dir.mkdir(parents=True)
        reviewed_markdown = (
            "# 第一章 正文\n\n"
            "人工复核后的开篇。\n\n"
            "## 分论\n\n"
            "人工复核后的结论。\n"
        )
        (reviewed_dir / "chapter.md").write_text(
            reviewed_markdown,
            encoding="utf-8",
        )
        payload = {
            "page_offset": 0,
            "entries": [
                TocEntry(
                    "chapter",
                    "第一章",
                    "正文",
                    1,
                    "chapter",
                    1,
                    pdf_page=1,
                ).__dict__
            ],
        }
        manifest, rows = compile_chapters(
            self.pdf_path,
            output,
            [PageRecord(1, "逐页 OCR 坏文本", language="ja")],
            payload,
            granularity="chapter",
            require_translation=True,
        )

        rendered = (output / "chapters" / manifest[0]["filename"]).read_text(
            encoding="utf-8"
        )
        self.assertEqual(rendered, reviewed_markdown)
        self.assertTrue(manifest[0]["reviewed_override"])
        self.assertNotIn("source-pdf", rendered)
        self.assertNotIn("pdf-pages", rendered)
        self.assertNotIn("PDF_PAGE", rendered)
        self.assertNotIn("pagebreak", rendered)
        knowledge_text = "\n".join(row["content"] for row in rows)
        self.assertIn("人工复核后的开篇", knowledge_text)
        self.assertIn("人工复核后的结论", knowledge_text)
        self.assertNotIn("逐页 OCR 坏文本", knowledge_text)
        self.assertNotIn("# 第一章 正文", knowledge_text)
        for row in rows:
            self.assertNotIn("source_pdf", row)
            self.assertNotIn("pdf_page_start", row)
            self.assertNotIn("pdf_page_end", row)
            self.assertNotIn("printed_page", row)

    def test_reviewed_docx_preserves_inline_emphasis_and_underline(self) -> None:
        output = self.root / "reviewed-docx-inline"
        chapter_dir = output / "chapters"
        chapter_dir.mkdir(parents=True)
        filename = "001_第一章_正文.md"
        (chapter_dir / filename).write_text(
            "# 第一章 正文\n\n"
            "普通文字、<u>下划线</u>、**粗体**与*斜体*。\n",
            encoding="utf-8",
        )
        manifest = [
            {
                "sequence": 1,
                "id": "chapter",
                "display_title": "第一章 正文",
                "filename": filename,
                "reviewed_override": True,
            }
        ]
        docx_path = output / "inline.docx"
        build_docx(
            docx_path,
            chapter_dir,
            manifest,
            book_title="测试书",
        )

        from docx import Document

        document = Document(docx_path)
        body = next(
            paragraph
            for paragraph in document.paragraphs
            if paragraph.text.startswith("普通文字")
        )
        runs = {run.text: run for run in body.runs if run.text}
        self.assertTrue(runs["下划线"].underline)
        self.assertTrue(runs["粗体"].bold)
        self.assertTrue(runs["斜体"].italic)
        self.assertFalse(bool(runs["普通文字、"].underline))

    def test_docx_inline_runs_preserve_word_boundaries(self) -> None:
        output = self.root / "docx-inline-spacing"
        chapter_dir = output / "chapters"
        chapter_dir.mkdir(parents=True)
        filename = "001_chapter.md"
        (chapter_dir / filename).write_text(
            "# Chapter\n\nEnglish *emphasized* tail.\n",
            encoding="utf-8",
        )
        manifest = [
            {
                "sequence": 1,
                "id": "chapter",
                "display_title": "Chapter",
                "filename": filename,
                "reviewed_override": True,
            }
        ]
        docx_path = output / "inline-spacing.docx"
        build_docx(
            docx_path,
            chapter_dir,
            manifest,
            book_title="Test Book",
        )

        from docx import Document

        document = Document(docx_path)
        paragraph = next(
            paragraph
            for paragraph in document.paragraphs
            if paragraph.text.startswith("English")
        )
        self.assertEqual(paragraph.text, "English emphasized tail.")

    def test_docx_embeds_chapter_images(self) -> None:
        output = self.root / "docx-image-embed"
        chapter_dir = output / "chapters"
        media_dir = output / "media"
        media_dir.mkdir(parents=True)
        chapter_dir.mkdir(parents=True)
        Image.new("RGB", (480, 320), color=(255, 255, 255)).save(
            media_dir / "fig_small.png"
        )
        Image.new("RGB", (2400, 900), color=(0, 0, 0)).save(
            media_dir / "fig_wide.png"
        )
        filename = "001_第一章.md"
        (chapter_dir / filename).write_text(
            "# 第一章\n\n"
            "正文段落。\n\n"
            "![](../media/fig_small.png)\n\n"
            "![](../media/fig_wide.png)\n\n"
            "后续段落。\n",
            encoding="utf-8",
        )
        manifest = [
            {
                "sequence": 1,
                "id": "chapter",
                "display_title": "第一章",
                "filename": filename,
                "reviewed_override": False,
            }
        ]
        docx_path = output / "figures.docx"
        build_docx(docx_path, chapter_dir, manifest, book_title="测试书")

        from docx import Document

        document = Document(docx_path)
        self.assertEqual(len(document.inline_shapes), 2)
        widths = sorted(shape.width.cm for shape in document.inline_shapes)
        self.assertAlmostEqual(widths[0], 480 / 96 * 2.54, places=2)
        self.assertAlmostEqual(widths[1], 14.5, places=2)
        figure_paragraphs = [
            paragraph
            for paragraph in document.paragraphs
            if not paragraph.text.strip() and paragraph.runs
        ]
        centered = [
            paragraph
            for paragraph in figure_paragraphs
            if paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER
        ]
        self.assertEqual(len(centered), 2)
        for paragraph in centered:
            self.assertEqual(paragraph.paragraph_format.first_line_indent, 0)

    def test_docx_missing_image_fails_loudly(self) -> None:
        output = self.root / "docx-image-missing"
        chapter_dir = output / "chapters"
        chapter_dir.mkdir(parents=True)
        filename = "001_第一章.md"
        (chapter_dir / filename).write_text(
            "# 第一章\n\n![](../media/absent.png)\n",
            encoding="utf-8",
        )
        manifest = [
            {
                "sequence": 1,
                "id": "chapter",
                "display_title": "第一章",
                "filename": filename,
                "reviewed_override": False,
            }
        ]
        with self.assertRaises(FileNotFoundError):
            build_docx(
                output / "missing-figure.docx",
                chapter_dir,
                manifest,
                book_title="测试书",
            )

    def test_docx_wrap_lines_are_merged(self) -> None:
        output = self.root / "docx-wrap-merge"
        chapter_dir = output / "chapters"
        chapter_dir.mkdir(parents=True)
        filename = "001_第一章.md"
        (chapter_dir / filename).write_text(
            "# 第一章\n\n"
            "第一段文本，后面的内容被\n"
            "错误拆分成了两行。\n\n"
            "第二段。\n",
            encoding="utf-8",
        )
        manifest = [
            {
                "sequence": 1,
                "id": "chapter",
                "display_title": "第一章",
                "filename": filename,
                "reviewed_override": False,
            }
        ]
        docx_path = output / "wrapped.docx"
        build_docx(
            docx_path,
            chapter_dir,
            manifest,
            book_title="测试书",
        )

        from docx import Document

        document = Document(docx_path)
        paragraph = next(
            paragraph
            for paragraph in document.paragraphs
            if paragraph.text.startswith("第一段文本")
        )
        self.assertNotIn("\n", paragraph.text)
        self.assertEqual(paragraph.style.name, "Normal")
        self.assertEqual(
            paragraph._p.findall(
                ".//w:br",
                {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"},
            ),
            [],
        )

    def test_docx_explicit_html_break_is_preserved(self) -> None:
        output = self.root / "docx-explicit-break"
        chapter_dir = output / "chapters"
        chapter_dir.mkdir(parents=True)
        filename = "001_第一章.md"
        (chapter_dir / filename).write_text(
            "# 第一章\n\n第一行。<br>第二行。\n",
            encoding="utf-8",
        )
        manifest = [
            {
                "sequence": 1,
                "id": "chapter",
                "display_title": "第一章",
                "filename": filename,
                "reviewed_override": True,
            }
        ]
        docx_path = output / "explicit-break.docx"
        build_docx(
            docx_path,
            chapter_dir,
            manifest,
            book_title="测试书",
        )

        from docx import Document

        document = Document(docx_path)
        paragraph = next(
            paragraph
            for paragraph in document.paragraphs
            if paragraph.text.startswith("第一行")
        )
        self.assertEqual(paragraph.text, "第一行。\n第二行。")
        self.assertEqual(
            len(
                paragraph._p.findall(
                    ".//w:br",
                    {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"},
                )
            ),
            1,
        )

    def test_docx_omits_ocr_blank_page_labels(self) -> None:
        output = self.root / "docx-blank-page-label"
        chapter_dir = output / "chapters"
        chapter_dir.mkdir(parents=True)
        filename = "001_第一章.md"
        (chapter_dir / filename).write_text(
            "# 第一章\n\n正文第一段。\n\n[空白页]\n\n正文第二段。\n",
            encoding="utf-8",
        )
        manifest = [
            {
                "sequence": 1,
                "id": "chapter",
                "display_title": "第一章",
                "filename": filename,
                "reviewed_override": False,
            }
        ]
        docx_path = output / "blank-page-label.docx"
        build_docx(
            docx_path,
            chapter_dir,
            manifest,
            book_title="测试书",
        )

        from docx import Document

        document = Document(docx_path)
        texts = [paragraph.text for paragraph in document.paragraphs]
        self.assertNotIn("[空白页]", texts)
        self.assertIn("正文第一段。", texts)
        self.assertIn("正文第二段。", texts)

    def test_docx_omits_duplicate_epub_titlepage_and_cover_scaffolds(self) -> None:
        output = self.root / "docx-epub-scaffolds"
        chapter_dir = output / "chapters"
        chapter_dir.mkdir(parents=True)
        chapters = (
            ("001_titlepage.md", "# titlepage\n", "titlepage"),
            (
                "002_Cover_Image_cover_jpeg.md",
                "# ![Cover Image](../../cover.jpeg)\n",
                "![Cover Image](../../cover.jpeg)",
            ),
            ("003_第一章.md", "# 第一章\n\n正文。\n", "第一章"),
        )
        manifest = []
        for sequence, (filename, markdown, title) in enumerate(chapters, start=1):
            (chapter_dir / filename).write_text(markdown, encoding="utf-8")
            manifest.append(
                {
                    "sequence": sequence,
                    "id": f"epub-{sequence:04d}",
                    "display_title": title,
                    "filename": filename,
                    "source_href": f"EPUB/xhtml/{filename[:-3]}.xhtml",
                    "reviewed_override": False,
                }
            )
        docx_path = output / "epub-scaffolds.docx"
        build_docx(
            docx_path,
            chapter_dir,
            manifest,
            book_title="测试书",
        )

        from docx import Document

        document = Document(docx_path)
        self.assertEqual(document.core_properties.author, "")
        headings = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.style.name == "Heading 1"
        ]
        self.assertEqual(headings, ["第一章"])

    def test_docx_book_layout_has_controlled_front_matter_and_footer(self) -> None:
        output = self.root / "docx-book-layout"
        chapter_dir = output / "chapters"
        chapter_dir.mkdir(parents=True)
        chapters = (
            ("001_第一章.md", "# 第一章\n\n正文。\n"),
            ("002_第二章.md", "# 第二章\n\n后文。\n"),
        )
        for filename, markdown in chapters:
            (chapter_dir / filename).write_text(markdown, encoding="utf-8")
        manifest = [
            {
                "sequence": index,
                "id": f"chapter-{index}",
                "display_title": f"第{'一' if index == 1 else '二'}章",
                "filename": filename,
                "reviewed_override": True,
            }
            for index, (filename, _markdown) in enumerate(chapters, start=1)
        ]
        path = output / "book.docx"
        build_docx(
            path,
            chapter_dir,
            manifest,
            book_title="测试书",
            author="测试作者",
        )

        from docx import Document

        document = Document(path)
        self.assertEqual(document.paragraphs[0].style.name, "Codex Book Title")
        self.assertEqual(document.paragraphs[0].text, "测试书")
        self.assertEqual(document.paragraphs[1].text, "测试作者")
        self.assertEqual(document.core_properties.author, "测试作者")
        self.assertEqual(
            [
                paragraph.text
                for paragraph in document.paragraphs
                if paragraph.style.name == "Heading 1"
            ],
            ["第一章", "第二章"],
        )

        namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        with zipfile.ZipFile(path) as archive:
            body = ET.fromstring(archive.read("word/document.xml"))
            styles = ET.fromstring(archive.read("word/styles.xml"))
            footer = ET.fromstring(archive.read("word/footer1.xml"))
        self.assertEqual(len(body.findall(".//w:br[@w:type='page']", namespace)), 1)
        heading_style = styles.find(
            ".//w:style[@w:styleId='Heading1']", namespace
        )
        normal_style = styles.find(
            ".//w:style[@w:styleId='Normal']", namespace
        )
        footnote_reference_style = styles.find(
            ".//w:style[@w:styleId='FootnoteReference']", namespace
        )
        source_note_style = styles.find(
            ".//w:style[@w:styleId='SourceNote']", namespace
        )
        self.assertIsNotNone(heading_style)
        self.assertIsNotNone(normal_style)
        self.assertIsNotNone(footnote_reference_style)
        self.assertIsNotNone(source_note_style)
        self.assertIsNotNone(heading_style.find(".//w:pageBreakBefore", namespace))
        self.assertEqual(
            normal_style.find("w:rPr/w:rFonts", namespace).get(
                f"{{{namespace['w']}}}eastAsia"
            ),
            "SimSun",
        )
        self.assertEqual(
            heading_style.find("w:rPr/w:rFonts", namespace).get(
                f"{{{namespace['w']}}}eastAsia"
            ),
            "Microsoft YaHei",
        )
        self.assertEqual(
            normal_style.find("w:rPr/w:spacing", namespace).get(
                f"{{{namespace['w']}}}val"
            ),
            "0",
        )
        self.assertEqual(
            normal_style.find("w:pPr/w:jc", namespace).get(
                f"{{{namespace['w']}}}val"
            ),
            "both",
        )
        self.assertEqual(
            footnote_reference_style.find("w:rPr/w:vertAlign", namespace).get(
                f"{{{namespace['w']}}}val"
            ),
            "superscript",
        )
        self.assertEqual(
            source_note_style.find("w:pPr/w:jc", namespace).get(
                f"{{{namespace['w']}}}val"
            ),
            "left",
        )
        self.assertEqual(
            [field.get(f"{{{namespace['w']}}}instr") for field in footer.findall(".//w:fldSimple", namespace)],
            ["PAGE"],
        )

    def test_docx_book_layout_styles_backmatter_and_table_geometry(self) -> None:
        output = self.root / "docx-book-backmatter"
        chapter_dir = output / "chapters"
        chapter_dir.mkdir(parents=True)
        chapters = (
            ("001_主要参考书目.md", "# 主要参考书目\n\nSmith, A., A Book, 2001.\n"),
            ("002_索引.md", "# 索引\n\n阿伦特 12, 18\n"),
            (
                "003_表格.md",
                "# 表格\n\n| 项目 | 很长的说明列 |\n| --- | --- |\n| A | 说明文字 |\n",
            ),
        )
        manifest = [
            {
                "sequence": index,
                "id": identity,
                "display_title": title,
                "filename": filename,
                "reviewed_override": True,
            }
            for index, (identity, title, (filename, _markdown)) in enumerate(
                zip(
                    ("bibliography", "index", "chapter"),
                    ("主要参考书目", "索引", "表格"),
                    chapters,
                ),
                start=1,
            )
        ]
        for filename, markdown in chapters:
            (chapter_dir / filename).write_text(markdown, encoding="utf-8")
        path = output / "book.docx"
        build_docx(path, chapter_dir, manifest, book_title="测试书")

        from docx import Document

        document = Document(path)
        paragraphs = {paragraph.text: paragraph.style.name for paragraph in document.paragraphs}
        self.assertEqual(paragraphs["Smith, A., A Book, 2001."], "Bibliography Entry")
        self.assertEqual(paragraphs["阿伦特 12, 18"], "Index Entry")

        namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        attr = lambda name: f"{{{namespace['w']}}}{name}"
        with zipfile.ZipFile(path) as archive:
            body = ET.fromstring(archive.read("word/document.xml"))
        table = body.find(".//w:tbl", namespace)
        self.assertIsNotNone(table)
        table_width = table.find("w:tblPr/w:tblW", namespace)
        table_indent = table.find("w:tblPr/w:tblInd", namespace)
        grid = [
            int(column.get(attr("w")))
            for column in table.findall("w:tblGrid/w:gridCol", namespace)
        ]
        self.assertEqual(table_width.get(attr("type")), "dxa")
        self.assertEqual(table_indent.get(attr("w")), "120")
        self.assertEqual(sum(grid), int(table_width.get(attr("w"))))
        for row in table.findall("w:tr", namespace):
            self.assertEqual(
                [
                    int(cell.find("w:tcPr/w:tcW", namespace).get(attr("w")))
                    for cell in row.findall("w:tc", namespace)
                ],
                grid,
            )

    def test_reviewed_override_round_trips_middle_dot_title_and_body(self) -> None:
        output = self.root / "reviewed-middle-dot"
        reviewed_dir = output / "reviewed_chapters"
        reviewed_dir.mkdir(parents=True)
        title = "米哈伊尔·罗亚·巴尔达姆约恩研究"
        reviewed_markdown = (
            f"# {title}\n\n"
            "## 罗亚的生平\n\n"
            "罗亚仍是正文中必须反复出现的合法人物名。\n\n"
            "## 樹木\n\n"
            "> 合法引文。\n>\n> ——《月姬》\n"
        )
        (reviewed_dir / "chapter.md").write_text(
            reviewed_markdown,
            encoding="utf-8",
        )
        payload = {
            "entries": [
                TocEntry(
                    "chapter",
                    "",
                    title,
                    1,
                    "chapter",
                    1,
                    pdf_page=1,
                ).__dict__
            ]
        }

        manifest, rows = compile_chapters(
            self.pdf_path,
            output,
            [PageRecord(1, "不会进入成品的 OCR 文本")],
            payload,
            granularity="chapter",
        )

        rendered = (output / "chapters" / manifest[0]["filename"]).read_text(
            encoding="utf-8"
        )
        self.assertEqual(rendered, reviewed_markdown)
        self.assertIn("## 罗亚的生平", rendered)
        self.assertIn("## 樹木", rendered)
        self.assertIn("> ——《月姬》", rendered)
        knowledge_text = "\n".join(row["content"] for row in rows)
        self.assertIn("罗亚仍是正文", knowledge_text)
        self.assertIn("——《月姬》", knowledge_text)

        epub_path = output / "reviewed.epub"
        build_epub(
            epub_path,
            output / "chapters",
            manifest,
            book_title="测试书",
            language="zh-CN",
        )
        with zipfile.ZipFile(epub_path) as archive:
            epub_text = "\n".join(
                archive.read(name).decode("utf-8", errors="ignore")
                for name in archive.namelist()
                if name.endswith(".xhtml")
            )
        self.assertIn("罗亚仍是正文", epub_text)
        self.assertIn("罗亚的生平", epub_text)
        self.assertIn("——《月姬》", epub_text)

        docx_path = output / "reviewed.docx"
        build_docx(
            docx_path,
            output / "chapters",
            manifest,
            book_title="测试书",
        )
        from docx import Document

        document = Document(docx_path)
        docx_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("罗亚仍是正文", docx_text)
        self.assertIn("罗亚的生平", docx_text)
        self.assertIn("——《月姬》", docx_text)

    def test_reviewed_override_strips_pdf_page_comment_number(self) -> None:
        source = """# 第一章 正文

<!-- PDF_PAGE: 7 -->

7

正文。
"""
        result = strip_reviewed_publication_metadata(source)
        self.assertEqual(result, "# 第一章 正文\n\n正文。\n")

    def test_reviewed_override_strips_only_explicit_source_page_markers(self) -> None:
        output = self.root / "reviewed-explicit-markers"
        reviewed_dir = output / "reviewed_chapters"
        reviewed_dir.mkdir(parents=True)
        reviewed_markdown = (
            "# 第一章 正文\n\n"
            "<!-- source-pdf: source.pdf -->\n"
            "<!-- pdf-pages: 1-1 -->\n"
            '<span epub:type="pagebreak" id="pdf-page-1" title="1"></span>\n'
            "<!-- PDF_PAGE: 1 -->\n\n"
            "1\n\n"
            '<span id="pdf-page-2" title="2" epub:type="pagebreak"></span>\n'
            "<!-- PDF_PAGE: 2 -->\n\n"
            "2022\n\n"
            "## 合法小标题\n\n"
            "> 合法引文第一段。\n>\n> ——原著第二段。\n\n"
            "Twitter ID：@Grand_Order_RTA\n\n"
            "| 名称 | 值 |\n"
            "| --- | --- |\n"
            "| 罗亚 | 保留 |\n\n"
            "1. 第一条脚注。\n"
            "2. 第二条脚注。\n"
        )
        (reviewed_dir / "chapter.md").write_text(
            reviewed_markdown,
            encoding="utf-8",
        )
        payload = {
            "entries": [
                TocEntry(
                    "chapter",
                    "第一章",
                    "正文",
                    1,
                    "chapter",
                    1,
                    pdf_page=1,
                ).__dict__
            ]
        }

        manifest, rows = compile_chapters(
            self.pdf_path,
            output,
            [PageRecord(1, "OCR 正文")],
            payload,
            granularity="chapter",
        )

        rendered = (output / "chapters" / manifest[0]["filename"]).read_text(
            encoding="utf-8"
        )
        self.assertNotIn("source-pdf", rendered)
        self.assertNotIn("pdf-pages", rendered)
        self.assertNotIn("PDF_PAGE", rendered)
        self.assertNotIn("pagebreak", rendered)
        self.assertNotIn("\n1\n", rendered)
        self.assertIn("\n2022\n", rendered)
        self.assertIn("## 合法小标题", rendered)
        self.assertIn("> ——原著第二段。", rendered)
        self.assertIn("@Grand_Order_RTA", rendered)
        self.assertIn("| 罗亚 | 保留 |", rendered)

        knowledge_text = "\n".join(row["content"] for row in rows)
        for marker in ("source-pdf", "pdf-pages", "PDF_PAGE", "pagebreak"):
            self.assertNotIn(marker, knowledge_text)
        self.assertIn("2022", knowledge_text)
        self.assertIn("合法引文第一段", knowledge_text)
        self.assertIn("@Grand_Order_RTA", knowledge_text)
        self.assertIn("| 罗亚 | 保留 |", knowledge_text)

        epub_path = output / "reviewed-explicit-markers.epub"
        build_epub(
            epub_path,
            output / "chapters",
            manifest,
            book_title="测试书",
            language="zh-CN",
        )
        with zipfile.ZipFile(epub_path) as archive:
            epub_text = "\n".join(
                archive.read(name).decode("utf-8", errors="ignore")
                for name in archive.namelist()
                if name.endswith(".xhtml") and name != "OEBPS/nav.xhtml"
            )
        for marker in ("source-pdf", "pdf-pages", "PDF_PAGE", "pagebreak"):
            self.assertNotIn(marker, epub_text)
        self.assertIn("2022", epub_text)
        self.assertIn("<blockquote>", epub_text)
        self.assertIn("@Grand_Order_RTA", epub_text)
        self.assertIn("<table>", epub_text)

        docx_path = output / "reviewed-explicit-markers.docx"
        build_docx(
            docx_path,
            output / "chapters",
            manifest,
            book_title="测试书",
        )
        from docx import Document

        document = Document(docx_path)
        docx_paragraphs = [paragraph.text for paragraph in document.paragraphs]
        docx_text = "\n".join(docx_paragraphs)
        for marker in ("source-pdf", "pdf-pages", "PDF_PAGE", "pagebreak"):
            self.assertNotIn(marker, docx_text)
        self.assertIn("2022", docx_paragraphs)
        self.assertIn("Twitter ID：@Grand_Order_RTA", docx_paragraphs)
        self.assertIn("合法引文第一段。", docx_paragraphs)
        self.assertIn("——原著第二段。", docx_paragraphs)
        self.assertFalse(any(text.startswith(">") for text in docx_paragraphs))
        numbered = [
            paragraph
            for paragraph in document.paragraphs
            if paragraph.style.name.startswith("List Number")
        ]
        self.assertEqual(
            [paragraph.text for paragraph in numbered],
            ["第一条脚注。", "第二条脚注。"],
        )
        self.assertEqual(len(document.tables), 1)
        self.assertEqual(
            [[cell.text for cell in row.cells] for row in document.tables[0].rows],
            [["名称", "值"], ["罗亚", "保留"]],
        )

    def test_reviewed_publication_metadata_removes_standalone_printed_numbers(self) -> None:
        source = """# 第一章

<span epub:type="pagebreak" id="pdf-page-1" title="1"></span>
<!-- PDF_PAGE: 1 -->

序言第一段会出现明显正文交接，需要修复这一行的排版。

1

2

接续正文，这里开始后续段落。

2022

"""
        cleaned = strip_reviewed_publication_metadata(source)
        self.assertIn("序言第一段会出现明显正文交接，需要修复这一行的排版。", cleaned)
        self.assertIn("接续正文，这里开始后续段落。", cleaned)
        self.assertIn("2022", cleaned)
        self.assertNotIn("\n1\n", cleaned)
        self.assertNotIn("\n2\n", cleaned)

    def test_docx_reviewed_wrap_and_alignment_are_normalized(self) -> None:
        output = self.root / "docx-review-layout"
        chapter_dir = output / "chapters"
        chapter_dir.mkdir(parents=True)
        filename = "001_第一章.md"
        (chapter_dir / filename).write_text(
            """# 第一章

<span epub:type="pagebreak" id="pdf-page-1" title="1"></span>
<!-- PDF_PAGE: 1 -->

1

正文开头第一段正文交接到下一行。

<span epub:type="pagebreak" id="pdf-page-2" title="2"></span>
<!-- PDF_PAGE: 2 -->

2

正文第二段承接第一页内容。

""",
            encoding="utf-8",
        )
        manifest = [
            {
                "sequence": 1,
                "id": "chapter",
                "display_title": "第一章",
                "filename": filename,
                "reviewed_override": True,
            }
        ]
        docx_path = output / "book.docx"
        build_docx(docx_path, chapter_dir, manifest, book_title="测试书")

        from docx import Document

        document = Document(docx_path)
        docx_paragraphs = [paragraph.text for paragraph in document.paragraphs]
        self.assertNotIn("1", docx_paragraphs)
        self.assertNotIn("2", docx_paragraphs)
        self.assertIn("正文开头第一段正文交接到下一行。", docx_paragraphs)
        self.assertIn("正文第二段承接第一页内容。", docx_paragraphs)
        self.assertEqual(document.styles["Normal"].paragraph_format.alignment, WD_ALIGN_PARAGRAPH.JUSTIFY)
        self.assertEqual(document.styles["Quote"].paragraph_format.alignment, WD_ALIGN_PARAGRAPH.JUSTIFY)

    def test_docx_source_notes_are_separated_from_wrapped_body(self) -> None:
        output = self.root / "docx-source-notes"
        chapter_dir = output / "chapters"
        chapter_dir.mkdir(parents=True)
        filename = "001_第一章.md"
        (chapter_dir / filename).write_text(
            """# 第一章

正文通过把变
Bazard,Doctrine Saint-Simonienne-Exposition,Paris1854,pp.123f.145.
①
Ibid.,p.124.
②
Ibid.,p.127.
③迁视为实存的特有形式，正文继续。

①第二次世界大战时希特勒曾建立集中营，监
禁和屠杀爱国者和战俘。——译者

后文。
""",
            encoding="utf-8",
        )
        manifest = [
            {
                "sequence": 1,
                "id": "chapter",
                "display_title": "第一章",
                "filename": filename,
                "reviewed_override": True,
            }
        ]
        docx_path = output / "book.docx"
        build_docx(docx_path, chapter_dir, manifest, book_title="测试书")

        from docx import Document

        document = Document(docx_path)
        source_notes = [
            paragraph
            for paragraph in document.paragraphs
            if paragraph.style.name == "Source Note"
        ]
        self.assertEqual(len(source_notes), 2)
        self.assertIn("Bazard,Doctrine", source_notes[0].text)
        self.assertIn("① Ibid.,p.124.", source_notes[0].text)
        self.assertIn("——译者", source_notes[1].text)
        self.assertEqual(
            document.styles["Source Note"].paragraph_format.alignment,
            WD_ALIGN_PARAGRAPH.LEFT,
        )
        normal_text = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.style.name == "Normal"
        ]
        self.assertIn("正文通过把变", normal_text)
        self.assertIn("③迁视为实存的特有形式，正文继续。", normal_text)

    def test_reviewed_override_still_requires_ocr_range(self) -> None:
        output = self.root / "reviewed-missing-ocr"
        reviewed_dir = output / "reviewed_chapters"
        reviewed_dir.mkdir(parents=True)
        (reviewed_dir / "chapter.md").write_text(
            "# 第一章 正文\n\n人工复核正文。\n",
            encoding="utf-8",
        )
        payload = {
            "entries": [
                TocEntry(
                    "chapter",
                    "第一章",
                    "正文",
                    1,
                    "chapter",
                    1,
                    pdf_page=1,
                ).__dict__
            ]
        }

        with self.assertRaisesRegex(ValueError, "Missing OCR pages"):
            compile_chapters(
                self.pdf_path,
                output,
                [PageRecord(2, "第二页")],
                payload,
                granularity="chapter",
                require_translation=True,
            )

    def test_invalid_reviewed_override_preserves_previous_chapters(self) -> None:
        output = self.root / "invalid-reviewed-override"
        chapter_dir = output / "chapters"
        chapter_dir.mkdir(parents=True)
        previous = chapter_dir / "001_previous.md"
        previous.write_text("# 上次成功产物\n", encoding="utf-8")
        reviewed_dir = output / "reviewed_chapters"
        reviewed_dir.mkdir(parents=True)
        override_path = reviewed_dir / "chapter.md"
        override_path.write_text("# 错误标题\n\n正文。\n", encoding="utf-8")
        payload = {
            "entries": [
                TocEntry(
                    "chapter",
                    "第一章",
                    "正文",
                    1,
                    "chapter",
                    1,
                    pdf_page=1,
                ).__dict__
            ]
        }

        with self.assertRaisesRegex(ValueError, "H1 must match"):
            compile_chapters(
                self.pdf_path,
                output,
                [PageRecord(1, "OCR 正文")],
                payload,
                granularity="chapter",
            )
        self.assertEqual(previous.read_text(encoding="utf-8"), "# 上次成功产物\n")

        override_path.write_text("\n", encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "override is empty"):
            compile_chapters(
                self.pdf_path,
                output,
                [PageRecord(1, "OCR 正文")],
                payload,
                granularity="chapter",
            )
        self.assertEqual(previous.read_text(encoding="utf-8"), "# 上次成功产物\n")

        override_path.write_text(
            "# 第一章 正文\n\n"
            "<!-- source-pdf: source.pdf -->\n"
            "<!-- pdf-pages: 1-1 -->\n"
            '<span epub:type="pagebreak" id="pdf-page-1" title="1"></span>\n'
            "<!-- PDF_PAGE: 1 -->\n\n"
            "1\n",
            encoding="utf-8",
        )
        with self.assertRaisesRegex(
            ValueError,
            "body is empty after publication metadata removal",
        ):
            compile_chapters(
                self.pdf_path,
                output,
                [PageRecord(1, "OCR 正文")],
                payload,
                granularity="chapter",
            )
        self.assertEqual(previous.read_text(encoding="utf-8"), "# 上次成功产物\n")

    def test_compile_preflight_preserves_previous_chapters_on_failure(self) -> None:
        output = self.root / "preflight-preserves-output"
        chapter_dir = output / "chapters"
        chapter_dir.mkdir(parents=True)
        old = chapter_dir / "001_previous.md"
        old.write_text("# 上次成功产物\n", encoding="utf-8")
        payload = {
            "page_offset": 0,
            "entries": [
                TocEntry(
                    "chapter",
                    "第一章",
                    "正文",
                    1,
                    "chapter",
                    1,
                    pdf_page=1,
                ).__dict__
            ],
        }
        with self.assertRaisesRegex(ValueError, "Missing or stale translation"):
            compile_chapters(
                self.pdf_path,
                output,
                [PageRecord(1, "日本語", language="ja")],
                payload,
                granularity="chapter",
                require_translation=True,
            )
        self.assertEqual(old.read_text(encoding="utf-8"), "# 上次成功产物\n")

    def test_epub_and_bookmarked_pdf(self) -> None:
        output = self.root / "book"
        chapter_dir = output / "chapters"
        chapter_dir.mkdir(parents=True)
        (chapter_dir / "001_第一章.md").write_text(
            """# 第一章

<!-- source-pdf: sample.pdf -->
<!-- pdf-pages: 4-5 -->

<span epub:type="pagebreak" id="pdf-page-4" title="4"></span>
<!-- PDF_PAGE: 4 -->

正文第一页。

1

<span epub:type="pagebreak" id="pdf-page-5" title="5"></span>
<!-- PDF_PAGE: 5 -->

正文第二页。

2
""",
            encoding="utf-8",
        )
        manifest = [
            {
                "sequence": 1,
                "display_title": "第一章",
                "filename": "001_第一章.md",
            }
        ]
        epub_path = output / "sample.epub"
        build_epub(epub_path, chapter_dir, manifest, book_title="测试书", language="zh-CN")
        with zipfile.ZipFile(epub_path) as archive:
            self.assertEqual(archive.namelist()[0], "mimetype")
            self.assertEqual(archive.read("mimetype"), b"application/epub+zip")
            self.assertIn("OEBPS/nav.xhtml", archive.namelist())
            ET.fromstring(archive.read("OEBPS/nav.xhtml"))
            chapter_xml = archive.read("OEBPS/001_第一章.xhtml")
            ET.fromstring(chapter_xml)
            self.assertNotIn(b"source-pdf", chapter_xml)
            self.assertNotIn(b"PDF_PAGE", chapter_xml)
            self.assertNotIn(b"pdf-page-", chapter_xml)
            self.assertNotIn(b">1<", chapter_xml)

        docx_path = output / "sample.docx"
        build_docx(docx_path, chapter_dir, manifest, book_title="测试书")
        with zipfile.ZipFile(docx_path) as archive:
            document_xml = archive.read("word/document.xml")
            ET.fromstring(document_xml)
            self.assertNotIn(b"source-pdf", document_xml)
            self.assertNotIn(b"PDF_PAGE", document_xml)
            self.assertNotIn(b"pdf-page-", document_xml)
            self.assertEqual(document_xml.count(b'<w:br w:type="page"'), 1)

        toc_payload = {
            "entries": [
                {
                    "id": "toc-1",
                    "index": "第一章",
                    "title": "起点",
                    "level": 1,
                    "kind": "chapter",
                    "printed_page": 1,
                    "pdf_page": 4,
                    "end_pdf_page": 8,
                },
                {
                    "id": "toc-end",
                    "index": "",
                    "title": "封底",
                    "level": 1,
                    "kind": "part",
                    "printed_page": None,
                    "pdf_page": 8,
                    "end_pdf_page": 8,
                }
            ]
        }
        bookmarked = output / "bookmarked.pdf"
        build_bookmarked_pdf(self.pdf_path, bookmarked, toc_payload)
        with fitz.open(bookmarked) as document:
            self.assertEqual(document.get_toc()[0][1:], ["第一章 起点", 4])
            self.assertEqual(len(document.get_toc()), 1)

    def test_cli_manual_toc_to_all_outputs(self) -> None:
        output = self.root / "cli-output"
        for record in self.sample_records():
            record.ocr_model = "fixture-ocr"
            record.language = "zh"
            save_page_record(output, record)
        manual_toc = self.root / "manual-toc.json"
        manual_toc.write_text(
            json.dumps(
                {
                    "toc_pdf_pages": [2],
                    "entries": [entry.__dict__ for entry in self.sample_entries()],
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        self.assertEqual(
            main(
                [
                    str(self.pdf_path),
                    "-o",
                    str(output),
                    "--phase",
                    "toc",
                    "--toc-json",
                    str(manual_toc),
                    "--page-offset",
                    "3",
                ]
            ),
            0,
        )
        self.assertEqual(
            main(
                [
                    str(self.pdf_path),
                    "-o",
                    str(output),
                    "--phase",
                    "compile",
                    "--granularity",
                    "chapter",
                    "--no-verify",
                ]
            ),
            0,
        )
        self.assertTrue((output / "knowledge_base.jsonl").exists())
        self.assertTrue((output / "sample.epub").exists())
        self.assertTrue((output / "sample.docx").exists())
        self.assertTrue((output / "sample_带目录.pdf").exists())


if __name__ == "__main__":
    unittest.main()
