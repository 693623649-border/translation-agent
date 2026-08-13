import hashlib
import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock

import fitz

from book_pipeline import (
    PageRecord,
    TocEntry,
    _main_unlocked,
    build_bookmarked_pdf,
    build_docx,
    build_epub,
    compile_chapters,
    save_page_record,
    write_knowledge_base,
)
from publication_verifier import (
    _citation_inventory,
    _docx_markdown_body,
    _docx_positive_footnote_texts,
    _markdown_visible_text,
    verify_publication,
)


class PublicationVerifierTests(unittest.TestCase):
    book_title = "测试书"

    def test_markdown_visible_text_mirrors_trailing_page_discard(self) -> None:
        source = "# 章节\n\n正文。\n\n4\n0\n"
        self.assertEqual(
            _markdown_visible_text(source),
            "章节\n正文。",
        )

    def test_docx_body_strips_footnote_after_ascii_exclamation(self) -> None:
        source = "# 章节\n\n人心![^note]\n\n[^note]: 注释。\n"
        self.assertEqual(_docx_markdown_body(source), "# 章节\n\n人心!\n")

    def test_standard_footnote_inventory_keeps_legacy_markers_out_of_contract(self) -> None:
        inventory = _citation_inventory(
            "# 章节\n\n标准注释[^note]，遗留角标〔7〕与方括号 [8]。\n\n"
            "[^note]: 标准定义。\n"
        )

        self.assertEqual(inventory["references"], ["note"])
        self.assertEqual(inventory["definitions"], ["note"])
        self.assertEqual(inventory["missing_definitions"], [])
        self.assertEqual(inventory["unused_definitions"], [])
        self.assertEqual(
            inventory["legacy_reference_markers"],
            ["〔7〕", "[8]"],
        )

    def setUp(self) -> None:
        self.render_gate = mock.patch(
            "publication_verifier._check_docx_render",
            return_value={
                "summary": "fixture render passed",
                "metrics": {"page_count": 1},
                "issues": [],
                "warnings": [],
            },
        )
        self.render_gate_mock = self.render_gate.start()
        self.addCleanup(self.render_gate.stop)
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.source_pdf = self.root / "source.pdf"
        self.output = self.root / "output"
        self._build_source_pdf()
        self._build_complete_bundle()

    def tearDown(self) -> None:
        self.temp.cleanup()

    def _build_source_pdf(self) -> None:
        document = fitz.open()
        for page_number in range(1, 7):
            page = document.new_page()
            page.insert_text((72, 72), f"ORIGINAL PAGE {page_number}")
        document.save(self.source_pdf)
        document.close()

    @staticmethod
    def _toc_entries() -> list[TocEntry]:
        return [
            TocEntry("cover", "", "封面", 1, "part", None, pdf_page=1),
            TocEntry("ch-1", "第一章", "审定稿", 1, "chapter", 1, pdf_page=2),
            TocEntry("sec-1", "第一节", "结构验证", 2, "section", 2, pdf_page=3),
            TocEntry("sub-1", "一、", "细节", 3, "subsection", 3, pdf_page=4),
            TocEntry("ch-2", "第二章", "续论", 1, "chapter", 4, pdf_page=5),
            TocEntry("back-cover", "", "封底", 1, "part", None, pdf_page=6),
        ]

    @staticmethod
    def _chapter_one() -> str:
        return (
            "# 第一章 审定稿\n\n"
            "## 第一节 结构验证\n\n"
            "普通文字与**粗体文字**、*斜体文字*、<u>下划线文字</u>，年份 [2022] 必须保留。"
            "正文含引注[^101]和[^102]。\n\n"
            "> 引文第一段。\n>\n> 引文第二段。\n\n"
            "| 名称 | 值 |\n"
            "| --- | --- |\n"
            "| 罗亚 | 保留 |\n\n"
            "[^101]: 与正文一一对应的完整注释。\n\n"
            "[^102]: 第二条完整注释。\n"
        )

    @staticmethod
    def _chapter_two() -> str:
        # Exceed the 4000-character knowledge-base chunk size so coverage also
        # proves that a chapter split into multiple rows is reconstructed.
        return (
            "# 第二章 续论\n\n"
            "## 第二节 长文\n\n"
            "第二章开篇。\n\n"
            + "乙" * 4100
            + "\n"
        )

    def _build_complete_bundle(self) -> None:
        reviewed_dir = self.output / "reviewed_chapters"
        reviewed_dir.mkdir(parents=True)
        (reviewed_dir / "ch-1.md").write_text(
            self._chapter_one(), encoding="utf-8"
        )
        (reviewed_dir / "ch-2.md").write_text(
            self._chapter_two(), encoding="utf-8"
        )

        entries = self._toc_entries()
        toc_payload = {
            "page_offset": 0,
            "printed_pages_per_pdf_page": 1,
            "entries": [entry.__dict__ for entry in entries],
        }
        (self.output / "toc.json").write_text(
            json.dumps(toc_payload, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        records = [
            PageRecord(page, f"不会进入审定成品的第 {page} 页 OCR 文字。")
            for page in range(1, 7)
        ]
        for record in records:
            record.language = "zh"
            record.ocr_model = "fixture-ocr"
            save_page_record(self.output, record)
        manifest, knowledge_rows = compile_chapters(
            self.source_pdf,
            self.output,
            records,
            toc_payload,
            granularity="chapter",
        )
        write_knowledge_base(
            self.output / "knowledge_base.jsonl", knowledge_rows
        )
        build_epub(
            self.output / f"{self.book_title}.epub",
            self.output / "chapters",
            manifest,
            book_title=self.book_title,
            language="zh-CN",
        )
        build_docx(
            self.output / f"{self.book_title}.docx",
            self.output / "chapters",
            manifest,
            book_title=self.book_title,
        )
        build_bookmarked_pdf(
            self.source_pdf,
            self.output / f"{self.book_title}_带目录.pdf",
            toc_payload,
        )

    @staticmethod
    def _checks(report: dict[str, object]) -> dict[str, dict[str, object]]:
        return {
            str(item["id"]): item
            for item in report["checks"]  # type: ignore[index, union-attr]
        }

    def _verify(
        self,
        *,
        chapter_ids: list[str] | None = None,
        report_name: str = "release-report.json",
    ) -> dict[str, object]:
        return verify_publication(
            self.output,
            source_pdf=self.source_pdf,
            book_title=self.book_title,
            expected_language="zh-CN",
            require_all_reviewed=True,
            chapter_ids=chapter_ids,
            report_path=self.output / "audit" / report_name,
        )

    def _refresh_semantic_markdown_digest(self, chapter_id: str) -> None:
        manifest = json.loads(
            (self.output / "chapters.json").read_text(encoding="utf-8")
        )
        manifest_item = next(
            item for item in manifest if str(item.get("id") or "") == chapter_id
        )
        chapter_path = self.output / "chapters" / manifest_item["filename"]
        audit_path = self.output / "audit" / "semantic-reconstruction.json"
        audit = json.loads(audit_path.read_text(encoding="utf-8"))
        audit_item = next(
            item
            for item in audit["chapters"]
            if str(item.get("chapter_id") or "") == chapter_id
        )
        audit_item["markdown_sha256"] = hashlib.sha256(
            chapter_path.read_bytes()
        ).hexdigest()
        audit_path.write_text(
            json.dumps(audit, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )

    @staticmethod
    def _rewrite_zip_member(path: Path, member: str, transform) -> None:
        temporary = path.with_suffix(path.suffix + ".rewrite")
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
            temporary, "w"
        ) as destination:
            for info in source.infolist():
                payload = source.read(info.filename)
                if info.filename == member:
                    payload = transform(payload)
                destination.writestr(info, payload)
        temporary.replace(path)

    def _install_true_footnotes(
        self,
        *,
        reference_ids: tuple[int, ...] = (1,),
        definition_ids: tuple[int, ...] = (1,),
        separator_type: str | None = "separator",
        continuation_type: str | None = "continuationSeparator",
    ) -> None:
        path = self.output / f"{self.book_title}.docx"
        references = "".join(
            f'<w:footnoteReference w:id="{note_id}"/>'
            for note_id in reference_ids
        ).encode()
        self._rewrite_zip_member(
            path,
            "word/document.xml",
            lambda payload: payload.replace(
                b"</w:t>", b"</w:t>" + references, 1
            ),
        )
        self._rewrite_zip_member(
            path,
            "word/_rels/document.xml.rels",
            lambda payload: payload.replace(
                b"</Relationships>",
                (
                    b'<Relationship Id="rIdFootnotes" '
                    b'Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                    b'relationships/footnotes" Target="footnotes.xml"/>'
                    b"</Relationships>"
                ),
                1,
            ),
        )
        self._rewrite_zip_member(
            path,
            "[Content_Types].xml",
            lambda payload: payload.replace(
                b"</Types>",
                (
                    b'<Override PartName="/word/footnotes.xml" '
                    b'ContentType="application/vnd.openxmlformats-officedocument.'
                    b'wordprocessingml.footnotes+xml"/></Types>'
                ),
                1,
            ),
        )
        separator_attribute = (
            f' w:type="{separator_type}"' if separator_type is not None else ""
        )
        continuation_attribute = (
            f' w:type="{continuation_type}"'
            if continuation_type is not None
            else ""
        )
        definitions = "".join(
            (
                f'<w:footnote w:id="{note_id}"><w:p><w:r>'
                f'<w:footnoteRef/><w:t>脚注 {note_id}</w:t>'
                f"</w:r></w:p></w:footnote>"
            )
            for note_id in definition_ids
        )
        footnotes_xml = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/'
            'wordprocessingml/2006/main">'
            f'<w:footnote w:id="-1"{separator_attribute}><w:p><w:r>'
            '<w:separator/></w:r></w:p></w:footnote>'
            f'<w:footnote w:id="0"{continuation_attribute}><w:p><w:r>'
            '<w:continuationSeparator/></w:r></w:p></w:footnote>'
            f"{definitions}</w:footnotes>"
        ).encode()
        with zipfile.ZipFile(path, "a") as archive:
            archive.writestr("word/footnotes.xml", footnotes_xml)

    def test_complete_bundle_passes_all_publication_checks_and_writes_report(self) -> None:
        report = self._verify()

        self.assertTrue(report["ok"], report.get("errors"))
        self.assertEqual(report["status"], "passed")
        self.assertEqual(report["mode"], "full")
        checks = self._checks(report)
        for check_id in (
            "checkpoints.complete",
            "manifest.valid",
            "chapters.files",
            "reviewed.exact",
            "semantics.integrity",
            "citations.integrity",
            "content.hygiene",
            "epub.structure",
            "docx.structure",
            "docx.render",
            "knowledge_base.structure",
            "pdf.bookmarks",
        ):
            with self.subTest(check_id=check_id):
                self.assertEqual(checks[check_id]["status"], "passed")

        report_path = self.output / "audit" / "release-report.json"
        self.assertTrue(report_path.is_file())
        saved = json.loads(report_path.read_text(encoding="utf-8"))
        self.assertEqual(saved["status"], "passed")
        self.assertEqual(saved["report_path"], str(report_path.resolve()))
        self.assertEqual(saved["checks"], report["checks"])

    def test_docx_render_failure_blocks_release_report(self) -> None:
        self.render_gate_mock.return_value = {
            "summary": "rendered Word output lost visible CJK glyphs",
            "metrics": {"page_count": 1},
            "issues": [
                {
                    "code": "docx_render_cjk_glyphs_missing",
                    "message": "CJK text exists in the PDF layer but is not visibly rendered.",
                }
            ],
            "warnings": [],
        }

        report = self._verify(report_name="render-failed.json")
        render_check = self._checks(report)["docx.render"]

        self.assertFalse(report["ok"])
        self.assertEqual(report["status"], "failed")
        self.assertEqual(render_check["status"], "failed")
        self.assertIn(
            "docx_render_cjk_glyphs_missing",
            {issue["code"] for issue in render_check["issues"]},
        )

    def test_reviewed_byte_mismatch_is_a_targeted_failure(self) -> None:
        manifest = json.loads(
            (self.output / "chapters.json").read_text(encoding="utf-8")
        )
        chapter_path = self.output / "chapters" / manifest[0]["filename"]
        chapter_path.write_text(
            chapter_path.read_text(encoding="utf-8").replace(
                "普通文字", "被篡改的文字", 1
            ),
            encoding="utf-8",
        )

        report = self._verify(chapter_ids=["ch-1"], report_name="mismatch.json")
        checks = self._checks(report)
        self.assertFalse(report["ok"])
        self.assertEqual(report["mode"], "chapters")
        self.assertEqual(checks["reviewed.exact"]["status"], "failed")
        self.assertTrue(checks["reviewed.exact"].get("issues"))
        for check_id in (
            "epub.structure",
            "docx.structure",
            "docx.render",
            "knowledge_base.structure",
            "pdf.bookmarks",
        ):
            self.assertEqual(checks[check_id]["status"], "skipped")

    def test_semantic_markdown_digest_missing_blocks_release(self) -> None:
        audit_path = self.output / "audit" / "semantic-reconstruction.json"
        audit = json.loads(audit_path.read_text(encoding="utf-8"))
        audit["chapters"][0].pop("markdown_sha256")
        audit_path.write_text(
            json.dumps(audit, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )

        report = self._verify(
            chapter_ids=["ch-1"],
            report_name="semantic-digest-missing.json",
        )
        check = self._checks(report)["semantics.integrity"]

        self.assertFalse(report["ok"])
        self.assertIn(
            "semantic_markdown_digest_stale",
            {issue["code"] for issue in check["issues"]},
        )

    def test_semantic_markdown_digest_mismatch_blocks_release(self) -> None:
        audit_path = self.output / "audit" / "semantic-reconstruction.json"
        audit = json.loads(audit_path.read_text(encoding="utf-8"))
        audit["chapters"][0]["markdown_sha256"] = "0" * 64
        audit_path.write_text(
            json.dumps(audit, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )

        report = self._verify(
            chapter_ids=["ch-1"],
            report_name="semantic-digest-mismatch.json",
        )
        check = self._checks(report)["semantics.integrity"]

        self.assertFalse(report["ok"])
        self.assertIn(
            "semantic_markdown_digest_stale",
            {issue["code"] for issue in check["issues"]},
        )

    def test_semantic_audit_top_level_release_flags_block_release(self) -> None:
        audit_path = self.output / "audit" / "semantic-reconstruction.json"
        original = json.loads(audit_path.read_text(encoding="utf-8"))
        variants = {
            "status": lambda audit: audit.update(status="blocked"),
            "release_blocked": lambda audit: audit.update(release_blocked=True),
            "summary.release_blocked": lambda audit: audit["summary"].update(
                release_blocked=True
            ),
        }

        for name, mutate in variants.items():
            with self.subTest(signal=name):
                audit = json.loads(json.dumps(original))
                mutate(audit)
                audit_path.write_text(
                    json.dumps(audit, ensure_ascii=False, indent=2) + "\n",
                    encoding="utf-8",
                )

                report = self._verify(
                    chapter_ids=["ch-1"],
                    report_name=f"semantic-top-level-{name}.json",
                )
                check = self._checks(report)["semantics.integrity"]

                self.assertFalse(report["ok"])
                matching = [
                    issue
                    for issue in check["issues"]
                    if issue["code"] == "semantic_audit_release_blocked"
                ]
                self.assertEqual(len(matching), 1)
                self.assertIn(name, matching[0]["evidence"]["signals"])

    def test_reviewed_source_bom_is_normalized_but_generated_bytes_are_strict(self) -> None:
        manifest = json.loads(
            (self.output / "chapters.json").read_text(encoding="utf-8")
        )
        reviewed_path = self.output / "reviewed_chapters" / "ch-1.md"
        reviewed_path.write_text("\ufeff  \n" + self._chapter_one(), encoding="utf-8")
        report = self._verify(chapter_ids=["ch-1"], report_name="bom.json")
        self.assertTrue(report["ok"], report.get("errors"))

        chapter_path = self.output / "chapters" / manifest[0]["filename"]
        chapter_path.write_bytes(
            chapter_path.read_text(encoding="utf-8").replace("\n", "\r\n").encode(
                "utf-8"
            )
        )
        report = self._verify(chapter_ids=["ch-1"], report_name="strict-bytes.json")
        check = self._checks(report)["reviewed.exact"]
        self.assertFalse(report["ok"])
        self.assertIn(
            "reviewed_content_mismatch",
            {issue["code"] for issue in check["issues"]},
        )

    def test_orphan_body_citation_fails_integrity_check(self) -> None:
        manifest = json.loads(
            (self.output / "chapters.json").read_text(encoding="utf-8")
        )
        orphaned = self._chapter_one().replace(
            "\n\n[^101]: 与正文一一对应的完整注释。",
            "",
        )
        reviewed_path = self.output / "reviewed_chapters" / "ch-1.md"
        chapter_path = self.output / "chapters" / manifest[0]["filename"]
        reviewed_path.write_text(orphaned, encoding="utf-8")
        chapter_path.write_text(orphaned, encoding="utf-8")

        report = self._verify(chapter_ids=["ch-1"], report_name="orphan.json")
        checks = self._checks(report)
        self.assertFalse(report["ok"])
        self.assertEqual(checks["reviewed.exact"]["status"], "passed")
        self.assertEqual(checks["citations.integrity"]["status"], "failed")
        self.assertTrue(checks["citations.integrity"].get("issues"))

    def test_reviewed_chapter_orphan_definition_fails_bidirectional_integrity(self) -> None:
        manifest = json.loads(
            (self.output / "chapters.json").read_text(encoding="utf-8")
        )
        orphaned = self._chapter_one().replace("和[^102]", "", 1)
        (self.output / "reviewed_chapters" / "ch-1.md").write_text(
            orphaned, encoding="utf-8"
        )
        (self.output / "chapters" / manifest[0]["filename"]).write_text(
            orphaned, encoding="utf-8"
        )

        report = self._verify(chapter_ids=["ch-1"], report_name="unused.json")
        check = self._checks(report)["citations.integrity"]
        self.assertFalse(report["ok"])
        self.assertIn(
            "unused_note_definitions",
            {issue["code"] for issue in check["issues"]},
        )

    def test_standard_footnotes_with_legacy_markers_only_warn(self) -> None:
        manifest = json.loads(
            (self.output / "chapters.json").read_text(encoding="utf-8")
        )
        mixed = self._chapter_one().replace(
            "正文含引注[^101]和[^102]。",
            "正文含引注[^101]和[^102]，另有待迁移旧标记〔7〕和 [8]。",
            1,
        )
        (self.output / "reviewed_chapters" / "ch-1.md").write_text(
            mixed, encoding="utf-8"
        )
        (self.output / "chapters" / manifest[0]["filename"]).write_text(
            mixed, encoding="utf-8"
        )
        self._refresh_semantic_markdown_digest("ch-1")

        report = self._verify(
            chapter_ids=["ch-1"],
            report_name="standard-with-legacy-markers.json",
        )
        check = self._checks(report)["citations.integrity"]
        chapter = check["metrics"]["chapters"]["ch-1"]
        warning_codes = {warning["code"] for warning in check["warnings"]}

        self.assertTrue(report["ok"], report.get("errors"))
        self.assertEqual(check["status"], "passed")
        self.assertEqual(chapter["reference_count"], 2)
        self.assertEqual(chapter["definition_count"], 2)
        self.assertEqual(chapter["missing_definitions"], [])
        self.assertEqual(chapter["unused_definitions"], [])
        self.assertEqual(chapter["legacy_reference_markers"], ["〔7〕", "[8]"])
        self.assertIn("legacy_citation_markers_ignored", warning_codes)
        self.assertNotIn("ambiguous_plain_numeric_markers", warning_codes)

    def test_legacy_interleaved_square_notes_are_blocked_by_semantic_gate(self) -> None:
        manifest = json.loads(
            (self.output / "chapters.json").read_text(encoding="utf-8")
        )
        legacy = (
            "# 第一章 审定稿\n\n"
            "正文引用真实注释 [167]，但年份 [2022] 只是年份。\n\n"
            "167 旧式行首注释定义。\n\n"
            "正文在逐页注释之后继续。\n"
        )
        (self.output / "reviewed_chapters" / "ch-1.md").write_text(
            legacy, encoding="utf-8"
        )
        (self.output / "chapters" / manifest[0]["filename"]).write_text(
            legacy, encoding="utf-8"
        )

        report = self._verify(chapter_ids=["ch-1"], report_name="legacy-notes.json")
        checks = self._checks(report)
        check = checks["citations.integrity"]
        chapter = check["metrics"]["chapters"]["ch-1"]
        self.assertFalse(report["ok"])
        self.assertIn(
            "semantic_legacy_notes_unresolved",
            {issue["code"] for issue in checks["semantics.integrity"]["issues"]},
        )
        self.assertEqual(chapter["reference_count"], 1)
        self.assertEqual(chapter["definition_count"], 1)
        self.assertEqual(chapter["missing_definitions"], [])
        self.assertEqual(chapter["unused_definitions"], [])
        self.assertEqual(chapter["unmatched_plain_numeric_markers"], [])

    def test_manifest_sequence_can_select_incremental_chapter(self) -> None:
        report = self._verify(chapter_ids=["1"], report_name="sequence.json")

        self.assertTrue(report["ok"], report.get("errors"))
        self.assertEqual(report["summary"]["selected_chapter_count"], 1)

    def test_manifest_path_escape_is_rejected_before_file_read(self) -> None:
        path = self.output / "chapters.json"
        manifest = json.loads(path.read_text(encoding="utf-8"))
        manifest[0]["filename"] = "../outside.md"
        path.write_text(
            json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        (self.output / "outside.md").write_text(
            "# 第一章 审定稿\n\n不得读取。\n", encoding="utf-8"
        )

        report = self._verify(chapter_ids=["ch-1"], report_name="unsafe-path.json")
        check = self._checks(report)["manifest.valid"]
        self.assertFalse(report["ok"])
        self.assertIn(
            "chapter_filename_invalid",
            {issue["code"] for issue in check["issues"]},
        )

    def test_manifest_cannot_drop_a_toc_selected_chapter(self) -> None:
        path = self.output / "chapters.json"
        manifest = json.loads(path.read_text(encoding="utf-8"))[:1]
        path.write_text(
            json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )

        report = self._verify(report_name="toc-coverage.json")
        check = self._checks(report)["manifest.valid"]
        self.assertFalse(report["ok"])
        self.assertIn(
            "manifest_toc_coverage_mismatch",
            {issue["code"] for issue in check["issues"]},
        )

    def test_manifest_duplicate_display_titles_are_rejected(self) -> None:
        path = self.output / "chapters.json"
        manifest = json.loads(path.read_text(encoding="utf-8"))
        manifest[1]["display_title"] = manifest[0]["display_title"]
        path.write_text(
            json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )

        report = self._verify(report_name="duplicate-title.json")
        check = self._checks(report)["manifest.valid"]
        self.assertFalse(report["ok"])
        self.assertIn(
            "duplicate_chapter_titles",
            {issue["code"] for issue in check["issues"]},
        )

    def test_epub_metadata_heading_and_body_tampering_are_blocked(self) -> None:
        epub = self.output / f"{self.book_title}.epub"
        self._rewrite_zip_member(
            epub,
            "OEBPS/package.opf",
            lambda payload: payload.replace(
                self.book_title.encode(), b"WRONG TITLE", 1
            ).replace(b"<dc:language>zh-CN</dc:language>", b"<dc:language>xx</dc:language>", 1),
        )
        manifest = json.loads(
            (self.output / "chapters.json").read_text(encoding="utf-8")
        )
        first_xhtml = Path(manifest[0]["filename"]).with_suffix(".xhtml").name
        self._rewrite_zip_member(
            epub,
            f"OEBPS/{first_xhtml}",
            lambda payload: payload.replace(b"<h2>", b"<p>", 1).replace(
                b"</h2>", b"</p>", 1
            ).replace("普通文字".encode(), "篡改文字".encode(), 1),
        )

        report = self._verify(report_name="epub-tamper.json")
        check = self._checks(report)["epub.structure"]
        codes = {issue["code"] for issue in check["issues"]}
        self.assertFalse(report["ok"])
        self.assertIn("epub_book_title_mismatch", codes)
        self.assertIn("epub_language_mismatch", codes)
        self.assertIn("epub_chapter_language_mismatch", codes)
        self.assertIn("epub_heading_structure_mismatch", codes)
        self.assertIn("epub_chapter_text_mismatch", codes)

    def test_docx_front_matter_heading_quote_and_inline_style_tampering_are_blocked(self) -> None:
        from docx import Document

        path = self.output / f"{self.book_title}.docx"
        document = Document(path)
        document.sections[0].header.paragraphs[0].text = "来源PDF页码: 123"
        first_heading = next(
            paragraph
            for paragraph in document.paragraphs
            if paragraph.style.name == "Heading 1"
        )
        first_heading.insert_paragraph_before("EXTRA FRONT MATTER")
        next(
            paragraph
            for paragraph in document.paragraphs
            if paragraph.style.name == "Heading 2"
        ).style = "Normal"
        quote = next(
            paragraph
            for paragraph in document.paragraphs
            if paragraph.style.name == "Quote"
        )
        normal = next(
            paragraph
            for paragraph in document.paragraphs
            if paragraph.style.name == "Normal" and paragraph.text.strip()
        )
        quote.style = "Normal"
        normal.style = "Quote"
        for paragraph in document.paragraphs:
            for run in paragraph.runs:
                if run.bold is True:
                    run.bold = False
                if run.italic is True:
                    run.italic = False
                if bool(run.underline):
                    run.underline = False
        document.save(path)

        report = self._verify(report_name="docx-tamper.json")
        check = self._checks(report)["docx.structure"]
        codes = {issue["code"] for issue in check["issues"]}
        self.assertFalse(report["ok"])
        self.assertIn("docx_front_matter_mismatch", codes)
        self.assertIn("docx_unexpected_header_footer_or_notes", codes)
        self.assertIn("docx_heading_structure_mismatch", codes)
        self.assertIn("docx_quote_structure_mismatch", codes)
        self.assertIn("docx_inline_style_mismatch", codes)

    def test_docx_sanctioned_book_layout_and_page_footer_pass(self) -> None:
        from docx import Document
        from docx.enum.style import WD_STYLE_TYPE
        from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        path = self.output / f"{self.book_title}.docx"
        document = Document(path)
        title = next(
            paragraph
            for paragraph in document.paragraphs
            if paragraph.style.name in {"Title", "Codex Book Title"}
        )
        if "Codex Book Title" not in {
            item.name for item in document.styles
        }:
            style = document.styles.add_style(
                "Codex Book Title", WD_STYLE_TYPE.PARAGRAPH
            )
            style.base_style = document.styles["Title"]
            title.style = style
        had_author = bool(str(document.core_properties.author or "").strip())
        document.core_properties.author = "测试作者"
        existing_author = next(
            (
                paragraph
                for paragraph in document.paragraphs
                if paragraph.style.name == "Normal"
                and paragraph.text.strip() == "测试作者"
            ),
            None,
        )
        first_heading = next(
            paragraph
            for paragraph in document.paragraphs
            if paragraph.style.name == "Heading 1"
        )
        if existing_author is None and not had_author:
            author = first_heading.insert_paragraph_before("测试作者")
            author.alignment = WD_ALIGN_PARAGRAPH.CENTER
        first_heading_index = next(
            index
            for index, paragraph in enumerate(document.paragraphs)
            if paragraph._p is first_heading._p
        )
        existing_title_break = any(
            paragraph._p.xpath(".//w:br[@w:type='page']")
            for paragraph in document.paragraphs[:first_heading_index]
        )
        if not existing_title_break:
            title_break = first_heading.insert_paragraph_before()
            title_break.add_run().add_break(WD_BREAK.PAGE)
        for paragraph in document.paragraphs:
            if paragraph.style.name != "Heading 1":
                continue
            page_break_before = OxmlElement("w:pageBreakBefore")
            paragraph._p.get_or_add_pPr().append(page_break_before)

        footer = document.sections[0].footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if not footer._p.xpath(".//w:fldSimple[contains(@w:instr, 'PAGE')]"):
            field = OxmlElement("w:fldSimple")
            field.set(qn("w:instr"), "PAGE")
            run = OxmlElement("w:r")
            text = OxmlElement("w:t")
            text.text = "1"
            run.append(text)
            field.append(run)
            footer._p.append(field)
        document.save(path)

        report = self._verify(report_name="docx-book-layout.json")
        check = self._checks(report)["docx.structure"]
        self.assertTrue(report["ok"], check["issues"])
        self.assertEqual(check["metrics"]["chapter_page_break_before_count"], 2)
        self.assertEqual(check["metrics"]["title_page_break_count"], 1)
        self.assertEqual(check["metrics"]["unexpected_page_break_count"], 0)
        self.assertEqual(check["metrics"]["page_number_footer_count"], 1)

    def test_docx_page_footer_with_leaked_source_text_is_blocked(self) -> None:
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        path = self.output / f"{self.book_title}.docx"
        document = Document(path)
        footer = document.sections[0].footer.paragraphs[0]
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "1"
        run.append(text)
        field.append(run)
        footer._p.append(field)
        footer.add_run(" 来源PDF页码: 123")
        document.save(path)

        report = self._verify(report_name="docx-footer-source-leak.json")
        check = self._checks(report)["docx.structure"]
        codes = {issue["code"] for issue in check["issues"]}
        self.assertFalse(report["ok"])
        self.assertIn("docx_unexpected_header_footer_or_notes", codes)

    def test_docx_unregistered_body_page_break_is_blocked(self) -> None:
        from docx import Document
        from docx.enum.text import WD_BREAK

        path = self.output / f"{self.book_title}.docx"
        document = Document(path)
        normal = next(
            paragraph
            for paragraph in document.paragraphs
            if paragraph.style.name == "Normal" and paragraph.text.strip()
        )
        normal.add_run().add_break(WD_BREAK.PAGE)
        document.save(path)

        report = self._verify(report_name="docx-unregistered-page-break.json")
        check = self._checks(report)["docx.structure"]
        self.assertFalse(report["ok"])
        self.assertIn(
            "docx_page_break_present",
            {issue["code"] for issue in check["issues"]},
        )

    def test_docx_true_footnotes_pass_package_contract(self) -> None:
        report = self._verify(report_name="docx-footnotes.json")
        check = self._checks(report)["docx.structure"]
        self.assertTrue(report["ok"], check["issues"])
        self.assertEqual(check["status"], "passed")
        self.assertEqual(check["metrics"]["footnote_reference_count"], 2)
        self.assertEqual(check["metrics"]["footnote_definition_count"], 2)
        self.assertEqual(check["metrics"]["footnote_text_match_count"], 2)
        self.assertEqual(check["metrics"]["footnote_reserved_node_count"], 2)
        self.assertTrue(check["metrics"]["footnotes_relationship_valid"])
        self.assertTrue(check["metrics"]["footnotes_content_type_valid"])

    def test_docx_footnote_text_inventory_follows_reference_order(self) -> None:
        path = self.output / f"{self.book_title}.docx"
        with zipfile.ZipFile(path) as archive:
            original = _docx_positive_footnote_texts(archive)

        def swap_first_two_reference_ids(payload: bytes) -> bytes:
            first = b'<w:footnoteReference w:id="1"/>'
            second = b'<w:footnoteReference w:id="2"/>'
            sentinel = b'<w:footnoteReference w:id="__SWAP__"/>'
            return (
                payload.replace(first, sentinel, 1)
                .replace(second, first, 1)
                .replace(sentinel, second, 1)
            )

        self._rewrite_zip_member(
            path,
            "word/document.xml",
            swap_first_two_reference_ids,
        )

        with zipfile.ZipFile(path) as archive:
            self.assertEqual(
                _docx_positive_footnote_texts(archive),
                [original[1], original[0]],
            )

    def test_docx_missing_reserved_footnote_type_is_blocked(self) -> None:
        path = self.output / f"{self.book_title}.docx"
        self._rewrite_zip_member(
            path,
            "word/footnotes.xml",
            lambda payload: payload.replace(b' w:type="separator"', b"", 1),
        )

        report = self._verify(report_name="docx-footnote-reserved-type.json")
        check = self._checks(report)["docx.structure"]
        self.assertFalse(report["ok"])
        self.assertIn(
            "docx_footnote_reserved_type_invalid",
            {issue["code"] for issue in check["issues"]},
        )

    def test_docx_orphan_footnote_reference_is_blocked_and_counted(self) -> None:
        path = self.output / f"{self.book_title}.docx"
        self._rewrite_zip_member(
            path,
            "word/document.xml",
            lambda payload: payload.replace(
                b"</w:t>",
                b'</w:t><w:r><w:footnoteReference w:id="999"/></w:r>',
                1,
            ),
        )

        report = self._verify(report_name="docx-footnote-orphan-reference.json")
        check = self._checks(report)["docx.structure"]
        self.assertFalse(report["ok"])
        self.assertIn(
            "docx_footnote_orphan_reference",
            {issue["code"] for issue in check["issues"]},
        )
        self.assertEqual(check["metrics"]["footnote_orphan_reference_count"], 1)

    def test_docx_endnotes_remain_outside_publication_contract(self) -> None:
        path = self.output / f"{self.book_title}.docx"
        with zipfile.ZipFile(path, "a") as archive:
            archive.writestr(
                "word/endnotes.xml",
                (
                    '<w:endnotes xmlns:w="http://schemas.openxmlformats.org/'
                    'wordprocessingml/2006/main"/>'
                ),
            )

        report = self._verify(report_name="docx-endnotes.json")
        check = self._checks(report)["docx.structure"]
        self.assertFalse(report["ok"])
        self.assertIn(
            "docx_unexpected_endnotes",
            {issue["code"] for issue in check["issues"]},
        )
        self.assertTrue(check["metrics"]["endnotes_present"])

    def test_knowledge_base_content_tampering_is_blocked(self) -> None:
        path = self.output / "knowledge_base.jsonl"
        rows = [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines()]
        rows[0]["content"] = "被截断"
        rows[0]["id"] = "0" * 40
        path.write_text(
            "\n".join(json.dumps(row, ensure_ascii=False) for row in rows) + "\n",
            encoding="utf-8",
        )

        report = self._verify(report_name="kb-tamper.json")
        check = self._checks(report)["knowledge_base.structure"]
        self.assertFalse(report["ok"])
        self.assertEqual(check["status"], "failed")
        self.assertIn(
            "knowledge_base_stable_ids_mismatch",
            {issue["code"] for issue in check["issues"]},
        )

    def test_missing_page_checkpoint_blocks_full_release(self) -> None:
        (self.output / "pages" / "page_0006.json").unlink()

        report = self._verify(report_name="missing-checkpoint.json")
        check = self._checks(report)["checkpoints.complete"]
        self.assertFalse(report["ok"])
        self.assertIn(
            "checkpoint_page_coverage_mismatch",
            {issue["code"] for issue in check["issues"]},
        )

    def test_required_non_chinese_translation_must_be_source_fresh(self) -> None:
        manifest_path = self.output / "chapters.json"
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        manifest[0]["reviewed_override"] = False
        manifest_path.write_text(
            json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        page_path = self.output / "pages" / "page_0002.json"
        page = json.loads(page_path.read_text(encoding="utf-8"))
        page["language"] = "ja"
        page_path.write_text(
            json.dumps(page, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )

        report = verify_publication(
            self.output,
            source_pdf=self.source_pdf,
            book_title=self.book_title,
            expected_language="zh-CN",
            require_translation=True,
            report_path=self.output / "audit" / "translation-required.json",
        )
        check = self._checks(report)["checkpoints.complete"]
        self.assertFalse(report["ok"])
        self.assertIn(
            "checkpoint_translation_stale_or_missing",
            {issue["code"] for issue in check["issues"]},
        )

    def test_full_report_with_skipped_formats_is_partial_not_release_ready(self) -> None:
        report = verify_publication(
            self.output,
            source_pdf=self.source_pdf,
            book_title=self.book_title,
            expected_language="zh-CN",
            require_epub=False,
            require_docx=False,
            require_knowledge_base=False,
            require_bookmarked_pdf=False,
            require_all_reviewed=True,
            report_path=self.output / "audit" / "partial.json",
        )

        self.assertFalse(report["ok"])
        self.assertFalse(report["release_ready"])
        self.assertEqual(report["status"], "partial")
        self.assertEqual(report["summary"]["skipped"], 5)

    def test_word_profile_is_release_ready_without_unselected_containers(self) -> None:
        report = verify_publication(
            self.output,
            source_pdf=self.source_pdf,
            book_title=self.book_title,
            expected_language="zh-CN",
            require_epub=False,
            require_docx=True,
            require_docx_render=True,
            require_knowledge_base=False,
            require_bookmarked_pdf=False,
            require_all_reviewed=True,
            publication_profile="word",
            report_path=self.output / "audit" / "word-release-report.json",
        )

        self.assertTrue(report["ok"], report["errors"])
        self.assertTrue(report["release_ready"])
        self.assertEqual(report["status"], "passed")
        self.assertEqual(report["publication_profile"], "word")
        self.assertEqual(report["summary"]["skipped"], 3)
        checks = self._checks(report)
        self.assertEqual(checks["docx.structure"]["status"], "passed")
        self.assertEqual(checks["docx.render"]["status"], "passed")

    def test_word_profile_cli_writes_release_report_and_exits_zero(self) -> None:
        with mock.patch("builtins.print"):
            exit_code = _main_unlocked(
                [
                    str(self.source_pdf),
                    "--output-dir",
                    str(self.output),
                    "--phase",
                    "verify",
                    "--title",
                    self.book_title,
                    "--verification-profile",
                    "word",
                    "--no-epub",
                    "--no-kb",
                    "--no-bookmarked-pdf",
                    "--require-all-reviewed",
                ]
            )

        report_path = self.output / "audit" / "word-release-report.json"
        report = json.loads(report_path.read_text(encoding="utf-8"))
        self.assertEqual(exit_code, 0)
        self.assertEqual(report["publication_profile"], "word")
        self.assertTrue(report["release_ready"])
        self.assertEqual(report["status"], "passed")

    def test_report_write_failure_keeps_returned_counts_consistent(self) -> None:
        with mock.patch(
            "publication_verifier._write_report",
            side_effect=OSError("simulated write failure"),
        ):
            report = self._verify(report_name="unwritable.json")

        self.assertFalse(report["ok"])
        self.assertFalse(report["release_ready"])
        self.assertEqual(report["status"], "failed")
        self.assertEqual(report["summary"]["failed"], 1)
        self.assertEqual(report["metrics"]["failed_check_count"], 1)
        self.assertEqual(report["metrics"]["check_count"], len(report["checks"]))
        self.assertEqual(report["metrics"]["error_count"], len(report["errors"]))
        self.assertEqual(report["checks"][-1]["id"], "report.write")
        self.assertEqual(report["errors"][-1]["check_id"], "report.write")

    def test_unmanifested_epub_xhtml_is_rejected(self) -> None:
        epub = self.output / f"{self.book_title}.epub"
        with zipfile.ZipFile(epub, "a") as archive:
            archive.writestr(
                "OEBPS/stale.xhtml",
                "<html xmlns='http://www.w3.org/1999/xhtml'><body>stale</body></html>",
            )

        report = self._verify(report_name="epub-extra.json")
        check = self._checks(report)["epub.structure"]
        self.assertFalse(report["ok"])
        self.assertIn(
            "epub_archive_chapters_mismatch",
            {issue["code"] for issue in check["issues"]},
        )

    def test_same_page_count_modified_pdf_is_blocked(self) -> None:
        path = self.output / f"{self.book_title}_带目录.pdf"
        temporary = self.output / "modified.pdf"
        with fitz.open(path) as document:
            document[0].insert_text((72, 120), "TAMPERED")
            document.save(temporary)
        temporary.replace(path)

        report = self._verify(report_name="pdf-tamper.json")
        check = self._checks(report)["pdf.bookmarks"]
        codes = {issue["code"] for issue in check["issues"]}
        self.assertFalse(report["ok"])
        self.assertIn("pdf_page_render_mismatch", codes)
        self.assertIn("pdf_text_layer_mismatch", codes)

    def test_internal_page_token_is_blocked(self) -> None:
        manifest = json.loads(
            (self.output / "chapters.json").read_text(encoding="utf-8")
        )
        contaminated = self._chapter_one().replace("普通文字", "普通文字⟦P0001⟧")
        (self.output / "reviewed_chapters" / "ch-1.md").write_text(
            contaminated, encoding="utf-8"
        )
        (self.output / "chapters" / manifest[0]["filename"]).write_text(
            contaminated, encoding="utf-8"
        )

        report = self._verify(chapter_ids=["ch-1"], report_name="token.json")
        check = self._checks(report)["content.hygiene"]
        self.assertFalse(report["ok"])
        self.assertIn(
            "internal_placeholder",
            {issue["code"] for issue in check["issues"]},
        )

    def test_abandoned_page_image_directory_is_blocked(self) -> None:
        (self.output / "_page_images_1234_deadbeef").mkdir()

        report = self._verify(report_name="temp-dir.json")
        check = self._checks(report)["runtime.hygiene"]
        self.assertFalse(report["ok"])
        self.assertIn(
            "temporary_artifacts_present",
            {issue["code"] for issue in check["issues"]},
        )

    def test_forbidden_page_metadata_fails_content_hygiene(self) -> None:
        manifest = json.loads(
            (self.output / "chapters.json").read_text(encoding="utf-8")
        )
        contaminated = self._chapter_one().replace(
            "\n\n## 第一节", "\n\n<!-- PDF_PAGE: 2 -->\n\n## 第一节", 1
        )
        reviewed_path = self.output / "reviewed_chapters" / "ch-1.md"
        chapter_path = self.output / "chapters" / manifest[0]["filename"]
        reviewed_path.write_text(contaminated, encoding="utf-8")
        chapter_path.write_text(contaminated, encoding="utf-8")

        report = self._verify(chapter_ids=["ch-1"], report_name="metadata.json")
        checks = self._checks(report)
        self.assertFalse(report["ok"])
        self.assertEqual(checks["reviewed.exact"]["status"], "failed")
        self.assertEqual(checks["content.hygiene"]["status"], "failed")
        self.assertTrue(checks["content.hygiene"].get("issues"))

    def test_decorated_ocr_page_number_fails_content_hygiene(self) -> None:
        manifest = json.loads(
            (self.output / "chapters.json").read_text(encoding="utf-8")
        )
        contaminated = self._chapter_one().replace(
            "普通文字",
            "普通文字\n\n●30出血的代价。",
            1,
        )
        (self.output / "reviewed_chapters" / "ch-1.md").write_text(
            contaminated,
            encoding="utf-8",
        )
        (self.output / "chapters" / manifest[0]["filename"]).write_text(
            contaminated,
            encoding="utf-8",
        )

        report = self._verify(chapter_ids=["ch-1"], report_name="decorated-page.json")
        check = self._checks(report)["content.hygiene"]
        self.assertFalse(report["ok"])
        self.assertIn(
            "decorated_page_number",
            {issue["code"] for issue in check["issues"]},
        )

    def test_incremental_gate_uses_a_separate_default_report(self) -> None:
        report = verify_publication(
            self.output,
            source_pdf=self.source_pdf,
            book_title=self.book_title,
            require_all_reviewed=True,
            chapter_ids=["ch-1"],
        )

        self.assertTrue(report["ok"], report.get("errors"))
        self.assertEqual(report["mode"], "chapters")
        self.assertEqual(
            Path(str(report["report_path"])).name,
            "chapter-report.json",
        )

    def test_full_pdf_gate_requires_the_source_pdf(self) -> None:
        report = verify_publication(
            self.output,
            source_pdf=None,
            book_title=self.book_title,
            require_all_reviewed=True,
            report_path=self.output / "audit" / "missing-source.json",
        )

        checks = self._checks(report)
        self.assertFalse(report["ok"])
        self.assertEqual(checks["pdf.bookmarks"]["status"], "failed")
        self.assertIn(
            "source_pdf_required",
            {issue["code"] for issue in checks["pdf.bookmarks"]["issues"]},
        )

    @unittest.skipUnless(__import__("sys").platform != "win32", "POSIX flock test")
    def test_active_stage_lock_blocks_full_publication_gate(self) -> None:
        import fcntl

        lock_dir = self.output / ".stage_locks"
        lock_dir.mkdir(parents=True, exist_ok=True)
        lock_path = lock_dir / "ocr.lock"
        with lock_path.open("a+", encoding="utf-8") as handle:
            fcntl.flock(handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
            try:
                report = self._verify(report_name="active-lock.json")
            finally:
                fcntl.flock(handle.fileno(), fcntl.LOCK_UN)

        checks = self._checks(report)
        self.assertFalse(report["ok"])
        self.assertEqual(checks["runtime.hygiene"]["status"], "failed")
        self.assertIn(
            "active_stage_locks",
            {issue["code"] for issue in checks["runtime.hygiene"]["issues"]},
        )


if __name__ == "__main__":
    unittest.main()
