"""Tests for the source-to-knowledge-base ingest tool and its fidelity gate."""

from __future__ import annotations

import contextlib
import io
import json
import json as _json
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT / "tools" / "kb_ingest_plugin") not in sys.path:
    sys.path.insert(0, str(ROOT / "tools" / "kb_ingest_plugin"))

import kb_ingest  # noqa: E402


CHINESE_BODY = (
    "　　翻译在明治时期不只是一项学术工作，而是国家建设的一部分。"
    "福泽谕吉在《西洋事情》中所做的，正是把西方的制度与观念"
    "重新组装成日本人能够理解的叙述框架。"
)


def _write(path: Path, text: str) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")
    return path


def _docx_with_line_break(path: Path, text: str) -> Path:
    """Build a DOCX whose body paragraph carries a non-page ``w:br``."""

    from docx import Document
    from docx.enum.text import WD_BREAK

    document = Document()
    document.add_heading("第一章", level=1)
    paragraph = document.add_paragraph()
    paragraph.add_run("前半句被硬换行切断，")
    paragraph.add_run().add_break(WD_BREAK.LINE)
    paragraph.add_run("后半句继续。")
    document.save(path)
    return path


class SourceAdapterTests(unittest.TestCase):
    def test_markdown_splits_on_shallowest_heading(self) -> None:
        chapters = kb_ingest.parse_markdown_chapters(
            "# 书\n\n## 第一章\n\n甲\n\n## 第二章\n\n乙\n"
        )
        self.assertEqual([c["title"] for c in chapters], ["书"])
        # A single shallowest heading means the document is one chapter.
        self.assertIn("第一章", chapters[0]["body"])

    def test_markdown_uses_h2_when_it_is_shallowest(self) -> None:
        chapters = kb_ingest.parse_markdown_chapters(
            "## 第一章\n\n甲\n\n## 第二章\n\n乙\n"
        )
        self.assertEqual([c["title"] for c in chapters], ["第一章", "第二章"])

    def test_markdown_without_headings_is_one_chapter(self) -> None:
        chapters = kb_ingest.parse_markdown_chapters("只有正文。\n")
        self.assertEqual(len(chapters), 1)
        self.assertEqual(chapters[0]["title"], "全文")

    def test_markdown_ignores_headings_inside_code_fence(self) -> None:
        chapters = kb_ingest.parse_markdown_chapters(
            "# 真标题\n\n```\n# 不是标题\n```\n\n正文。\n"
        )
        self.assertEqual([c["title"] for c in chapters], ["真标题"])

    def test_plaintext_requires_short_standalone_heading(self) -> None:
        chapters = kb_ingest.parse_plaintext_chapters(
            "一、第一節\n\n甲文。\n\n二、第二節\n\n乙文。\n"
        )
        self.assertEqual([c["title"] for c in chapters], ["一、第一節", "二、第二節"])

    def test_plaintext_keeps_sentence_starting_with_number(self) -> None:
        # A long line merely beginning with a numeral is prose, not a heading.
        text = "一、这是正文而不是标题，因为它很长并且以句号结尾。\n"
        chapters = kb_ingest.parse_plaintext_chapters(text)
        self.assertEqual(chapters[0]["title"], "全文")

    def test_epub_without_spine_fails_closed(self) -> None:
        import zipfile

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "book.epub"
            with zipfile.ZipFile(path, "w") as archive:
                archive.writestr(
                    "META-INF/container.xml",
                    '<?xml version="1.0"?><container><rootfiles>'
                    '<rootfile full-path="OEBPS/content.opf"/></rootfiles></container>',
                )
                archive.writestr(
                    "OEBPS/content.opf",
                    '<?xml version="1.0"?><package><manifest>'
                    '<item id="c1" href="c1.xhtml"/></manifest></package>',
                )
                archive.writestr("OEBPS/c1.xhtml", "<html><body><p>正文</p></body></html>")
            with self.assertRaises(kb_ingest.IngestError) as caught:
                kb_ingest.parse_epub_chapters(path)
        self.assertIn("spine", str(caught.exception))

    def test_epub_reads_spine_order(self) -> None:
        import zipfile

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "book.epub"
            with zipfile.ZipFile(path, "w") as archive:
                archive.writestr(
                    "META-INF/container.xml",
                    '<?xml version="1.0"?><container><rootfiles>'
                    '<rootfile full-path="OEBPS/content.opf"/></rootfiles></container>',
                )
                archive.writestr(
                    "OEBPS/content.opf",
                    '<?xml version="1.0"?><package><manifest>'
                    '<item id="c1" href="c1.xhtml"/>'
                    '<item id="c2" href="c2.xhtml"/></manifest>'
                    '<spine><itemref idref="c2"/><itemref idref="c1"/></spine></package>',
                )
                archive.writestr(
                    "OEBPS/c1.xhtml",
                    "<html><body><h1>後</h1><p>第二篇正文。</p></body></html>",
                )
                archive.writestr(
                    "OEBPS/c2.xhtml",
                    "<html><body><h1>先</h1><p>第一篇正文。</p></body></html>",
                )
            chapters = kb_ingest.parse_epub_chapters(path)
        self.assertEqual([c["title"] for c in chapters], ["先", "後"])

    def test_epub_heading_is_not_duplicated_into_the_body(self) -> None:
        """A heading names the chapter and must not also open its body.

        Leaving it in the body duplicates the title and then trips the
        publisher's running-title stripper, which deletes a prose line that
        merely resembles the heading.
        """
        import zipfile

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "book.epub"
            with zipfile.ZipFile(path, "w") as archive:
                archive.writestr(
                    "META-INF/container.xml",
                    '<?xml version="1.0"?><container><rootfiles>'
                    '<rootfile full-path="OEBPS/content.opf"/></rootfiles></container>',
                )
                archive.writestr(
                    "OEBPS/content.opf",
                    '<?xml version="1.0"?><package><manifest>'
                    '<item id="c1" href="c1.xhtml"/></manifest>'
                    '<spine><itemref idref="c1"/></spine></package>',
                )
                archive.writestr(
                    "OEBPS/c1.xhtml",
                    "<html><body><h1>第一章 汉文直读与训读</h1>"
                    "<p>江户时代的读书人面对汉文，发展出两种处理方式。</p>"
                    "</body></html>",
                )
            chapters = kb_ingest.parse_epub_chapters(path)
        self.assertEqual(len(chapters), 1)
        self.assertEqual(chapters[0]["title"], "第一章 汉文直读与训读")
        self.assertNotIn("第一章", chapters[0]["body"])
        self.assertEqual(chapters[0]["body"], "江户时代的读书人面对汉文，发展出两种处理方式。")

    def test_epub_ingest_round_trips_without_losing_text(self) -> None:
        """The whole EPUB path must pass the fidelity gate, not just parse."""
        import zipfile

        with tempfile.TemporaryDirectory() as directory:
            workspace = Path(directory)
            path = workspace / "book.epub"
            with zipfile.ZipFile(path, "w") as archive:
                archive.writestr(
                    "META-INF/container.xml",
                    '<?xml version="1.0"?><container><rootfiles>'
                    '<rootfile full-path="OEBPS/content.opf"/></rootfiles></container>',
                )
                archive.writestr(
                    "OEBPS/content.opf",
                    '<?xml version="1.0"?><package><manifest>'
                    '<item id="c1" href="c1.xhtml"/></manifest>'
                    '<spine><itemref idref="c1"/></spine></package>',
                )
                archive.writestr(
                    "OEBPS/c1.xhtml",
                    "<html><body><h1>第一章 汉文直读与训读</h1>"
                    f"<p>{CHINESE_BODY}</p></body></html>",
                )
            report = kb_ingest.ingest_source(
                path, output_dir=workspace / "out", title="EPUB 测试书"
            )
        self.assertEqual(report["status"], "passed")
        self.assertEqual(report["adapter"], "epub-spine")
        fidelity = report["docx_fidelity"]
        self.assertEqual(fidelity["missing_characters"], 0)
        self.assertEqual(fidelity["extra_characters"], 0)
        self.assertEqual(fidelity["similarity"], 1.0)

    def test_heading_dropping_only_consumes_a_leading_run(self) -> None:
        # Each heading is consumed once and only from the front, so a heading
        # that legitimately recurs deeper in the prose survives.
        paragraphs = ["第一章", "正文甲。", "第一章", "正文乙。"]
        self.assertEqual(
            kb_ingest._drop_heading_paragraphs(paragraphs, ["第一章"]),
            ["正文甲。", "第一章", "正文乙。"],
        )
        # Nothing is dropped when the first paragraph is not a heading.
        self.assertEqual(
            kb_ingest._drop_heading_paragraphs(["正文。", "第一章"], ["第一章"]),
            ["正文。", "第一章"],
        )
        # Two stacked headings at the front are both consumed.
        self.assertEqual(
            kb_ingest._drop_heading_paragraphs(["上编", "第一章", "正文。"], ["上编", "第一章"]),
            ["正文。"],
        )

    def test_unsupported_suffix_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = _write(Path(directory) / "a.rtf", "x")
            with self.assertRaises(kb_ingest.IngestError):
                kb_ingest.read_source(path)

    def test_docx_without_headings_is_rejected(self) -> None:
        from docx import Document

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "flat.docx"
            document = Document()
            document.add_paragraph("没有标题结构。")
            document.save(path)
            with self.assertRaises(kb_ingest.IngestError) as caught:
                kb_ingest.read_source(path)
        self.assertIn("标题", str(caught.exception))

    def test_scanned_pdf_is_rejected_with_the_ocr_entrypoint(self) -> None:
        """An image-only PDF must fail closed and name the real command.

        Guessing here would silently put a blank or garbage corpus into the
        database, so the refusal is part of the contract rather than an
        incidental error.
        """
        try:
            import fitz
        except ImportError:  # pragma: no cover - dependency boundary.
            self.skipTest("PyMuPDF is not installed")

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "scan.pdf"
            document = fitz.open()
            document.new_page()
            document.save(str(path))
            document.close()
            with self.assertRaises(kb_ingest.IngestError) as caught:
                kb_ingest.read_source(path)
        message = str(caught.exception)
        self.assertIn("扫描件", message)
        self.assertIn("graph_pipeline.py", message)

    def test_unreadable_pdf_fails_with_a_named_cause(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "broken.pdf"
            path.write_bytes(b"%PDF-1.4\nnot actually a pdf\n")
            with self.assertRaises(kb_ingest.IngestError) as caught:
                kb_ingest.read_source(path)
        self.assertIn("PDF", str(caught.exception))

    def test_pdf_text_layer_is_read(self) -> None:
        try:
            import fitz
        except ImportError:  # pragma: no cover - dependency boundary.
            self.skipTest("PyMuPDF is not installed")

        font = next(
            (
                candidate
                for candidate in (
                    r"C:\Windows\Fonts\simsun.ttc",
                    r"C:\Windows\Fonts\msyh.ttc",
                    r"C:\Windows\Fonts\simhei.ttf",
                )
                if Path(candidate).is_file()
            ),
            None,
        )
        if font is None:
            self.skipTest("no CJK system font available for a text-layer fixture")

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "text.pdf"
            document = fitz.open()
            page = document.new_page()
            page.insert_textbox(
                fitz.Rect(50, 50, 550, 700),
                (CHINESE_BODY + "\n") * 6,
                fontsize=11,
                fontname="cjk",
                fontfile=font,
            )
            document.save(str(path))
            document.close()

            chapters, adapter = kb_ingest.read_source(path)
        self.assertEqual(adapter, "pdf-text-layer")
        self.assertTrue(chapters)
        self.assertTrue(any("翻译" in chapter["body"] for chapter in chapters))


class CorpusTests(unittest.TestCase):
    def _chapters(self) -> list[dict[str, object]]:
        return [
            {"title": "第一章 起点", "body": CHINESE_BODY},
            {"title": "第二章 展开", "body": CHINESE_BODY * 2},
        ]

    def test_rows_use_the_five_field_contract(self) -> None:
        source = Path("book.md")
        rows, sidecar = kb_ingest.build_corpus(
            source, self._chapters(), book_title="测试书", author="作者"
        )
        self.assertTrue(rows)
        for row in rows:
            self.assertEqual(
                sorted(row), ["chapter_id", "chapter_order", "content", "id", "title"]
            )
            self.assertEqual(len(row["id"]), 40)
            self.assertTrue(row["title"].startswith("[测试书] "))
        self.assertEqual(len(sidecar), len(rows))
        self.assertEqual(sidecar[0]["book_title"], "测试书")
        self.assertEqual(sidecar[0]["author"], "作者")

    def test_row_ids_are_stable_across_runs(self) -> None:
        source = Path("book.md")
        first, _ = kb_ingest.build_corpus(source, self._chapters(), book_title="测试书")
        second, _ = kb_ingest.build_corpus(source, self._chapters(), book_title="测试书")
        self.assertEqual([r["id"] for r in first], [r["id"] for r in second])

    def test_chunking_honours_the_paragraph_budget(self) -> None:
        long_body = "\n\n".join(CHINESE_BODY * 3 for _ in range(40))
        rows, _ = kb_ingest.build_corpus(
            Path("book.md"),
            [{"title": "长章", "body": long_body}],
            book_title="测试书",
            chunk_chars=4000,
        )
        self.assertGreater(len(rows), 1)
        for row in rows:
            self.assertLessEqual(len(row["content"]), 4000)

    def test_empty_source_is_rejected(self) -> None:
        with self.assertRaises(kb_ingest.IngestError):
            kb_ingest.build_corpus(
                Path("book.md"), [{"title": "空", "body": "   "}], book_title="测试书"
            )


class DiffTests(unittest.TestCase):
    def test_identical_texts_produce_no_difference(self) -> None:
        result = kb_ingest.diff_texts("甲乙丙", "甲乙丙")
        self.assertEqual(result["missing"], 0)
        self.assertEqual(result["extra"], 0)
        self.assertEqual(result["ratio"], 1.0)
        self.assertEqual(result["hunks"], [])

    def test_missing_characters_are_counted_and_located(self) -> None:
        result = kb_ingest.diff_texts("翻译与近代日本", "翻译近代日本")
        self.assertEqual(result["missing"], 1)
        self.assertEqual(result["extra"], 0)
        self.assertEqual(result["hunks"][0]["missing_text"], "与")
        self.assertIn("翻译", result["hunks"][0]["context"])

    def test_extra_characters_are_counted(self) -> None:
        result = kb_ingest.diff_texts("翻译与近代日本", "翻译与近代的日本")
        self.assertEqual(result["missing"], 0)
        self.assertEqual(result["extra"], 1)

    def test_reordering_is_not_hidden(self) -> None:
        # A character multiset would call this equal; an ordered diff must not.
        result = kb_ingest.diff_texts("甲乙", "乙甲")
        self.assertGreater(result["missing"], 0)
        self.assertGreater(result["extra"], 0)

    def test_hunk_list_is_bounded_but_counts_stay_exact(self) -> None:
        expected = "".join(f"{index}甲" for index in range(200))
        actual = "".join(f"{index}乙" for index in range(200))
        result = kb_ingest.diff_texts(expected, actual)
        self.assertEqual(result["missing"], 200)
        self.assertLessEqual(len(result["hunks"]), kb_ingest.MAX_HUNKS)
        self.assertTrue(result["truncated"])

    def test_whitespace_stripping_separates_layout_from_content(self) -> None:
        self.assertEqual(kb_ingest.strip_whitespace("甲 乙\n丙"), "甲乙丙")


class NormalizationTests(unittest.TestCase):
    def test_canonical_text_collapses_whitespace_and_folds_width(self) -> None:
        canonical = kb_ingest._canonical_from_markdown("甲\n\n\n乙　丙")
        self.assertEqual(canonical, "甲 乙 丙")

    def test_normalize_for_match_folds_punctuation_but_not_script(self) -> None:
        self.assertEqual(kb_ingest.normalize_for_match("，。！"), ",.!")
        # Traditional characters must survive: the gate checks the printed set.
        self.assertEqual(kb_ingest.normalize_for_match("翻譯"), "翻譯")
        self.assertEqual(kb_ingest.normalize_for_match("ABC"), "abc")

    def test_optionally_numbered_heading_is_not_a_heading(self) -> None:
        chapters = kb_ingest.parse_plaintext_chapters("一、短句。\n\n正文继续。\n")
        self.assertEqual(chapters[0]["title"], "全文")


class DocxInspectionTests(unittest.TestCase):
    def test_package_inspection_reports_styles_and_page_breaks(self) -> None:
        from docx import Document
        from docx.enum.text import WD_BREAK

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "doc.docx"
            document = Document()
            document.add_paragraph(style="Title").add_run("书名")
            document.add_paragraph("作者")
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            document.add_heading("第一章", level=1)
            document.add_paragraph("正文。")
            document.save(path)

            package = kb_ingest.inspect_docx_package(path)

        self.assertEqual(package["explicit_page_breaks"], 1)
        self.assertEqual(package["line_breaks"], [])
        self.assertEqual(kb_ingest.cover_prefix(package["paragraphs"]), "书名\n作者")
        body = kb_ingest._docx_body_only(package["paragraphs"])
        self.assertEqual([p["text"] for p in body], ["第一章", "正文。"])

    def test_soft_line_break_is_counted_as_abnormal(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = _docx_with_line_break(Path(directory) / "br.docx", "")
            package = kb_ingest.inspect_docx_package(path)
            report = kb_ingest.check_line_wrapping(package["paragraphs"])

        self.assertEqual(package["explicit_page_breaks"], 0)
        self.assertEqual(report["metrics"]["line_break_count"], 1)
        codes = {issue["code"] for issue in report["issues"]}
        self.assertIn("docx_line_break_present", codes)

    def test_cover_page_is_excluded_from_line_wrap_scope(self) -> None:
        from docx import Document
        from docx.enum.text import WD_BREAK

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "ok.docx"
            document = Document()
            document.add_paragraph(style="Title").add_run("书名")
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            document.add_heading("第一章", level=1)
            document.add_paragraph("没有任何硬换行的正文。")
            document.save(path)
            package = kb_ingest.inspect_docx_package(path)
            report = kb_ingest.check_line_wrapping(package["paragraphs"])

        self.assertEqual(report["metrics"]["line_break_count"], 0)
        self.assertEqual(report["issues"], [])

    def test_sentence_split_detection(self) -> None:
        # A paragraph that does not end a sentence, followed by one beginning
        # with a closing mark, is the signature of an unmerged hard wrap.
        self.assertTrue(kb_ingest._splits_sentence("前半句没有结束", "，后半句继续。"))
        # A legitimate new paragraph is never flagged.
        self.assertFalse(kb_ingest._splits_sentence("上一段结束了。", "下一段开始。"))
        self.assertFalse(kb_ingest._splits_sentence("这是一个标题", "正文另起一段。"))
        self.assertFalse(kb_ingest._splits_sentence("诗句一行", "诗句另一行"))


class FidelityGateTests(unittest.TestCase):
    def _publish(self, workspace: Path, body: str) -> dict[str, object]:
        source = _write(workspace / "src.md", f"# 第一章 测试\n\n{body}\n")
        return kb_ingest.ingest_source(
            source, output_dir=workspace / "out", title="测试书", author="作者"
        )

    def test_clean_source_passes_every_check(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            report = self._publish(Path(directory), CHINESE_BODY)

        self.assertEqual(report["status"], "passed")
        self.assertTrue(report["ok"])
        fidelity = report["docx_fidelity"]
        self.assertEqual(fidelity["missing_characters"], 0)
        self.assertEqual(fidelity["extra_characters"], 0)
        self.assertEqual(fidelity["similarity"], 1.0)
        self.assertEqual(fidelity["issues"], [])
        self.assertEqual(fidelity["metrics"]["paragraphs_with_line_breaks"], 0)
        self.assertEqual(report["language_gate"]["passed"], True)

    def test_corpus_and_sidecar_are_published(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            report = self._publish(Path(directory), CHINESE_BODY)
            out = Path(report["workspace"])
            corpus = out / "knowledge_base.jsonl"
            self.assertTrue(corpus.is_file())
            self.assertTrue((out / "knowledge_base.meta.jsonl").is_file())
            self.assertTrue((out / "knowledge_base.rag.json").is_file())
            rows = [
                json.loads(line)
                for line in corpus.read_text(encoding="utf-8").splitlines()
                if line.strip()
            ]
            self.assertEqual(len(rows), report["chunk_count"])
            self.assertEqual(
                sorted(rows[0]), ["chapter_id", "chapter_order", "content", "id", "title"]
            )

    def test_every_chapter_becomes_a_heading_one(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            workspace = Path(directory)
            source = _write(
                workspace / "src.md",
                "# 第一章 甲\n\n" + CHINESE_BODY + "\n\n# 第二章 乙\n\n" + CHINESE_BODY + "\n",
            )
            report = kb_ingest.ingest_source(
                source, output_dir=workspace / "out", title="双章书"
            )
        self.assertEqual(report["status"], "passed")
        self.assertEqual(report["chapter_titles"], ["第一章 甲", "第二章 乙"])

    def test_missing_characters_block_the_run(self) -> None:
        from docx import Document

        with tempfile.TemporaryDirectory() as directory:
            workspace = Path(directory)
            source = _write(workspace / "src.md", f"# 第一章\n\n{CHINESE_BODY}\n")
            report = kb_ingest.ingest_source(
                source, output_dir=workspace / "out", title="测试书"
            )
            docx_path = Path(report["artifacts"]["docx"])

            # Delete visible source characters from the Word body the way a
            # truncating converter would.
            document = Document(str(docx_path))
            target = next(
                p for p in document.paragraphs if "翻译在明治时期" in p.text
            )
            target.text = target.text[:10]
            document.save(str(docx_path))

            chapters = kb_ingest._read_chapters_from_output(workspace / "out")
            fidelity = kb_ingest.verify_word_document(
                docx_path, chapters, book_title="测试书"
            )

        self.assertEqual(fidelity["status"], "failed")
        self.assertGreater(fidelity["missing_characters"], 0)
        self.assertIn(
            "docx_text_mismatch", {issue["code"] for issue in fidelity["issues"]}
        )

    def test_injected_content_is_detected_as_extra(self) -> None:
        from docx import Document

        with tempfile.TemporaryDirectory() as directory:
            workspace = Path(directory)
            source = _write(workspace / "src.md", f"# 第一章\n\n{CHINESE_BODY}\n")
            report = kb_ingest.ingest_source(
                source, output_dir=workspace / "out", title="测试书"
            )
            docx_path = Path(report["artifacts"]["docx"])

            document = Document(str(docx_path))
            document.add_paragraph("这是模型凭空补写的一句。")
            document.save(str(docx_path))

            chapters = kb_ingest._read_chapters_from_output(workspace / "out")
            fidelity = kb_ingest.verify_word_document(
                docx_path, chapters, book_title="测试书"
            )

        self.assertEqual(fidelity["status"], "failed")
        self.assertGreater(fidelity["extra_characters"], 0)

    def test_placeholder_traces_are_reported(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            report = self._publish(
                Path(directory), CHINESE_BODY + "\n\nTRANSLATION_FAILED\n"
            )
        fidelity = report["docx_fidelity"]
        codes = {issue["code"] for issue in fidelity["issues"]}
        self.assertIn("docx_trace_internal_placeholder", codes)

    def test_heading_mismatch_is_reported(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            workspace = Path(directory)
            source = _write(workspace / "src.md", f"# 第一章 原名\n\n{CHINESE_BODY}\n")
            report = kb_ingest.ingest_source(
                source, output_dir=workspace / "out", title="测试书"
            )
            chapters = kb_ingest._read_chapters_from_output(workspace / "out")
            chapters[0]["title"] = "第一章 改过的名字"
            fidelity = kb_ingest.verify_word_document(
                Path(report["artifacts"]["docx"]), chapters, book_title="测试书"
            )

        self.assertIn(
            "docx_heading_mismatch", {issue["code"] for issue in fidelity["issues"]}
        )

    def test_soft_line_break_blocks_the_run_unless_allowed(self) -> None:
        from docx import Document
        from docx.enum.text import WD_BREAK

        with tempfile.TemporaryDirectory() as directory:
            workspace = Path(directory)
            source = _write(workspace / "src.md", f"# 第一章\n\n{CHINESE_BODY}\n")
            report = kb_ingest.ingest_source(
                source, output_dir=workspace / "out", title="测试书"
            )
            docx_path = Path(report["artifacts"]["docx"])

            document = Document(str(docx_path))
            paragraph = document.add_paragraph()
            paragraph.add_run("前半句，")
            paragraph.add_run().add_break(WD_BREAK.LINE)
            paragraph.add_run("后半句。")
            document.save(str(docx_path))

            chapters = kb_ingest._read_chapters_from_output(workspace / "out")
            strict = kb_ingest.verify_word_document(
                docx_path, chapters, book_title="测试书"
            )
            relaxed = kb_ingest.verify_word_document(
                docx_path, chapters, book_title="测试书", allow_line_breaks=True
            )

        self.assertEqual(strict["status"], "failed")
        self.assertIn(
            "docx_line_break_present", {issue["code"] for issue in strict["issues"]}
        )
        # Whitespace-only content changes: the break is a warning when allowed.
        self.assertIn(
            "docx_line_break_present",
            {warning["code"] for warning in relaxed["warnings"]},
        )

    def test_source_page_metadata_is_reported_not_silently_dropped(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            body = CHINESE_BODY + "\n\n<!-- PDF_PAGE: 42 -->\n\n" + CHINESE_BODY
            report = self._publish(Path(directory), body)
        fidelity = report["docx_fidelity"]
        # The comment cannot reach the page, so it must be surfaced explicitly.
        self.assertTrue(fidelity["ok"])
        self.assertEqual(fidelity["metrics"]["source_page_marker_count"], 1)
        warnings = {warning["code"]: warning for warning in fidelity["warnings"]}
        self.assertIn("docx_source_page_markers_present", warnings)
        self.assertEqual(
            warnings["docx_source_page_markers_present"]["evidence"]["page_comments"][0][
                "marker"
            ],
            "<!-- PDF_PAGE: 42 -->",
        )

    def test_standalone_source_page_numbers_are_reported(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            body = CHINESE_BODY + "\n\n42\n\n" + CHINESE_BODY
            report = self._publish(Path(directory), body)
        fidelity = report["docx_fidelity"]
        self.assertEqual(fidelity["metrics"]["source_page_marker_count"], 1)
        self.assertEqual(fidelity["missing_characters"], 0)
        self.assertEqual(fidelity["extra_characters"], 0)


class LanguageGateTests(unittest.TestCase):
    def test_foreign_prose_blocks_ingest(self) -> None:
        english = (
            "This chapter explains why translation became a state project "
            "during the Meiji period and how it reshaped modern Japanese prose."
        )
        with tempfile.TemporaryDirectory() as directory:
            workspace = Path(directory)
            source = _write(workspace / "en.md", f"# Chapter One\n\n{english}\n")
            report = kb_ingest.ingest_source(
                source, output_dir=workspace / "out", title="English Book"
            )
        self.assertEqual(report["status"], "blocked")
        self.assertFalse(report["ok"])
        self.assertEqual(report["language_gate"]["pending_count"], 1)
        self.assertIn("en", report["language_gate"]["languages"])
        self.assertFalse((Path(report["workspace"]) / "knowledge_base.jsonl").is_file())

    def test_allow_foreign_registers_the_corpus(self) -> None:
        english = (
            "This chapter explains why translation became a state project "
            "during the Meiji period and how it reshaped modern Japanese prose."
        )
        with tempfile.TemporaryDirectory() as directory:
            workspace = Path(directory)
            source = _write(workspace / "en.md", f"# Chapter One\n\n{english}\n")
            report = kb_ingest.ingest_source(
                source,
                output_dir=workspace / "out",
                title="English Book",
                allow_foreign=True,
            )
        self.assertTrue(report["ok"])
        self.assertTrue(report["language_gate"]["allowed_foreign"])

    def test_chinese_corpus_passes_the_gate(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            workspace = Path(directory)
            source = _write(workspace / "zh.md", f"# 第一章\n\n{CHINESE_BODY}\n")
            report = kb_ingest.ingest_source(
                source, output_dir=workspace / "out", title="中文书"
            )
        self.assertTrue(report["language_gate"]["passed"])
        self.assertEqual(report["language_gate"]["pending_count"], 0)


class CommandTests(unittest.TestCase):
    def _ingest(self, workspace: Path) -> dict:
        source = _write(workspace / "src.md", f"# 第一章\n\n{CHINESE_BODY}\n")
        return kb_ingest.ingest_source(
            source, output_dir=workspace / "out", title="测试书"
        )

    def test_main_returns_zero_on_success_and_writes_json(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            workspace = Path(directory)
            source = _write(workspace / "src.md", f"# 第一章\n\n{CHINESE_BODY}\n")
            buffer = io.StringIO()
            with contextlib.redirect_stdout(buffer):
                code = kb_ingest.main(
                    [
                        "ingest",
                        str(source),
                        "--output-dir",
                        str(workspace / "out"),
                        "--title",
                        "测试书",
                    ]
                )
        self.assertEqual(code, 0)
        payload = json.loads(buffer.getvalue().strip())
        self.assertTrue(payload["ok"])

    def test_main_reports_errors_as_json(self) -> None:
        buffer = io.StringIO()
        with contextlib.redirect_stdout(buffer):
            code = kb_ingest.main(["ingest", "missing.md"])
        self.assertEqual(code, 1)
        payload = json.loads(buffer.getvalue().strip())
        self.assertIn("error", payload)

    def test_main_exits_nonzero_on_a_failed_fidelity_gate(self) -> None:
        english = (
            "This chapter explains why translation became a state project "
            "during the Meiji period and how it reshaped modern Japanese prose."
        )
        with tempfile.TemporaryDirectory() as directory:
            workspace = Path(directory)
            source = _write(workspace / "en.md", f"# Chapter One\n\n{english}\n")
            buffer = io.StringIO()
            with contextlib.redirect_stdout(buffer):
                code = kb_ingest.main(
                    [
                        "ingest",
                        str(source),
                        "--output-dir",
                        str(workspace / "out"),
                        "--title",
                        "English Book",
                    ]
                )
        self.assertEqual(code, 1)
        payload = json.loads(buffer.getvalue().strip())
        self.assertEqual(payload["status"], "blocked")

    def test_status_reports_both_gates(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            workspace = Path(directory)
            report = self._ingest(workspace)
            status = kb_ingest.corpus_status(report["workspace"])
        self.assertTrue(status["gates"]["chinese"]["passed"])
        self.assertEqual(status["corpus"]["chunk_count"], report["chunk_count"])
        self.assertEqual(status["sidecar"]["coverage"], 1.0)
        self.assertEqual(status["gates"]["word_fidelity"]["status"], "passed")

    def test_verify_word_uses_the_recorded_title(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            workspace = Path(directory)
            report = self._ingest(workspace)
            verified = kb_ingest.verify_word(report["workspace"])
        self.assertEqual(verified["status"], "passed")
        # The workspace directory name must not be mistaken for the book title.
        self.assertNotIn(
            "docx_core_title_mismatch",
            {warning["code"] for warning in verified["warnings"]},
        )

    def test_verify_word_rejects_an_ambiguous_docx_set(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            workspace = Path(directory)
            report = self._ingest(workspace)
            (Path(report["workspace"]) / "另一个.docx").write_bytes(b"x")
            with self.assertRaises(kb_ingest.IngestError):
                kb_ingest.verify_word(report["workspace"])

    def test_stdin_request_dispatch(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            workspace = Path(directory)
            report = self._ingest(workspace)
            payload = kb_ingest.run_request(
                {"command": "status", "args": {"workspace": report["workspace"]}}
            )
        self.assertEqual(payload["corpus"]["chunk_count"], report["chunk_count"])

    def test_stdin_request_rejects_unknown_command(self) -> None:
        with self.assertRaises(kb_ingest.IngestError):
            kb_ingest.run_request({"command": "nope", "args": {}})

    def test_ingest_report_is_written_to_disk(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            workspace = Path(directory)
            report = self._ingest(workspace)
            report_path = Path(report["report_path"])
            self.assertTrue(report_path.is_file())
            stored = _json.loads(report_path.read_text(encoding="utf-8"))
        self.assertEqual(stored["status"], "passed")
        self.assertEqual(stored["book_title"], "测试书")


if __name__ == "__main__":
    unittest.main()
