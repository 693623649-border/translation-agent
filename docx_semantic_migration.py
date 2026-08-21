"""Migrate an accepted Word publication into semantic Markdown chapters.

The migration is intended for legacy publications whose reviewed DOCX is a
better semantic oracle than their page-shaped OCR Markdown.  It extracts the
document in body order, preserves each true Word footnote reference at its
exact location, and emits a closed one-reference/one-definition Markdown
contract.  It never guesses a reference position.

The implementation is deliberately fail-closed.  A malformed Word footnote
package, duplicate/orphan note IDs, an unexpected chapter heading, or a note
outside a manifest chapter aborts the migration before anything is written.
"""

from __future__ import annotations

from collections import Counter
from dataclasses import asdict, dataclass
import hashlib
import json
import os
from pathlib import Path
import re
import tempfile
import datetime as dt
from typing import Any, Iterable, Mapping, Sequence
import zipfile

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from lxml import etree

from docx_footnotes import W_NS, inspect_docx_footnotes
from publication_semantics import (
    markdown_footnote_contract_sha256,
    parse_markdown_footnotes,
    semantic_audit_summary,
)


W = "{%s}" % W_NS
_MARKER_TOKEN = re.compile(r"\x00FN:(\d+)\x00")


class DocxSemanticMigrationError(ValueError):
    """Raised when an accepted DOCX cannot be migrated without guessing."""


@dataclass(frozen=True)
class FootnoteProvenance:
    stable_id: str
    word_id: int
    occurrence: int
    chapter_id: str
    block_index: int
    block_kind: str
    context_before: str
    context_after: str
    text: str
    text_sha256: str


@dataclass(frozen=True)
class SemanticMarkdownChapter:
    chapter_id: str
    filename: str
    display_title: str
    markdown: str
    footnotes: tuple[FootnoteProvenance, ...]
    pdf_page: int | None = None
    end_pdf_page: int | None = None

    def audit_dict(self) -> dict[str, Any]:
        return {
            "chapter_id": self.chapter_id,
            "filename": self.filename,
            # The accepted DOCX is the reviewed semantic authority during a
            # legacy migration.  Keeping this true prevents a later compile
            # from silently replacing the migrated chapter with page OCR.
            "reviewed_override": True,
            "semantic_source": "accepted-docx-migration",
            "footnote_count": len(self.footnotes),
            "pages": [],
            "issues": [],
            "release_blocked": False,
            "markdown_sha256": hashlib.sha256(
                self.markdown.encode("utf-8")
            ).hexdigest(),
            "footnote_contract_sha256": markdown_footnote_contract_sha256(
                self.markdown
            ),
            "source_page_range": [self.pdf_page, self.end_pdf_page],
            "mappings": [asdict(item) for item in self.footnotes],
        }


@dataclass(frozen=True)
class DocxSemanticMigration:
    source_docx: Path
    source_sha256: str
    chapters: tuple[SemanticMarkdownChapter, ...]
    front_matter: tuple[str, ...]

    @property
    def footnote_count(self) -> int:
        return sum(len(chapter.footnotes) for chapter in self.chapters)

    @property
    def word_id_non_monotonic_positions(self) -> tuple[dict[str, Any], ...]:
        """Return adjacent Word-ID descents in body-reference order.

        Word permits reference IDs to appear out of numeric order.  Such a
        document is valid, but recording each descent makes it explicit that
        stable Markdown IDs were assigned by body occurrence and that no
        definition was accidentally re-sorted during migration.
        """

        mappings = [
            mapping
            for chapter in self.chapters
            for mapping in chapter.footnotes
        ]
        positions: list[dict[str, Any]] = []
        for previous, current in zip(mappings, mappings[1:]):
            if current.word_id >= previous.word_id:
                continue
            positions.append(
                {
                    "occurrence": current.occurrence,
                    "word_id": current.word_id,
                    "previous_occurrence": previous.occurrence,
                    "previous_word_id": previous.word_id,
                    "chapter_id": current.chapter_id,
                    "block_index": current.block_index,
                    "stable_id": current.stable_id,
                }
            )
        return tuple(positions)

    def audit_dict(self) -> dict[str, Any]:
        chapter_audits = [chapter.audit_dict() for chapter in self.chapters]
        return {
            "schema_version": 1,
            "status": "passed",
            "release_blocked": False,
            "generated_by": "docx_semantic_migration",
            "mode": "accepted-docx-semantic-extraction",
            "migration": {
                "mode": "accepted-docx-semantic-extraction",
                "source_docx": str(self.source_docx),
                "source_docx_sha256": self.source_sha256,
                "front_matter": list(self.front_matter),
                "reference_order": "document-occurrence",
                "word_id_non_monotonic_positions": list(
                    self.word_id_non_monotonic_positions
                ),
            },
            "summary": semantic_audit_summary(chapter_audits),
            "chapters": chapter_audits,
        }

    def updated_manifest(self, manifest: Sequence[Mapping[str, Any]]) -> list[dict[str, Any]]:
        by_id = {chapter.chapter_id: chapter for chapter in self.chapters}
        output: list[dict[str, Any]] = []
        for raw in manifest:
            item = dict(raw)
            chapter = by_id.get(str(item.get("id") or ""))
            if chapter is None:
                raise DocxSemanticMigrationError(
                    "migration result does not cover manifest chapter %r" % item.get("id")
                )
            item["semantic_footnote_count"] = len(chapter.footnotes)
            item["semantic_issue_count"] = 0
            item["semantic_source"] = "accepted-docx-migration"
            item["reviewed_override"] = True
            output.append(item)
        return output


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _xml(data: bytes, member: str) -> etree._Element:
    parser = etree.XMLParser(resolve_entities=False, no_network=True, recover=False)
    try:
        return etree.fromstring(data, parser=parser)
    except etree.XMLSyntaxError as exc:
        raise DocxSemanticMigrationError(
            "invalid XML in %s: %s" % (member, exc)
        ) from exc


def _manifest_items(manifest: Sequence[Mapping[str, Any]]) -> list[dict[str, Any]]:
    if isinstance(manifest, (str, bytes)) or not isinstance(manifest, Sequence):
        raise DocxSemanticMigrationError("manifest must be an ordered sequence")
    items = [dict(item) for item in manifest]
    if not items:
        raise DocxSemanticMigrationError("manifest is empty")
    required = ("id", "filename", "display_title")
    for index, item in enumerate(items):
        missing = [key for key in required if not str(item.get(key) or "").strip()]
        if missing:
            raise DocxSemanticMigrationError(
                "manifest item %d is missing %s" % (index, ", ".join(missing))
            )
    for key in required:
        values = [str(item[key]) for item in items]
        duplicates = sorted(value for value, count in Counter(values).items() if count > 1)
        if duplicates:
            raise DocxSemanticMigrationError(
                "manifest has duplicate %s values: %s" % (key, duplicates)
            )
    return items


def _footnote_definitions(path: Path) -> dict[int, str]:
    with zipfile.ZipFile(path, "r") as archive:
        try:
            root = _xml(archive.read("word/footnotes.xml"), "word/footnotes.xml")
        except KeyError as exc:
            raise DocxSemanticMigrationError("DOCX has no footnotes.xml") from exc
    definitions: dict[int, str] = {}
    for note in root.findall(W + "footnote"):
        raw_id = note.get(W + "id")
        try:
            note_id = int(raw_id) if raw_id is not None else None
        except ValueError as exc:
            raise DocxSemanticMigrationError(
                "invalid footnote definition ID %r" % raw_id
            ) from exc
        if note_id is None or note_id <= 0:
            continue
        paragraphs: list[str] = []
        for paragraph in note.findall(".//" + W + "p"):
            parts: list[str] = []
            for node in paragraph.iter():
                if node.tag == W + "t":
                    parts.append(node.text or "")
                elif node.tag == W + "tab":
                    parts.append("\t")
                elif node.tag in (W + "br", W + "cr"):
                    parts.append("\n")
            paragraphs.append("".join(parts).strip())
        value = "\n".join(part for part in paragraphs if part).strip()
        if not value:
            raise DocxSemanticMigrationError(
                "footnote definition %d has no visible text" % note_id
            )
        if note_id in definitions:
            raise DocxSemanticMigrationError(
                "duplicate footnote definition ID %d" % note_id
            )
        definitions[note_id] = value
    return definitions


def _paragraph_raw(paragraph: Paragraph) -> str:
    parts: list[str] = []
    for node in paragraph._p.iter():
        if node.tag == W + "t":
            parts.append(node.text or "")
        elif node.tag == W + "tab":
            parts.append("\t")
        elif node.tag in (W + "br", W + "cr"):
            if node.get(W + "type") != "page":
                parts.append("\n")
        elif node.tag == W + "footnoteReference":
            raw_id = node.get(W + "id")
            try:
                note_id = int(raw_id) if raw_id is not None else None
            except ValueError as exc:
                raise DocxSemanticMigrationError(
                    "invalid footnote reference ID %r" % raw_id
                ) from exc
            if note_id is None or note_id <= 0:
                raise DocxSemanticMigrationError(
                    "body contains non-positive footnote reference %r" % raw_id
                )
            parts.append("\x00FN:%d\x00" % note_id)
    return "".join(parts)


def _escape_markdown(value: str, *, table: bool = False) -> str:
    value = value.replace("\\", "\\\\")
    for character in ("`", "*", "_"):
        value = value.replace(character, "\\" + character)
    if table:
        value = value.replace("|", "\\|")
        value = value.replace("\n", "<br>")
    else:
        value = value.replace("\n", "  \n")
    return value


def _context(raw: str, start: int, end: int, width: int = 80) -> tuple[str, str]:
    def visible(value: str) -> str:
        return _MARKER_TOKEN.sub("", value).replace("\n", " ").strip()

    return visible(raw[max(0, start - width) : start]), visible(raw[end : end + width])


def _serialize_inline(
    raw: str,
    *,
    table: bool,
    chapter_id: str,
    block_index: int,
    block_kind: str,
    definitions: Mapping[int, str],
    seen_word_ids: set[int],
    occurrence_start: int,
) -> tuple[str, list[FootnoteProvenance], int]:
    output: list[str] = []
    mappings: list[FootnoteProvenance] = []
    cursor = 0
    occurrence = occurrence_start
    for match in _MARKER_TOKEN.finditer(raw):
        output.append(_escape_markdown(raw[cursor : match.start()], table=table))
        word_id = int(match.group(1))
        if word_id in seen_word_ids:
            raise DocxSemanticMigrationError(
                "footnote reference ID %d appears more than once" % word_id
            )
        if word_id not in definitions:
            raise DocxSemanticMigrationError(
                "footnote reference ID %d has no definition" % word_id
            )
        seen_word_ids.add(word_id)
        occurrence += 1
        stable_id = "accepted-docx-fn-%06d" % occurrence
        before, after = _context(raw, match.start(), match.end())
        text = definitions[word_id]
        mappings.append(
            FootnoteProvenance(
                stable_id=stable_id,
                word_id=word_id,
                occurrence=occurrence,
                chapter_id=chapter_id,
                block_index=block_index,
                block_kind=block_kind,
                context_before=before,
                context_after=after,
                text=text,
                text_sha256=hashlib.sha256(text.encode("utf-8")).hexdigest(),
            )
        )
        output.append("[^%s]" % stable_id)
        cursor = match.end()
    output.append(_escape_markdown(raw[cursor:], table=table))
    return "".join(output), mappings, occurrence


def _iter_body_blocks(document: Any) -> Iterable[Paragraph | Table]:
    for child in document.element.body.iterchildren():
        if child.tag == W + "p":
            yield Paragraph(child, document)
        elif child.tag == W + "tbl":
            yield Table(child, document)


def extract_docx_semantic_markdown(
    docx_path: os.PathLike[str] | str,
    manifest: Sequence[Mapping[str, Any]],
) -> DocxSemanticMigration:
    """Extract manifest-aligned Markdown chapters from a reviewed DOCX.

    Stable IDs follow *reference occurrence*, not numeric Word ID order.  Word
    IDs are retained in :class:`FootnoteProvenance`, so even a valid document
    containing out-of-order IDs cannot exchange note definitions.
    """

    path = Path(docx_path).resolve()
    items = _manifest_items(manifest)
    inventory = inspect_docx_footnotes(path)
    if not inventory.valid:
        raise DocxSemanticMigrationError(
            "invalid Word footnote package: %s" % ", ".join(inventory.problems)
        )
    definitions = _footnote_definitions(path)
    if Counter(inventory.definition_ids) != Counter(definitions.keys()):
        raise DocxSemanticMigrationError("footnote definition inventory changed while reading")

    document = Document(path)
    expected_titles = [str(item["display_title"]).strip() for item in items]
    actual_titles = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if getattr(paragraph.style, "name", "") == "Heading 1"
    ]
    if actual_titles != expected_titles:
        raise DocxSemanticMigrationError(
            "Heading 1/manifest mismatch: expected=%r actual=%r"
            % (expected_titles, actual_titles)
        )

    # Each item is one Markdown block.  Table rows deliberately stay together;
    # blank lines between them would turn a real table into pipe-delimited prose.
    chapter_lines: list[list[str]] = [[] for _ in items]
    chapter_maps: list[list[FootnoteProvenance]] = [[] for _ in items]
    block_counts = [0 for _ in items]
    front_matter: list[str] = []
    current = -1
    occurrence = 0
    seen_word_ids: set[int] = set()

    for block in _iter_body_blocks(document):
        if isinstance(block, Paragraph):
            style = getattr(block.style, "name", "") or "Normal"
            raw = _paragraph_raw(block)
            if style == "Heading 1":
                current += 1
                if current >= len(items) or raw.strip() != expected_titles[current]:
                    raise DocxSemanticMigrationError("chapter boundary changed during extraction")
                chapter_lines[current].append("# " + _escape_markdown(raw.strip()))
                block_counts[current] += 1
                continue
            if current < 0:
                if _MARKER_TOKEN.search(raw):
                    raise DocxSemanticMigrationError(
                        "footnote reference occurs before the first manifest chapter"
                    )
                if raw.strip():
                    front_matter.append(raw.strip())
                continue
            if not raw.strip() and not _MARKER_TOKEN.search(raw):
                continue
            block_counts[current] += 1
            kind = style
            inline, mappings, occurrence = _serialize_inline(
                raw,
                table=False,
                chapter_id=str(items[current]["id"]),
                block_index=block_counts[current],
                block_kind=kind,
                definitions=definitions,
                seen_word_ids=seen_word_ids,
                occurrence_start=occurrence,
            )
            inline = inline.strip()
            if not inline:
                continue
            chapter_maps[current].extend(mappings)
            if style == "Heading 2":
                line = "## " + inline
            elif style == "Heading 3":
                line = "### " + inline
            elif style.startswith("List Bullet"):
                line = "- " + inline
            elif style.startswith("List Number"):
                line = "1. " + inline
            elif style == "Quote":
                line = "> " + inline.replace("\n", "\n> ")
            else:
                line = inline
            chapter_lines[current].append(line)
            continue

        if current < 0:
            raise DocxSemanticMigrationError("table occurs before the first manifest chapter")
        block_counts[current] += 1
        block_index = block_counts[current]
        rows: list[list[str]] = []
        table_mappings: list[FootnoteProvenance] = []
        for row in block.rows:
            cells: list[str] = []
            seen_cells: set[int] = set()
            for cell in row.cells:
                cell_identity = id(cell._tc)
                if cell_identity in seen_cells:
                    cells.append("")
                    continue
                seen_cells.add(cell_identity)
                parts: list[str] = []
                for paragraph in cell.paragraphs:
                    raw = _paragraph_raw(paragraph)
                    inline, mappings, occurrence = _serialize_inline(
                        raw,
                        table=True,
                        chapter_id=str(items[current]["id"]),
                        block_index=block_index,
                        block_kind="Table",
                        definitions=definitions,
                        seen_word_ids=seen_word_ids,
                        occurrence_start=occurrence,
                    )
                    if inline.strip():
                        parts.append(inline.strip())
                    table_mappings.extend(mappings)
                cells.append("<br>".join(parts))
            rows.append(cells)
        if not rows:
            continue
        width = max(len(row) for row in rows)
        rows = [row + [""] * (width - len(row)) for row in rows]
        chapter_maps[current].extend(table_mappings)
        table_lines = [
            "| " + " | ".join(rows[0]) + " |",
            "| " + " | ".join(["---"] * width) + " |",
            *(
                "| " + " | ".join(row) + " |" for row in rows[1:]
            ),
        ]
        chapter_lines[current].append(
            "\n".join(table_lines)
        )

    if current != len(items) - 1:
        raise DocxSemanticMigrationError("not all manifest chapters were extracted")
    orphan_ids = sorted(set(definitions) - seen_word_ids)
    if orphan_ids:
        raise DocxSemanticMigrationError(
            "orphan Word footnote definitions: %s" % orphan_ids
        )

    chapters: list[SemanticMarkdownChapter] = []
    for index, item in enumerate(items):
        lines = chapter_lines[index]
        mappings = chapter_maps[index]
        definitions_md: list[str] = []
        for mapping in mappings:
            note_lines = mapping.text.splitlines() or [mapping.text]
            definition = "[^%s]: %s" % (
                mapping.stable_id,
                note_lines[0],
            )
            if len(note_lines) > 1:
                definition += "\n" + "\n".join(
                    "    " + line for line in note_lines[1:]
                )
            definitions_md.append(definition)
        markdown = "\n\n".join(lines).rstrip()
        if definitions_md:
            markdown += "\n\n" + "\n\n".join(definitions_md)
        markdown += "\n"
        md_inventory = parse_markdown_footnotes(markdown)
        if not md_inventory.valid or len(md_inventory.definitions) != len(mappings):
            raise DocxSemanticMigrationError(
                "generated Markdown footnote contract is not closed for %s"
                % item["id"]
            )
        chapters.append(
            SemanticMarkdownChapter(
                chapter_id=str(item["id"]),
                filename=str(item["filename"]),
                display_title=str(item["display_title"]),
                markdown=markdown,
                footnotes=tuple(mappings),
                pdf_page=item.get("pdf_page"),
                end_pdf_page=item.get("end_pdf_page"),
            )
        )

    if occurrence != len(definitions):
        raise DocxSemanticMigrationError(
            "reference/definition total mismatch after extraction"
        )
    return DocxSemanticMigration(
        source_docx=path,
        source_sha256=_sha256_file(path),
        chapters=tuple(chapters),
        front_matter=tuple(front_matter),
    )


def write_semantic_migration(
    migration: DocxSemanticMigration,
    chapter_dir: os.PathLike[str] | str,
    *,
    audit_path: os.PathLike[str] | str | None = None,
) -> None:
    """Atomically write extracted chapters and, optionally, their audit JSON."""

    target_dir = Path(chapter_dir)
    target_dir.mkdir(parents=True, exist_ok=True)
    payloads: list[tuple[Path, bytes]] = [
        (target_dir / chapter.filename, chapter.markdown.encode("utf-8"))
        for chapter in migration.chapters
    ]
    if audit_path is not None:
        target_audit = Path(audit_path)
        payloads.append(
            (
                target_audit,
                (json.dumps(migration.audit_dict(), ensure_ascii=False, indent=2) + "\n").encode(
                    "utf-8"
                ),
            )
        )
    staged: list[tuple[Path, Path]] = []
    try:
        for target, data in payloads:
            target.parent.mkdir(parents=True, exist_ok=True)
            descriptor, name = tempfile.mkstemp(prefix=".%s." % target.name, dir=target.parent)
            temporary = Path(name)
            with os.fdopen(descriptor, "wb") as handle:
                handle.write(data)
                handle.flush()
                os.fsync(handle.fileno())
            staged.append((temporary, target))
        for temporary, target in staged:
            os.replace(temporary, target)
    finally:
        for temporary, _ in staged:
            temporary.unlink(missing_ok=True)


def backup_semantic_migration_sources(
    output_dir: os.PathLike[str] | str,
    backup_path: os.PathLike[str] | str,
) -> Path:
    """Create a non-overwriting recovery archive before applying a migration.

    The archive deliberately lives outside ``output_dir`` so the publication
    runtime-hygiene gate never mistakes it for a release artifact.  Only files
    that the semantic migration can replace, plus the accepted DOCX and TOC
    needed to audit a rollback, are included.
    """

    root = Path(output_dir).expanduser().resolve()
    destination = Path(backup_path).expanduser().resolve()
    try:
        destination.relative_to(root)
    except ValueError:
        pass
    else:
        raise DocxSemanticMigrationError(
            "semantic migration backup must be outside the publication output"
        )
    if destination.exists():
        raise DocxSemanticMigrationError(
            "semantic migration backup already exists: %s" % destination
        )

    candidates: list[Path] = []
    for relative in ("chapters.json", "toc.json"):
        path = root / relative
        if path.is_file():
            candidates.append(path)
    for directory in ("chapters", "reviewed_chapters"):
        candidates.extend(sorted((root / directory).glob("*.md")))
    for name in (
        "release-report.json",
        "graph-release-report.json",
        "word-release-report.json",
        "semantic-reconstruction.json",
        "docx-semantic-migration.json",
    ):
        path = root / "audit" / name
        if path.is_file():
            candidates.append(path)
    candidates.extend(sorted(root.glob("*.docx")))

    inventory = []
    for path in candidates:
        resolved = path.resolve()
        try:
            relative = resolved.relative_to(root)
        except ValueError as exc:
            raise DocxSemanticMigrationError(
                "backup source resolves outside the publication output: %s" % path
            ) from exc
        inventory.append(
            {
                "path": relative.as_posix(),
                "sha256": _sha256_file(resolved),
                "size": resolved.stat().st_size,
            }
        )

    destination.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=".%s." % destination.name,
        suffix=".zip",
        dir=destination.parent,
    )
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        manifest = {
            "schema_version": 1,
            "created_at": dt.datetime.now(dt.timezone.utc).isoformat(),
            "source_output": str(root),
            "reviewed_chapters_existed": (root / "reviewed_chapters").is_dir(),
            "files": inventory,
        }
        with zipfile.ZipFile(
            temporary,
            "w",
            compression=zipfile.ZIP_DEFLATED,
            compresslevel=6,
        ) as archive:
            archive.writestr(
                "backup-manifest.json",
                json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
            )
            for item in inventory:
                archive.write(root / item["path"], item["path"])
        # Windows implements ``os.fsync`` with ``_commit``, which rejects a
        # read-only descriptor even though POSIX accepts one.
        with temporary.open("rb+") as handle:
            os.fsync(handle.fileno())
        os.replace(temporary, destination)
    finally:
        temporary.unlink(missing_ok=True)
    return destination


def apply_semantic_migration(
    migration: DocxSemanticMigration,
    output_dir: os.PathLike[str] | str,
    *,
    manifest: Sequence[Mapping[str, Any]] | None = None,
    backup_path: os.PathLike[str] | str | None = None,
) -> dict[str, Any]:
    """Apply a reviewed-DOCX migration as one fail-closed publication bundle.

    This is the production entrypoint.  It keeps canonical chapters,
    ``reviewed_chapters/<id>.md``, both semantic audit records, and
    ``chapters.json`` in sync.  Each file replacement is atomic; when a backup
    path is supplied, a complete recovery archive is created before staging.
    """

    root = Path(output_dir).expanduser().resolve()
    manifest_path = root / "chapters.json"
    if manifest is None:
        try:
            value = json.loads(manifest_path.read_text(encoding="utf-8"))
        except (OSError, UnicodeError, json.JSONDecodeError) as exc:
            raise DocxSemanticMigrationError(
                "cannot read publication chapter manifest: %s" % exc
            ) from exc
        if not isinstance(value, list):
            raise DocxSemanticMigrationError(
                "publication chapter manifest must be a JSON array"
            )
        source_manifest: Sequence[Mapping[str, Any]] = value
    else:
        source_manifest = manifest

    if _sha256_file(migration.source_docx) != migration.source_sha256:
        raise DocxSemanticMigrationError(
            "accepted DOCX changed after semantic extraction"
        )
    updated_manifest = migration.updated_manifest(source_manifest)
    expected_filenames = {chapter.filename for chapter in migration.chapters}
    existing_filenames = {
        path.name for path in (root / "chapters").glob("*.md") if path.is_file()
    }
    unexpected = sorted(existing_filenames - expected_filenames)
    if unexpected:
        raise DocxSemanticMigrationError(
            "publication has unmanifested chapter Markdown: %s" % unexpected
        )

    backup: Path | None = None
    if backup_path is not None:
        backup = backup_semantic_migration_sources(root, backup_path)

    def child(directory: Path, name: str, label: str) -> Path:
        target = (directory / name).resolve()
        try:
            target.relative_to(directory.resolve())
        except ValueError as exc:
            raise DocxSemanticMigrationError(
                "%s resolves outside its publication directory: %r" % (label, name)
            ) from exc
        if target.parent != directory.resolve():
            raise DocxSemanticMigrationError(
                "%s must be a single filename: %r" % (label, name)
            )
        return target

    audit_payload = (
        json.dumps(migration.audit_dict(), ensure_ascii=False, indent=2) + "\n"
    ).encode("utf-8")
    payloads: list[tuple[Path, bytes]] = []
    for chapter in migration.chapters:
        payloads.append(
            (
                child(root / "chapters", chapter.filename, "chapter filename"),
                chapter.markdown.encode("utf-8"),
            )
        )
        payloads.append(
            (
                child(
                    root / "reviewed_chapters",
                    "%s.md" % chapter.chapter_id,
                    "reviewed chapter id",
                ),
                chapter.markdown.encode("utf-8"),
            )
        )
    payloads.extend(
        [
            (root / "audit" / "semantic-reconstruction.json", audit_payload),
            (root / "audit" / "docx-semantic-migration.json", audit_payload),
            (
                manifest_path,
                (
                    json.dumps(updated_manifest, ensure_ascii=False, indent=2) + "\n"
                ).encode("utf-8"),
            ),
        ]
    )

    staged: list[tuple[Path, Path]] = []
    try:
        for target, data in payloads:
            target.parent.mkdir(parents=True, exist_ok=True)
            descriptor, name = tempfile.mkstemp(
                prefix=".%s." % target.name,
                dir=target.parent,
            )
            temporary = Path(name)
            with os.fdopen(descriptor, "wb") as handle:
                handle.write(data)
                handle.flush()
                os.fsync(handle.fileno())
            staged.append((temporary, target))
        for temporary, target in staged:
            os.replace(temporary, target)
    finally:
        for temporary, _target in staged:
            temporary.unlink(missing_ok=True)

    return {
        "chapter_count": len(migration.chapters),
        "footnote_count": migration.footnote_count,
        "source_docx_sha256": migration.source_sha256,
        "backup_path": str(backup) if backup is not None else None,
        "semantic_audit_path": str(
            root / "audit" / "semantic-reconstruction.json"
        ),
        "migration_audit_path": str(
            root / "audit" / "docx-semantic-migration.json"
        ),
    }


__all__ = [
    "apply_semantic_migration",
    "backup_semantic_migration_sources",
    "DocxSemanticMigration",
    "DocxSemanticMigrationError",
    "FootnoteProvenance",
    "SemanticMarkdownChapter",
    "extract_docx_semantic_markdown",
    "write_semantic_migration",
]
