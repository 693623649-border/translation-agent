"""Deterministic release verification for compiled book publications.

The verifier deliberately does not call a language model.  It turns the
repeatable checks that used to be performed by one-off scripts into a small,
stable API that can be run after every compilation.  A failed check is data in
the returned report, not an exception from :func:`verify_publication`.
"""

from __future__ import annotations

import datetime as dt
import hashlib
import html
import json
import os
import posixpath
import re
import tempfile
import unicodedata
import zipfile
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Callable, Iterable
from xml.etree import ElementTree as ET

import rag_knowledge_base
from publication_semantics import (
    markdown_footnote_contract_sha256,
    parse_markdown_footnotes,
)

try:
    import fcntl
except ImportError:  # pragma: no cover - Windows compatibility path.
    fcntl = None


SCHEMA_VERSION = "1.2"
CHECK_IDS = (
    "checkpoints.complete",
    "manifest.valid",
    "chapters.files",
    "reviewed.exact",
    "semantics.integrity",
    "citations.integrity",
    "content.hygiene",
    "epub.structure",
    "docx.structure",
    "docx.render",
    "knowledge_base.structure",
    "pdf.bookmarks",
    "runtime.hygiene",
)
KB_FIELDS = frozenset({"id", "title", "chapter_id", "chapter_order", "content"})
DOCX_BOOK_TITLE_STYLES = frozenset({"Title", "Codex Book Title"})


def _issue(
    code: str,
    message: str,
    *,
    path: Path | str | None = None,
    chapter_id: str | None = None,
    **evidence: Any,
) -> dict[str, Any]:
    value: dict[str, Any] = {"code": code, "message": message}
    if path is not None:
        value["path"] = str(path)
    if chapter_id is not None:
        value["chapter_id"] = chapter_id
    if evidence:
        value["evidence"] = evidence
    return value


def _result(
    summary: str,
    *,
    metrics: dict[str, Any] | None = None,
    issues: list[dict[str, Any]] | None = None,
    warnings: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    return {
        "summary": summary,
        "metrics": metrics or {},
        "issues": issues or [],
        "warnings": warnings or [],
    }


def _sha256_text(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _read_utf8_exact(path: Path) -> str:
    """Decode UTF-8 without Python's universal-newline conversion."""

    with path.open("r", encoding="utf-8", newline="") as handle:
        return handle.read()


def _local_name(element: ET.Element) -> str:
    tag = element.tag
    return tag.rsplit("}", 1)[-1].lower() if isinstance(tag, str) else ""


def _visible_text(element: ET.Element) -> str:
    return "".join(element.itertext()).strip()


def _slugify(value: str, max_len: int = 100) -> str:
    value = unicodedata.normalize("NFKC", value)
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", value)
    value = re.sub(r"\s+", "_", value.strip())
    value = re.sub(r"_+", "_", value).strip(" ._")
    return (value[:max_len] or "untitled").strip(" ._")


def _display_title(item: dict[str, Any]) -> str:
    explicit = str(item.get("display_title") or "").strip()
    if explicit:
        return explicit
    return " ".join(
        value
        for value in (
            str(item.get("index") or "").strip(),
            str(item.get("title") or "").strip(),
        )
        if value
    )


def _select_toc_entries(
    entries: list[dict[str, Any]], granularity: str
) -> list[dict[str, Any]]:
    """Mirror ``book_pipeline.select_entries`` without importing it circularly."""

    with_pages = [item for item in entries if item.get("pdf_page") is not None]
    if granularity == "all":
        return with_pages
    if granularity == "chapter":
        selected = [
            item
            for item in with_pages
            if item.get("kind") in {"frontmatter", "chapter"}
            or (item.get("kind") == "other" and item.get("level") == 1)
        ]
        if selected:
            return selected
        non_parts = [
            item
            for item in with_pages
            if item.get("kind") not in {"part", "frontmatter"}
        ]
        if not non_parts:
            return with_pages
        level = min(int(item.get("level") or 1) for item in non_parts)
        return [item for item in non_parts if int(item.get("level") or 1) == level]
    wanted_kind = "subsection" if granularity == "subsection" else "section"
    selected = [item for item in with_pages if item.get("kind") == wanted_kind]
    if selected:
        return selected
    if granularity == "section":
        selected = [item for item in with_pages if item.get("kind") == "subsection"]
        if selected:
            return selected
    leaves: list[dict[str, Any]] = []
    for index, item in enumerate(with_pages):
        next_item = with_pages[index + 1] if index + 1 < len(with_pages) else None
        if next_item is None or int(next_item.get("level") or 1) <= int(
            item.get("level") or 1
        ):
            leaves.append(item)
    return leaves or with_pages


def _toc_manifest_signature(item: dict[str, Any]) -> tuple[Any, ...]:
    return (
        str(item.get("id") or ""),
        _display_title(item),
        item.get("pdf_page"),
        item.get("level"),
        item.get("kind"),
    )


def _expected_chapter_end_pages(
    entries: list[dict[str, Any]],
    selected: list[dict[str, Any]],
    *,
    granularity: str,
    last_pdf_page: int,
    printed_pages_per_pdf_page: int,
    toc_pages: Iterable[int] = (),
) -> dict[str, int]:
    positions = {
        str(item.get("id") or ""): index for index, item in enumerate(entries)
    }
    normalized_toc_pages = sorted(int(page) for page in toc_pages)
    result: dict[str, int] = {}
    for sequence, item in enumerate(selected):
        start = int(item.get("pdf_page") or 0)
        next_item: dict[str, Any] | None = None
        if granularity == "all":
            if sequence + 1 < len(selected):
                next_item = selected[sequence + 1]
        elif granularity == "chapter" and item.get("kind") in {"frontmatter", "other"}:
            if sequence + 1 < len(selected):
                next_item = selected[sequence + 1]
        else:
            position = positions.get(str(item.get("id") or ""), -1)
            level = int(item.get("level") or 1)
            for candidate in entries[position + 1 :]:
                if candidate.get("pdf_page") is None:
                    continue
                if int(candidate.get("level") or 1) > level:
                    continue
                if int(candidate.get("pdf_page") or 0) >= start:
                    next_item = candidate
                    break
        if next_item is None:
            end = last_pdf_page
        else:
            next_start = int(next_item.get("pdf_page") or start)
            same_level = int(next_item.get("level") or 1) == int(
                item.get("level") or 1
            )
            overlap = bool(
                same_level
                and (
                    granularity in {"section", "subsection"}
                    or (
                        printed_pages_per_pdf_page > 1
                        and next_item.get("printed_page") is not None
                        and int(next_item["printed_page"])
                        % printed_pages_per_pdf_page
                        != 0
                    )
                )
            )
            end = max(start, next_start if overlap else next_start - 1)
        if (
            normalized_toc_pages
            and item.get("kind") in {"frontmatter", "other"}
            and start < normalized_toc_pages[0] <= end
        ):
            end = normalized_toc_pages[0] - 1
        result[str(item.get("id") or "")] = end
    return result


def _path_is_within(path: Path, root: Path) -> bool:
    try:
        path.resolve().relative_to(root.resolve())
    except (OSError, ValueError):
        return False
    return True


def _normalize_cover_title(value: str) -> str:
    value = unicodedata.normalize("NFKC", value).casefold()
    return re.sub(r"[^\w\u3400-\u9fff\u3040-\u30ff]+", "", value)


def _chapter_body(markdown_text: str) -> str:
    lines = markdown_text.splitlines()
    return "\n".join(lines[1:]).strip() if lines else ""


def _strip_reviewed_publication_metadata(markdown_text: str) -> str:
    """Mirror the pipeline's intentionally narrow reviewed-source cleaner."""

    def pagebreak_numbers(line: str) -> set[int] | None:
        anchor = re.fullmatch(
            r"<span\b(?P<attrs>[^>]*)>\s*</span>", line, flags=re.I
        ) or re.fullmatch(r"<span\b(?P<attrs>[^>]*)/\s*>", line, flags=re.I)
        if anchor is None:
            return None
        attributes = anchor.group("attrs")
        if not re.search(
            r"\bepub:type\s*=\s*(['\"])pagebreak\1", attributes, flags=re.I
        ):
            return None
        numbers = {
            int(value)
            for value in re.findall(
                r"\b(?:pdf|printed)[-_]page[-_](\d{1,6})\b",
                attributes,
                flags=re.I,
            )
        }
        title = re.search(
            r"\btitle\s*=\s*(['\"])(\d{1,6})\1", attributes, flags=re.I
        )
        if title is not None:
            numbers.add(int(title.group(2)))
        return numbers

    output: list[str] = []
    removed_marker = False
    pending_page_numbers: set[int] | None = None
    for line in markdown_text.splitlines():
        stripped = line.strip()
        anchor_numbers = pagebreak_numbers(stripped)
        if anchor_numbers is not None:
            removed_marker = True
            pending_page_numbers = anchor_numbers
            continue
        metadata = re.fullmatch(
            r"<!--\s*(?P<key>source[-_ ]pdf|pdf[-_ ]pages|pdf[-_ ]page)"
            r"\s*:\s*(?P<value>.*?)\s*-->",
            stripped,
            flags=re.I,
        )
        if metadata is not None:
            removed_marker = True
            key = re.sub(r"[- ]", "_", metadata.group("key").lower())
            if key == "pdf_page":
                number = re.fullmatch(
                    r"[-—–\s]*(\d{1,6})[-—–\s]*",
                    metadata.group("value"),
                )
                if number is not None:
                    if pending_page_numbers is None:
                        pending_page_numbers = set()
                    pending_page_numbers.add(int(number.group(1)))
            continue
        if pending_page_numbers is not None:
            if not stripped:
                output.append(line)
                continue
            adjacent = re.fullmatch(r"[-—–\s]*(\d{1,6})[-—–\s]*", stripped)
            if adjacent is not None and int(adjacent.group(1)) in pending_page_numbers:
                removed_marker = True
                pending_page_numbers = None
                continue
            pending_page_numbers = None
        output.append(line)

    if not removed_marker:
        return markdown_text
    cleaned = "\n".join(output)
    if markdown_text.endswith(("\n", "\r")):
        cleaned += "\n"
    return cleaned


def _canonical_reviewed_markdown(markdown_text: str) -> str:
    """Mirror compile-time BOM/outer-whitespace and metadata normalization."""

    normalized = (
        markdown_text.replace("\r\n", "\n").replace("\r", "\n")
        .lstrip("\ufeff")
        .strip()
        .rstrip()
        + "\n"
    )
    return _strip_reviewed_publication_metadata(normalized).rstrip() + "\n"


TRACE_PATTERNS: tuple[tuple[str, re.Pattern[str]], ...] = (
    (
        "source_page_marker",
        re.compile(
            r"(?:<!--\s*(?:source[-_ ]pdf|pdf[-_ ]pages?|PDF_PAGE)\s*:|"
            r"epub:type\s*=\s*['\"]pagebreak['\"]|"
            r"\bid\s*=\s*['\"](?:pdf|printed)[-_]page[-_]\d+['\"]|"
            r"(?:^|\n)\s*(?:source[-_ ]page|pdf[-_ ]pages?|physical[-_ ]page|"
            r"来源(?:PDF)?页(?:码)?|原PDF页码)\s*[:：]\s*\d+)",
            re.I,
        ),
    ),
    (
        "decorated_page_number",
        re.compile(
            r"(?m)^[ \t]*(?:[●©◎○◉◯⊙•·◆◇][ \t]*\d{1,3}|"
            r"\d{1,3}[ \t]*[●©◎○◉◯⊙•·◆◇])"
        ),
    ),
    (
        "internal_placeholder",
        re.compile(
            r"(?:TRANSLATION_FAILED|OCR_FAILED|TODO[_ -]?TRANSLATE|"
            r"__PLACEHOLDER__|\[(?:待翻译|待校对|待识别|OCR失败)\]|"
            r"⟦(?:(?:LEX|PROTECT|PLACEHOLDER)[^⟧]*|[RPN]\d{1,8})⟧)",
            re.I,
        ),
    ),
    (
        "model_preamble",
        re.compile(
            r"(?:^|\n)\s*(?:好的|当然|以下是)[，,:：]?\s*"
            r"(?:我将|我会|为您|根据要求|以下为)?[^\n]{0,40}"
            r"(?:翻译|转写|OCR|校对|整理)(?:结果|内容|文本|如下)?[：:]?\s*(?:\n|$)|"
            r"(?:作为(?:一个|一名)?\s*(?:AI|人工智能)|"
            r"I(?:'m| am) (?:an? )?AI|I (?:cannot|can't) (?:read|translate)|"
            r"由\s*(?:GLM|DeepSeek|大语言|多模态|视觉)\s*模型(?:生成|处理))",
            re.I | re.M,
        ),
    ),
    (
        "replacement_character",
        re.compile(r"[\ufffd\x00]|(?:□\s*){3,}|(?:Ã.|Â.|â€|ðŸ){2,}"),
    ),
)


def _trace_issues(
    text: str,
    *,
    path: Path | str,
    chapter_id: str | None = None,
    limit: int = 8,
) -> list[dict[str, Any]]:
    issues: list[dict[str, Any]] = []
    for code, pattern in TRACE_PATTERNS:
        matches = list(pattern.finditer(text))
        if not matches:
            continue
        snippets = [
            re.sub(r"\s+", " ", match.group(0)).strip()[:160]
            for match in matches[:limit]
        ]
        issues.append(
            _issue(
                code,
                f"检测到 {len(matches)} 处发布痕迹或异常文本。",
                path=path,
                chapter_id=chapter_id,
                count=len(matches),
                snippets=snippets,
            )
        )
    return issues


def _suspicious_english_paragraphs(text: str) -> list[str]:
    warnings: list[str] = []
    for paragraph in re.split(r"\n\s*\n", text):
        compact = re.sub(r"\s+", " ", paragraph).strip()
        if len(compact) < 500:
            continue
        letters = re.findall(r"[A-Za-z]", compact)
        cjk = re.findall(r"[\u3400-\u9fff]", compact)
        if len(letters) >= 350 and len(letters) > 12 * max(1, len(cjk)):
            warnings.append(compact[:180])
    return warnings


def _citation_parts(text: str) -> tuple[str, str]:
    boundaries: list[int] = []
    heading = re.search(
        r"^#{2,6}\s*(?:注释|脚注|尾注|注|notes?|endnotes?|footnotes?)\s*$",
        text,
        flags=re.I | re.M,
    )
    if heading is not None:
        boundaries.append(heading.start())
    separator = re.search(
        r"^---\s*$\s*(?=^(?:\*\*\d+\*\*|\[\^[^\]\s]+\]:))",
        text,
        flags=re.M,
    )
    if separator is not None:
        boundaries.append(separator.start())
    footnote_definition = re.search(
        r"^\[\^[^\]\s]+\]:\s*",
        text,
        flags=re.M,
    )
    if footnote_definition is not None:
        boundaries.append(footnote_definition.start())
    if not boundaries:
        return text, ""
    boundary = min(boundaries)
    return text[:boundary], text[boundary:]


def _citation_inventory(text: str) -> dict[str, Any]:
    standard = parse_markdown_footnotes(text)
    if standard.references or standard.definitions:
        # Once a chapter adopts standard Markdown footnotes, that semantic
        # namespace is the only release-blocking citation contract.  Legacy
        # markers can survive temporarily as visible migration residue, but
        # must not be merged with ``[^id]`` IDs and manufacture false missing
        # definitions in an otherwise closed standard-footnote set.
        legacy_markers: list[tuple[int, str]] = [
            (match.start(), match.group(0))
            for match in re.finditer(r"〔\d+〕", standard.body)
        ]
        legacy_plain_ids: list[str] = []
        for match in re.finditer(
            r"(?<![!^])\[(\d+)\](?!\s*\()",
            standard.body,
        ):
            value = match.group(1)
            if value.isdigit() and 1000 <= int(value) <= 2100:
                continue
            legacy_plain_ids.append(value)
            legacy_markers.append((match.start(), match.group(0)))

        references = list(standard.references)
        definitions = [note_id for note_id, _value in standard.definitions]
        reference_set = set(references)
        definition_set = set(definitions)
        sort_key = lambda value: (0, int(value)) if value.isdigit() else (1, value)
        return {
            "references": references,
            "definitions": definitions,
            "reference_count": len(references),
            "unique_reference_count": len(reference_set),
            "definition_count": len(definitions),
            "missing_definitions": sorted(
                reference_set - definition_set,
                key=sort_key,
            ),
            "unused_definitions": sorted(
                definition_set - reference_set,
                key=sort_key,
            ),
            "duplicate_definitions": sorted(
                (
                    note_id
                    for note_id, count in Counter(definitions).items()
                    if count > 1
                ),
                key=sort_key,
            ),
            "square_note_style": False,
            "standard_footnote_style": True,
            "legacy_reference_markers": [
                marker for _position, marker in sorted(legacy_markers)
            ],
            "unmatched_plain_numeric_markers": sorted(
                set(legacy_plain_ids),
                key=sort_key,
            ),
        }

    body, notes = _citation_parts(text)
    references: list[str] = re.findall(r"〔(\d+)〕", body)
    references.extend(re.findall(r"\[\^([^\]\s]+)\]", body))
    plain_numeric_markers = re.findall(
        r"(?<![!^])\[(\d+)\](?!\s*\()",
        body,
    )
    definitions: list[str] = []
    definitions.extend(
        value
        for value in re.findall(r"^\s*\*\*(\d+)\*\*\s+", notes, flags=re.M)
    )
    definitions.extend(
        value
        for value in re.findall(
            r"^\s*\[\^([^\]\s]+)\]:\s*",
            notes,
            flags=re.M,
        )
    )
    definitions.extend(
        value
        for value in re.findall(r"^\s*(\d+)[.．、]\s+", notes, flags=re.M)
    )
    # Page-translated dissertations often retain the source's legacy note
    # layout: ``[167]`` in the prose followed by a bare ``167 Note text``
    # line, with ordinary prose continuing afterwards.  There is no single
    # endnotes boundary to split on, so recognise only bare-number lines whose
    # number is proven by a same-chapter bracketed marker.  This avoids treating
    # years, page references, or arbitrary numbered prose as definitions.
    plain_numeric_set = {
        value
        for value in plain_numeric_markers
        if not (value.isdigit() and 1000 <= int(value) <= 2100)
    }
    definitions.extend(
        value
        for value in re.findall(
            r"^\s*(\d{1,4})[ \t\u3000]+(?=\S)",
            text,
            flags=re.M,
        )
        if value in plain_numeric_set
    )
    definition_set = set(definitions)
    # Bracketed numbers are ambiguous in prose: [2022] may be a literal year,
    # while books such as Binder use [1] as their citation syntax.  Treat a
    # plain numeric marker as a reference only when the same chapter defines
    # that note id.  This recovers real numeric citations without inventing a
    # missing definition for every bracketed year or catalog number.
    matched_plain_markers = [
        value for value in plain_numeric_markers if value in definition_set
    ]
    uses_square_note_style = bool(matched_plain_markers)
    if uses_square_note_style:
        references.extend(
            value
            for value in plain_numeric_markers
            if value in definition_set
            or not (value.isdigit() and 1000 <= int(value) <= 2100)
        )
    else:
        references.extend(matched_plain_markers)
    reference_set = set(references)
    sort_key = lambda value: (0, int(value)) if value.isdigit() else (1, value)
    return {
        "references": references,
        "definitions": definitions,
        "reference_count": len(references),
        "unique_reference_count": len(reference_set),
        "definition_count": len(definitions),
        "missing_definitions": sorted(reference_set - definition_set, key=sort_key),
        "unused_definitions": sorted(definition_set - reference_set, key=sort_key),
        "duplicate_definitions": sorted(
            (
                number
                for number, count in Counter(definitions).items()
                if count > 1
            ),
            key=sort_key,
        ),
        "square_note_style": uses_square_note_style,
        "standard_footnote_style": False,
        "legacy_reference_markers": [],
        "unmatched_plain_numeric_markers": sorted(
            {
                value
                for value in plain_numeric_markers
                if value not in definition_set
                and not (value.isdigit() and 1000 <= int(value) <= 2100)
            },
            key=sort_key,
        ),
    }


def _markdown_html(markdown_text: str) -> str:
    import markdown  # type: ignore[import-not-found]

    output = markdown.markdown(
        markdown_text,
        extensions=["extra", "sane_lists", "footnotes"],
        output_format="xhtml",
    )
    # Mirrors book_pipeline.markdown_to_html: python-markdown serializes some
    # mixed content as named entities ElementTree cannot parse; restore them.
    return re.sub(
        r"&(?!amp;|lt;|gt;|quot;|apos;|#\d+|#x[0-9A-Fa-f]+)([A-Za-z][A-Za-z0-9]+);",
        lambda match: html.entities.html5.get(f"{match.group(1)};", match.group(0)),
        output,
    )


def _canonical_visible_text(value: str) -> str:
    """Normalize reader-visible text for cross-format completeness checks."""

    return re.sub(r"\s+", " ", unicodedata.normalize("NFC", value)).strip()


def _detect_language(value: str) -> str:
    """Mirror the pipeline's deterministic page-language classifier."""

    han = len(re.findall(r"[\u3400-\u9fff]", value))
    kana = len(re.findall(r"[\u3040-\u30ff]", value))
    hangul = len(re.findall(r"[\uac00-\ud7af]", value))
    cyrillic = len(re.findall(r"[\u0400-\u04ff]", value))
    latin = len(re.findall(r"[A-Za-z]", value))
    significant = han + kana + hangul + cyrillic + latin
    if significant < 12:
        return "unknown"
    if kana >= 5 and kana / max(1, han + kana) >= 0.08:
        return "ja"
    if hangul >= max(5, significant * 0.2):
        return "ko"
    if cyrillic >= max(5, significant * 0.25):
        return "ru"
    if han >= max(8, latin * 0.35):
        return "zh"
    if latin >= max(8, significant * 0.6):
        return "en"
    return "other"


def _markdown_visible_text(markdown_text: str) -> str:
    # Mirror the release layer's trailing-page-number discard
    # (``discard_trailing_printed_page``) so EPUB/Word verification compares
    # against the same cleaned text the publishers produced.  A printed page
    # footer split by column-aware OCR can span several digit lines.
    lines = markdown_text.rstrip("\n").split("\n")
    while lines:
        candidate = lines[-1].strip()
        if candidate and re.fullmatch(r"[-—–\s]*\d{1,3}[-—–\s]*", candidate):
            lines.pop()
        else:
            break
    cleaned = "\n".join(lines).rstrip() + "\n"
    root = ET.fromstring(f"<document>{_markdown_html(cleaned)}</document>")
    return "".join(root.itertext())


def _docx_markdown_body(markdown_text: str) -> str:
    """Return the prose rendered into document.xml for a true-footnote DOCX."""

    inventory = parse_markdown_footnotes(markdown_text)
    return re.sub(
        r"\[\^[^\]\s]+\]",
        "",
        inventory.body,
    )


def _expected_docx_footnote_texts(markdown_text: str) -> list[str]:
    inventory = parse_markdown_footnotes(markdown_text)
    definitions = inventory.definition_map()
    return [
        _canonical_visible_text(_markdown_visible_text(definitions[note_id]))
        for note_id in inventory.references
        if note_id in definitions
    ]


def _markdown_heading_signature(
    markdown_text: str,
    *,
    docx_levels: bool = False,
) -> list[tuple[int, str]]:
    """Return every authored heading, preserving order and effective level."""

    root = ET.fromstring(f"<document>{_markdown_html(markdown_text)}</document>")
    signature: list[tuple[int, str]] = []
    for element in root.iter():
        tag = _local_name(element)
        if not re.fullmatch(r"h[1-6]", tag):
            continue
        level = int(tag[1])
        if docx_levels:
            level = min(3, level)
        signature.append((level, _visible_text(element)))
    return signature


def _html_inline_style_fragments(root: ET.Element) -> dict[str, list[str]]:
    """Mirror the DOCX renderer's recursive style inheritance per text node."""

    fragments: dict[str, list[str]] = {"bold": [], "italic": [], "underline": []}

    def append(value: str | None, styles: tuple[bool, bool, bool]) -> None:
        if not value:
            return
        normalized = _canonical_visible_text(value)
        if not normalized:
            return
        for name, enabled in zip(("bold", "italic", "underline"), styles):
            if enabled:
                fragments[name].append(normalized)

    def visit(
        element: ET.Element,
        *,
        bold: bool = False,
        italic: bool = False,
        underline: bool = False,
    ) -> None:
        tag = _local_name(element)
        styles = (
            bold or tag in {"b", "strong", "th"},
            italic or tag in {"em", "i"},
            underline or tag in {"u", "ins"},
        )
        append(element.text, styles)
        for child in element:
            if _local_name(child) != "br":
                visit(
                    child,
                    bold=styles[0],
                    italic=styles[1],
                    underline=styles[2],
                )
            append(child.tail, styles)

    visit(root)
    return fragments


def _docx_expectations(
    selected: list[dict[str, Any]], chapter_texts: dict[str, str]
) -> dict[str, Any]:
    tables: list[list[int]] = []
    quote_paragraphs = 0
    headings: dict[str, list[tuple[int, str]]] = {}
    inline_styles: dict[str, dict[str, list[str]]] = {}
    quotes: dict[str, list[str]] = {}
    for item in selected:
        chapter_id = str(item.get("id") or "")
        source = _docx_markdown_body(chapter_texts.get(chapter_id, ""))
        root = ET.fromstring(f"<document>{_markdown_html(source)}</document>")
        headings[chapter_id] = _markdown_heading_signature(
            source,
            docx_levels=True,
        )
        inline_styles[chapter_id] = _html_inline_style_fragments(root)
        quotes[chapter_id] = []
        for element in root.iter():
            tag = _local_name(element)
            if tag == "table":
                shape: list[int] = []
                for row in element.iter():
                    if _local_name(row) == "tr":
                        shape.append(
                            sum(
                                1
                                for cell in row
                                if _local_name(cell) in {"td", "th"}
                            )
                        )
                tables.append(shape)
            elif tag == "blockquote":
                values = [
                    _canonical_visible_text(_visible_text(node))
                    for node in element
                    if _local_name(node) == "p" and _visible_text(node)
                ]
                quote_paragraphs += len(values)
                quotes[chapter_id].extend(values)
    return {
        "table_shapes": tables,
        "quote_paragraphs": quote_paragraphs,
        "headings": headings,
        "inline_styles": inline_styles,
        "quotes": quotes,
    }


def _pick_artifact(
    output_dir: Path,
    suffix: str,
    *,
    book_title: str | None,
    bookmarked_pdf: bool = False,
) -> tuple[Path | None, list[Path]]:
    if bookmarked_pdf:
        candidates = sorted(output_dir.glob("*_带目录.pdf"))
        expected_suffix = "_带目录.pdf"
    else:
        candidates = sorted(output_dir.glob(f"*{suffix}"))
        expected_suffix = suffix
    if book_title:
        exact = output_dir / f"{_slugify(book_title)}{expected_suffix}"
        if exact in candidates:
            return exact, candidates
    if len(candidates) == 1:
        return candidates[0], candidates
    return None, candidates


class _VerificationContext:
    def __init__(
        self,
        output_dir: Path,
        *,
        source_pdf: Path | None,
        book_title: str | None,
        expected_language: str | None,
        expected_translation_fingerprint: str | None,
        require_translation: bool | None,
        require_all_reviewed: bool,
        chapter_ids: list[str] | None,
    ) -> None:
        self.output_dir = output_dir
        self.source_pdf = source_pdf
        self.book_title = book_title
        self.expected_language = expected_language
        self.expected_translation_fingerprint = expected_translation_fingerprint
        self.require_translation = bool(require_translation)
        self.require_all_reviewed = require_all_reviewed
        self.requested_chapter_ids = chapter_ids
        self.incremental = chapter_ids is not None
        self.artifact_titles: dict[str, str] = {}
        self.manifest: list[dict[str, Any]] = []
        self.selected: list[dict[str, Any]] = []
        self.chapter_texts: dict[str, str] = {}

    @property
    def chapter_dir(self) -> Path:
        return self.output_dir / "chapters"


def _check_checkpoints(context: _VerificationContext) -> dict[str, Any]:
    """Prove source-page OCR coverage and fresh required translations."""

    issues: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []
    pages_dir = context.output_dir / "pages"
    records: dict[int, dict[str, Any]] = {}
    published_unreviewed_pages = {
        page_number
        for item in context.manifest
        if not item.get("reviewed_override")
        for page_number in range(
            int(item.get("pdf_page") or 0),
            int(item.get("end_pdf_page") or item.get("pdf_page") or -1) + 1,
        )
        if page_number > 0
    }
    if not pages_dir.is_dir():
        return _result(
            "逐页检查点目录不存在。",
            issues=[
                _issue(
                    "checkpoint_directory_missing",
                    "完整发布必须保留 pages/page_XXXX.json 审计层。",
                    path=pages_dir,
                )
            ],
        )
    for path in sorted(pages_dir.glob("page_*.json")):
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, UnicodeError, json.JSONDecodeError) as exc:
            issues.append(
                _issue("checkpoint_invalid", str(exc), path=path)
            )
            continue
        if not isinstance(payload, dict):
            issues.append(
                _issue("checkpoint_invalid", "逐页检查点不是 JSON 对象。", path=path)
            )
            continue
        page_number = payload.get("pdf_page")
        if not isinstance(page_number, int) or isinstance(page_number, bool) or page_number < 1:
            issues.append(
                _issue(
                    "checkpoint_page_invalid",
                    "逐页检查点的 pdf_page 必须是正整数。",
                    path=path,
                    value=page_number,
                )
            )
            continue
        if page_number in records:
            issues.append(
                _issue(
                    "checkpoint_page_duplicate",
                    "同一 PDF 页出现多个检查点。",
                    path=path,
                    page=page_number,
                )
            )
            continue
        if not isinstance(payload.get("text"), str) or not payload["text"].strip():
            issues.append(
                _issue(
                    "checkpoint_text_empty",
                    "OCR 检查点缺少非空 text；空白页应保存显式空白页标记。",
                    path=path,
                    page=page_number,
                )
            )
        elif re.search(
            r"(?:\[无法辨认\]|OCR_FAILED|TRANSLATION_FAILED|TODO[_ -]?TRANSLATE|"
            r"⟦(?:(?:LEX|PROTECT|PLACEHOLDER)[^⟧]*|[RPN]\d{1,8})⟧|\ufffd)",
            payload["text"],
            flags=re.I,
        ):
            target = issues if page_number in published_unreviewed_pages else warnings
            target.append(
                _issue(
                    "checkpoint_failure_marker_present",
                    (
                        "进入非审定发布章的 OCR 检查点仍含失败标记、内部 token 或乱码。"
                        if target is issues
                        else "OCR 检查点含失败标记，但对应发布内容已由审定稿覆盖或未进入正文。"
                    ),
                    path=path,
                    page=page_number,
                )
            )
        if not isinstance(payload.get("ocr_model"), str) or not payload["ocr_model"].strip():
            issues.append(
                _issue(
                    "checkpoint_ocr_model_missing",
                    "OCR 检查点缺少模型来源。",
                    path=path,
                    page=page_number,
                )
            )
        records[page_number] = payload

    expected_page_count = 0
    if context.source_pdf is not None and context.source_pdf.is_file():
        try:
            import fitz  # type: ignore[import-not-found]

            with fitz.open(context.source_pdf) as document:
                expected_page_count = document.page_count
        except Exception as exc:
            issues.append(
                _issue(
                    "checkpoint_source_pdf_unreadable",
                    f"{type(exc).__name__}: {exc}",
                    path=context.source_pdf,
                )
            )
    else:
        warnings.append(
            _issue(
                "checkpoint_source_pdf_not_provided",
                "未提供源 PDF，只能检查现有检查点连续性，不能证明全书页数。",
                path=pages_dir,
            )
        )
        expected_page_count = max(records, default=0)

    expected_pages = set(range(1, expected_page_count + 1))
    actual_pages = set(records)
    missing_pages = sorted(expected_pages - actual_pages)
    extra_pages = sorted(actual_pages - expected_pages)
    if missing_pages or extra_pages:
        issues.append(
            _issue(
                "checkpoint_page_coverage_mismatch",
                "逐页 OCR 检查点没有精确覆盖源 PDF。",
                path=pages_dir,
                missing=missing_pages[:50],
                missing_count=len(missing_pages),
                extra=extra_pages[:50],
                extra_count=len(extra_pages),
            )
        )

    translated_page_count = 0
    translation_required_pages: list[int] = []
    stale_translation_pages: list[int] = []
    non_chinese_languages = {"ja", "en", "ko", "ru", "other"}
    for page_number, payload in sorted(records.items()):
        text = str(payload.get("text") or "")
        effective_text = text
        proofread_text = str(payload.get("proofread_text") or "")
        if (
            proofread_text.strip()
            and payload.get("proofread_source_sha256") == _sha256_text(text)
            and payload.get("proofread_provider")
            and payload.get("proofread_model")
            and payload.get("proofread_language")
        ):
            effective_text = proofread_text
        language = str(payload.get("language") or "").lower()
        if language in {"", "unknown"}:
            language = _detect_language(effective_text)
        if (
            not context.require_translation
            or page_number not in published_unreviewed_pages
            or language not in non_chinese_languages
        ):
            continue
        translation_required_pages.append(page_number)
        fresh = bool(
            str(payload.get("translated_text") or "").strip()
            and payload.get("translation_source_sha256")
            == _sha256_text(effective_text)
            and payload.get("translation_provider")
            and payload.get("translation_model")
            and str(payload.get("translation_target_language") or "")
            in {"简体中文", "zh-CN", "zh_CN"}
        )
        if (
            fresh
            and context.expected_translation_fingerprint
            and payload.get("translation_fingerprint")
            != context.expected_translation_fingerprint
        ):
            fresh = False
        if fresh:
            translated_page_count += 1
        else:
            stale_translation_pages.append(page_number)
    if stale_translation_pages:
        issues.append(
            _issue(
                "checkpoint_translation_stale_or_missing",
                "面向中文发布的非中文 OCR 页缺少当前、可追溯的译文。",
                path=pages_dir,
                pages=stale_translation_pages[:50],
                count=len(stale_translation_pages),
                fingerprint_required=bool(context.expected_translation_fingerprint),
            )
        )

    return _result(
        "逐页 OCR 覆盖完整，所有应译页均有新鲜译文。"
        if not issues
        else "逐页 OCR 或译文检查点不完整。",
        metrics={
            "expected_page_count": expected_page_count,
            "checkpoint_page_count": len(records),
            "translation_required": context.require_translation,
            "translation_required_page_count": len(translation_required_pages),
            "translated_page_count": translated_page_count,
        },
        issues=issues,
        warnings=warnings,
    )


def _check_manifest(context: _VerificationContext) -> dict[str, Any]:
    path = context.output_dir / "chapters.json"
    issues: list[dict[str, Any]] = []
    if not path.is_file():
        return _result(
            "章节清单不存在。",
            issues=[_issue("manifest_missing", "缺少 chapters.json。", path=path)],
        )
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        return _result(
            "章节清单无法读取。",
            issues=[_issue("manifest_invalid_json", str(exc), path=path)],
        )
    if not isinstance(payload, list) or not payload:
        return _result(
            "章节清单不是非空数组。",
            issues=[
                _issue(
                    "manifest_invalid_type",
                    "chapters.json 必须是非空 JSON 数组。",
                    path=path,
                )
            ],
        )

    valid_items: list[dict[str, Any]] = []
    for position, raw_item in enumerate(payload, start=1):
        if not isinstance(raw_item, dict):
            issues.append(
                _issue(
                    "manifest_item_invalid",
                    f"第 {position} 条章节记录不是对象。",
                    path=path,
                )
            )
            continue
        item = dict(raw_item)
        chapter_id = str(item.get("id") or "").strip()
        filename = str(item.get("filename") or "").strip()
        title = _display_title(item)
        sequence = item.get("sequence")
        safe_chapter_id = bool(chapter_id) and all(
            separator not in chapter_id for separator in ("/", "\\")
        ) and chapter_id not in {".", ".."}
        safe_filename = (
            bool(filename)
            and Path(filename).name == filename
            and "/" not in filename
            and "\\" not in filename
            and filename.endswith(".md")
        )
        if not chapter_id:
            issues.append(
                _issue(
                    "chapter_id_missing",
                    f"第 {position} 条记录缺少 id。",
                    path=path,
                )
            )
        elif not safe_chapter_id:
            issues.append(
                _issue(
                    "chapter_id_invalid",
                    f"章节 {chapter_id!r} 的 id 不是安全的路径组件。",
                    path=path,
                    chapter_id=chapter_id,
                )
            )
        if not safe_filename:
            issues.append(
                _issue(
                    "chapter_filename_invalid",
                    f"章节 {chapter_id or position} 的 filename 不是安全的 Markdown 文件名。",
                    path=path,
                    chapter_id=chapter_id or None,
                    filename=filename,
                )
            )
        if not title:
            issues.append(
                _issue(
                    "chapter_title_missing",
                    f"章节 {chapter_id or position} 缺少标题。",
                    path=path,
                    chapter_id=chapter_id or None,
                )
            )
        if not isinstance(sequence, int) or isinstance(sequence, bool):
            issues.append(
                _issue(
                    "chapter_sequence_invalid",
                    f"章节 {chapter_id or position} 的 sequence 不是整数。",
                    path=path,
                    chapter_id=chapter_id or None,
                )
            )
        start_page = item.get("pdf_page")
        end_page = item.get("end_pdf_page")
        if (
            not isinstance(start_page, int)
            or isinstance(start_page, bool)
            or start_page < 1
            or not isinstance(end_page, int)
            or isinstance(end_page, bool)
            or end_page < start_page
        ):
            issues.append(
                _issue(
                    "chapter_page_range_invalid",
                    "章节 PDF 起止页必须是正整数且 end_pdf_page >= pdf_page。",
                    path=path,
                    chapter_id=chapter_id or None,
                    pdf_page=start_page,
                    end_pdf_page=end_page,
                )
            )
        if (
            safe_chapter_id
            and safe_filename
            and title
            and isinstance(sequence, int)
            and not isinstance(sequence, bool)
        ):
            valid_items.append(item)

    ids = [str(item.get("id") or "") for item in valid_items]
    filenames = [str(item.get("filename") or "") for item in valid_items]
    titles = [_display_title(item) for item in valid_items]
    sequences = [item.get("sequence") for item in valid_items]
    duplicate_ids = sorted(value for value, count in Counter(ids).items() if value and count > 1)
    duplicate_filenames = sorted(
        value for value, count in Counter(filenames).items() if value and count > 1
    )
    duplicate_titles = sorted(
        value for value, count in Counter(titles).items() if value and count > 1
    )
    if duplicate_ids:
        issues.append(
            _issue(
                "duplicate_chapter_ids",
                "章节 id 不唯一。",
                path=path,
                values=duplicate_ids,
            )
        )
    if duplicate_filenames:
        issues.append(
            _issue(
                "duplicate_chapter_filenames",
                "章节文件名不唯一。",
                path=path,
                values=duplicate_filenames,
            )
        )
    if duplicate_titles:
        issues.append(
            _issue(
                "duplicate_chapter_titles",
                "章节显示标题不唯一，目录与阅读器定位会产生歧义。",
                path=path,
                values=duplicate_titles,
            )
        )
    if all(isinstance(value, int) and not isinstance(value, bool) for value in sequences):
        expected = list(range(1, len(valid_items) + 1))
        if sequences != expected:
            issues.append(
                _issue(
                    "chapter_sequence_non_contiguous",
                    "章节 sequence 必须按清单顺序从 1 连续递增。",
                    path=path,
                    expected=expected,
                    actual=sequences,
                )
            )

    toc_path = context.output_dir / "toc.json"
    try:
        toc_payload = json.loads(toc_path.read_text(encoding="utf-8"))
        toc_entries = toc_payload.get("entries") if isinstance(toc_payload, dict) else None
        if not isinstance(toc_entries, list) or not toc_entries or not all(
            isinstance(item, dict) for item in toc_entries
        ):
            raise ValueError("toc.json.entries 必须是非空对象数组。")
        granularities = {
            str(item.get("granularity") or "")
            for item in valid_items
            if str(item.get("granularity") or "")
        }
        if len(granularities) > 1 or not granularities.issubset(
            {"chapter", "section", "subsection", "all"}
        ):
            issues.append(
                _issue(
                    "manifest_granularity_invalid",
                    "manifest 中的 granularity 必须统一且为受支持值。",
                    path=path,
                    values=sorted(granularities),
                )
            )
            candidate_granularities: list[str] = []
        elif granularities:
            candidate_granularities = [next(iter(granularities))]
        else:
            # Older manifests did not persist this field; exact TOC matching
            # safely infers the historical choice without forcing a rebuild.
            candidate_granularities = [
                "chapter",
                "section",
                "subsection",
                "all",
            ]
        actual_signature = [_toc_manifest_signature(item) for item in valid_items]
        matching_granularities = []
        candidate_counts: dict[str, int] = {}
        selected_by_granularity: dict[str, list[dict[str, Any]]] = {}
        for granularity in candidate_granularities:
            selected_toc = _select_toc_entries(toc_entries, granularity)
            selected_by_granularity[granularity] = selected_toc
            expected_signature = [
                _toc_manifest_signature(item) for item in selected_toc
            ]
            candidate_counts[granularity] = len(expected_signature)
            if actual_signature == expected_signature:
                matching_granularities.append(granularity)
        if not matching_granularities:
            issues.append(
                _issue(
                    "manifest_toc_coverage_mismatch",
                    "章节清单没有完整对应 toc.json 在所选编译粒度下的标题集合与顺序。",
                    path=path,
                    manifest_count=len(actual_signature),
                    candidate_counts=candidate_counts,
                )
            )
        else:
            chosen_granularity = matching_granularities[0]
            last_pdf_page = max(
                (int(item.get("end_pdf_page") or 0) for item in valid_items),
                default=0,
            )
            if context.source_pdf is not None and context.source_pdf.is_file():
                try:
                    import fitz  # type: ignore[import-not-found]

                    with fitz.open(context.source_pdf) as source_document:
                        last_pdf_page = source_document.page_count
                except Exception as exc:
                    issues.append(
                        _issue(
                            "manifest_source_pdf_unreadable",
                            f"{type(exc).__name__}: {exc}",
                            path=context.source_pdf,
                        )
                    )
            expected_ends = _expected_chapter_end_pages(
                toc_entries,
                selected_by_granularity[chosen_granularity],
                granularity=chosen_granularity,
                last_pdf_page=last_pdf_page,
                printed_pages_per_pdf_page=int(
                    toc_payload.get("printed_pages_per_pdf_page") or 1
                ),
                toc_pages=toc_payload.get("toc_pdf_pages") or [],
            )
            mismatched_ranges = [
                {
                    "id": str(item.get("id") or ""),
                    "expected_end_pdf_page": expected_ends.get(
                        str(item.get("id") or "")
                    ),
                    "actual_end_pdf_page": item.get("end_pdf_page"),
                }
                for item in valid_items
                if item.get("end_pdf_page")
                != expected_ends.get(str(item.get("id") or ""))
            ]
            if mismatched_ranges:
                issues.append(
                    _issue(
                        "manifest_page_ranges_mismatch",
                        "章节结束页没有按 toc.json、编译粒度和源 PDF 正确派生。",
                        path=path,
                        values=mismatched_ranges[:50],
                        count=len(mismatched_ranges),
                    )
                )
            out_of_bounds = [
                str(item.get("id") or "")
                for item in valid_items
                if isinstance(item.get("end_pdf_page"), int)
                and int(item["end_pdf_page"]) > last_pdf_page
            ]
            if out_of_bounds:
                issues.append(
                    _issue(
                        "manifest_page_ranges_out_of_bounds",
                        "章节页范围超出源 PDF。",
                        path=path,
                        values=out_of_bounds,
                    )
                )
    except (OSError, UnicodeError, json.JSONDecodeError, ValueError) as exc:
        issues.append(
            _issue(
                "manifest_toc_invalid",
                str(exc),
                path=toc_path,
            )
        )

    context.manifest = valid_items
    if context.requested_chapter_ids is None:
        context.selected = list(valid_items)
    else:
        requested = list(dict.fromkeys(context.requested_chapter_ids))
        by_id = {str(item.get("id") or ""): item for item in valid_items}
        by_sequence = {
            str(item.get("sequence")): item
            for item in valid_items
            if isinstance(item.get("sequence"), int)
            and not isinstance(item.get("sequence"), bool)
        }
        resolved_ids: list[str] = []
        unknown: list[str] = []
        for selector in requested:
            item = by_id.get(selector)
            if item is None and selector.isdecimal():
                item = by_sequence.get(str(int(selector)))
            if item is None:
                unknown.append(selector)
            else:
                resolved_ids.append(str(item.get("id") or ""))
        if unknown:
            issues.append(
                _issue(
                    "chapter_selection_unknown",
                    "增量验证包含未知章节 id 或 sequence。",
                    path=path,
                    values=unknown,
                )
            )
        selected_ids = set(resolved_ids)
        context.selected = [
            item for item in valid_items if str(item.get("id") or "") in selected_ids
        ]
        if not requested:
            issues.append(
                _issue(
                    "chapter_selection_empty",
                    "chapter_ids 不能为空列表。",
                    path=path,
                )
            )

    return _result(
        "章节清单结构有效。" if not issues else "章节清单存在结构问题。",
        metrics={
            "chapter_count": len(valid_items),
            "selected_chapter_count": len(context.selected),
            "reviewed_chapter_count": sum(
                bool(item.get("reviewed_override")) for item in valid_items
            ),
        },
        issues=issues,
    )


def _check_chapter_files(context: _VerificationContext) -> dict[str, Any]:
    issues: list[dict[str, Any]] = []
    if not context.manifest:
        return _result(
            "无法在缺少有效清单时验证章节文件。",
            issues=[_issue("manifest_dependency_failed", "chapters.json 未通过基础解析。")],
        )
    if not context.chapter_dir.is_dir():
        return _result(
            "章节目录不存在。",
            issues=[
                _issue(
                    "chapter_directory_missing",
                    "缺少 chapters 目录。",
                    path=context.chapter_dir,
                )
            ],
        )

    for item in context.selected:
        chapter_id = str(item.get("id") or "")
        filename = str(item.get("filename") or "")
        path = context.chapter_dir / filename
        if not _path_is_within(path, context.chapter_dir):
            issues.append(
                _issue(
                    "chapter_path_outside_output",
                    "章节路径解析到了 chapters 目录之外。",
                    path=path,
                    chapter_id=chapter_id,
                )
            )
            continue
        if not path.is_file():
            issues.append(
                _issue(
                    "chapter_file_missing",
                    "清单中的章节文件不存在。",
                    path=path,
                    chapter_id=chapter_id,
                )
            )
            continue
        try:
            text = path.read_text(encoding="utf-8")
        except (OSError, UnicodeError) as exc:
            issues.append(
                _issue(
                    "chapter_file_unreadable",
                    str(exc),
                    path=path,
                    chapter_id=chapter_id,
                )
            )
            continue
        context.chapter_texts[chapter_id] = text
        h1_titles = re.findall(r"^#\s+(.+?)\s*$", text, flags=re.M)
        expected_title = _display_title(item)
        if h1_titles != [expected_title]:
            issues.append(
                _issue(
                    "chapter_h1_mismatch",
                    "章节必须且只能包含一个与清单标题完全一致的 H1。",
                    path=path,
                    chapter_id=chapter_id,
                    expected=[expected_title],
                    actual=h1_titles,
                )
            )
        if not _chapter_body(text):
            issues.append(
                _issue(
                    "chapter_body_empty",
                    "章节正文为空。",
                    path=path,
                    chapter_id=chapter_id,
                )
            )

    if not context.incremental:
        expected_files = {
            str(item.get("filename") or "") for item in context.manifest
        }
        actual_files = {path.name for path in context.chapter_dir.glob("*.md")}
        extras = sorted(actual_files - expected_files)
        if extras:
            issues.append(
                _issue(
                    "unmanifested_chapter_files",
                    "chapters 目录包含未登记的 Markdown。",
                    path=context.chapter_dir,
                    values=extras,
                )
            )

    return _result(
        "章节文件、标题与正文完整。" if not issues else "章节文件存在缺失或结构错误。",
        metrics={
            "selected_chapter_count": len(context.selected),
            "readable_chapter_count": len(context.chapter_texts),
            "character_count": sum(len(value) for value in context.chapter_texts.values()),
        },
        issues=issues,
    )


def _check_reviewed_exact(context: _VerificationContext) -> dict[str, Any]:
    issues: list[dict[str, Any]] = []
    reviewed_dir = context.output_dir / "reviewed_chapters"
    selected_reviewed = [item for item in context.selected if item.get("reviewed_override")]
    if context.require_all_reviewed:
        missing_reviewed = [
            str(item.get("id") or "")
            for item in context.selected
            if not item.get("reviewed_override")
        ]
        if missing_reviewed:
            issues.append(
                _issue(
                    "not_all_chapters_reviewed",
                    "要求全书审定，但仍有章节未标记 reviewed_override。",
                    path=context.output_dir / "chapters.json",
                    values=missing_reviewed,
                )
            )

    exact_count = 0
    for item in selected_reviewed:
        chapter_id = str(item.get("id") or "")
        source_path = reviewed_dir / f"{chapter_id}.md"
        generated_path = context.chapter_dir / str(item.get("filename") or "")
        if not _path_is_within(source_path, reviewed_dir) or not _path_is_within(
            generated_path, context.chapter_dir
        ):
            issues.append(
                _issue(
                    "reviewed_path_outside_output",
                    "审定稿或成品章节路径解析到了允许目录之外。",
                    path=source_path,
                    chapter_id=chapter_id,
                )
            )
            continue
        if not source_path.is_file():
            issues.append(
                _issue(
                    "reviewed_source_missing",
                    "reviewed_override 为真，但人工审定稿不存在。",
                    path=source_path,
                    chapter_id=chapter_id,
                )
            )
            continue
        try:
            reviewed = _read_utf8_exact(source_path)
            generated = (
                _read_utf8_exact(generated_path)
                if generated_path.is_file()
                else None
            )
        except (OSError, UnicodeError) as exc:
            issues.append(
                _issue(
                    "reviewed_source_unreadable",
                    str(exc),
                    path=source_path,
                    chapter_id=chapter_id,
                )
            )
            continue
        expected = _canonical_reviewed_markdown(reviewed)
        # The reviewed source may legitimately contain a BOM or outer editor
        # whitespace because compile normalizes those once.  The generated
        # publication chapter itself must already be the exact normalized
        # bytes; normalizing it again would hide downstream tampering.
        actual = generated
        if actual != expected:
            issues.append(
                _issue(
                    "reviewed_content_mismatch",
                    "成品章节与去除显式页码元数据后的人工审定稿不完全一致。",
                    path=generated_path,
                    chapter_id=chapter_id,
                    reviewed_sha256=_sha256_text(expected),
                    generated_sha256=(
                        _sha256_text(actual) if actual is not None else None
                    ),
                )
            )
        else:
            exact_count += 1

    return _result(
        "人工审定稿与章节成品逐字一致。"
        if not issues
        else "人工审定稿与章节成品不一致。",
        metrics={
            "selected_reviewed_count": len(selected_reviewed),
            "exact_match_count": exact_count,
            "require_all_reviewed": context.require_all_reviewed,
        },
        issues=issues,
    )


def _check_semantics(context: _VerificationContext) -> dict[str, Any]:
    """Validate the stable semantic-footnote layer before any publisher gate."""

    issues: list[dict[str, Any]] = []
    audit_path = context.output_dir / "audit" / "semantic-reconstruction.json"
    audit_payload: dict[str, Any] = {}
    try:
        value = json.loads(audit_path.read_text(encoding="utf-8"))
        if not isinstance(value, dict):
            raise ValueError("semantic audit root must be an object")
        audit_payload = value
    except (OSError, UnicodeError, json.JSONDecodeError, ValueError) as exc:
        issues.append(
            _issue(
                "semantic_audit_missing_or_invalid",
                "发布前必须存在可解析的语义重建审计记录。",
                path=audit_path,
                error=str(exc),
            )
        )

    summary_payload = audit_payload.get("summary")
    summary_release_blocked = bool(
        isinstance(summary_payload, dict)
        and summary_payload.get("release_blocked") is True
    )
    release_block_signals: list[str] = []
    if audit_payload.get("status") == "blocked":
        release_block_signals.append("status")
    if audit_payload.get("release_blocked") is True:
        release_block_signals.append("release_blocked")
    if summary_release_blocked:
        release_block_signals.append("summary.release_blocked")
    if release_block_signals:
        issues.append(
            _issue(
                "semantic_audit_release_blocked",
                "语义重建审计的顶层状态阻止发布。",
                path=audit_path,
                signals=release_block_signals,
                status=audit_payload.get("status"),
                release_blocked=audit_payload.get("release_blocked"),
                summary_release_blocked=summary_release_blocked,
            )
        )

    raw_chapters = audit_payload.get("chapters")
    audit_chapters = raw_chapters if isinstance(raw_chapters, list) else []
    audit_by_id: dict[str, dict[str, Any]] = {}
    duplicate_audit_ids: list[str] = []
    for raw in audit_chapters:
        if not isinstance(raw, dict):
            continue
        chapter_id = str(raw.get("chapter_id") or "")
        if not chapter_id:
            continue
        if chapter_id in audit_by_id:
            duplicate_audit_ids.append(chapter_id)
        else:
            audit_by_id[chapter_id] = raw
    if duplicate_audit_ids:
        issues.append(
            _issue(
                "semantic_audit_duplicate_chapters",
                "语义审计含重复章节记录。",
                path=audit_path,
                values=sorted(set(duplicate_audit_ids)),
            )
        )

    footnote_count = 0
    blocking_audit_issue_count = 0
    for item in context.selected:
        chapter_id = str(item.get("id") or "")
        chapter_path = context.chapter_dir / str(item.get("filename") or "")
        text = context.chapter_texts.get(chapter_id)
        if text is None:
            continue
        inventory = parse_markdown_footnotes(text)
        footnote_count += len(inventory.definitions)
        for code, message, values in (
            (
                "semantic_markdown_duplicate_definitions",
                "Markdown 脚注定义不唯一。",
                inventory.duplicate_definitions,
            ),
            (
                "semantic_markdown_missing_definitions",
                "Markdown 脚注引用缺少定义。",
                inventory.missing_definitions,
            ),
            (
                "semantic_markdown_unused_definitions",
                "Markdown 脚注定义没有正文引用。",
                inventory.unused_definitions,
            ),
            (
                "semantic_markdown_duplicate_references",
                "同一 Markdown 脚注被重复引用，未满足当前一对一发布契约。",
                inventory.duplicate_references,
            ),
        ):
            if values:
                issues.append(
                    _issue(
                        code,
                        message,
                        path=chapter_path,
                        chapter_id=chapter_id,
                        values=list(values),
                    )
                )

        legacy = _citation_inventory(text)
        if (
            not inventory.definitions
            and legacy["reference_count"] > 0
            and legacy["definition_count"] > 0
        ):
            issues.append(
                _issue(
                    "semantic_legacy_notes_unresolved",
                    "章节仍以正文字符串模拟脚注，必须先转换为标准语义脚注。",
                    path=chapter_path,
                    chapter_id=chapter_id,
                    reference_count=legacy["reference_count"],
                    definition_count=legacy["definition_count"],
                )
            )

        audited = audit_by_id.get(chapter_id)
        if audited is None:
            issues.append(
                _issue(
                    "semantic_audit_chapter_missing",
                    "语义审计未覆盖发布章节。",
                    path=audit_path,
                    chapter_id=chapter_id,
                )
            )
            continue
        expected_markdown_digest = _sha256_file(chapter_path)
        actual_markdown_digest = str(audited.get("markdown_sha256") or "")
        if actual_markdown_digest != expected_markdown_digest:
            issues.append(
                _issue(
                    "semantic_markdown_digest_stale",
                    "章节 Markdown 内容与语义审计摘要不一致。",
                    path=chapter_path,
                    chapter_id=chapter_id,
                    expected_sha256=expected_markdown_digest,
                    audit_sha256=actual_markdown_digest or None,
                )
            )
        expected_contract = markdown_footnote_contract_sha256(text)
        actual_contract = str(audited.get("footnote_contract_sha256") or "")
        if actual_contract != expected_contract:
            issues.append(
                _issue(
                    "semantic_footnote_contract_stale",
                    "章节脚注关系在语义审计后发生变化。",
                    path=chapter_path,
                    chapter_id=chapter_id,
                    expected_sha256=expected_contract,
                    audit_sha256=actual_contract or None,
                )
            )
        audited_count = audited.get("footnote_count")
        if not isinstance(audited_count, int) or audited_count != len(
            inventory.definitions
        ):
            issues.append(
                _issue(
                    "semantic_footnote_count_mismatch",
                    "语义审计脚注数与当前章节不一致。",
                    path=audit_path,
                    chapter_id=chapter_id,
                    expected=len(inventory.definitions),
                    actual=audited_count,
                )
            )
        blocking = [
            value
            for value in (audited.get("issues") or [])
            if isinstance(value, dict) and bool(value.get("blocking", True))
        ]
        if bool(audited.get("release_blocked")) and not blocking:
            blocking = [
                {
                    "code": "semantic_release_blocked",
                    "message": "语义审计将本章标记为不可发布。",
                }
            ]
        blocking_audit_issue_count += len(blocking)
        for value in blocking:
            issues.append(
                _issue(
                    "semantic_audit_blocking_issue",
                    "脚注引用落点仍有未解决的不确定性。",
                    path=audit_path,
                    chapter_id=chapter_id,
                    semantic_code=str(value.get("code") or "unknown"),
                    semantic_message=str(value.get("message") or ""),
                    source_page=value.get("source_page"),
                    note_label=value.get("note_label"),
                    semantic_evidence=value.get("evidence") or {},
                )
            )

    return _result(
        "语义脚注闭环且无待复核引用落点。"
        if not issues
        else "语义重建未达到发布条件。",
        metrics={
            "audit_path": str(audit_path),
            "audited_chapter_count": len(audit_by_id),
            "selected_chapter_count": len(context.selected),
            "footnote_count": footnote_count,
            "blocking_audit_issue_count": blocking_audit_issue_count,
        },
        issues=issues,
    )


def _check_citations(context: _VerificationContext) -> dict[str, Any]:
    issues: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []
    inventories: dict[str, Any] = {}
    for item in context.selected:
        chapter_id = str(item.get("id") or "")
        text = context.chapter_texts.get(chapter_id)
        if text is None:
            continue
        inventory = _citation_inventory(text)
        inventories[chapter_id] = {
            key: value
            for key, value in inventory.items()
            if key not in {"references", "definitions"}
        }
        standard_footnotes = parse_markdown_footnotes(text)
        strict = bool(item.get("reviewed_override")) or bool(
            standard_footnotes.references or standard_footnotes.definitions
        )
        if inventory["legacy_reference_markers"]:
            warnings.append(
                _issue(
                    "legacy_citation_markers_ignored",
                    "标准 Markdown 脚注已启用；正文遗留的旧式引注标记未纳入发布阻断闭环，请迁移或人工确认。",
                    path=context.chapter_dir / str(item.get("filename") or ""),
                    chapter_id=chapter_id,
                    values=inventory["legacy_reference_markers"],
                )
            )
        if inventory["duplicate_definitions"]:
            target = issues if strict else warnings
            target.append(
                _issue(
                    "duplicate_note_definitions",
                    "同一尾注编号出现多次。",
                    path=context.chapter_dir / str(item.get("filename") or ""),
                    chapter_id=chapter_id,
                    values=inventory["duplicate_definitions"],
                )
            )
        if inventory["missing_definitions"]:
            target = issues if strict else warnings
            target.append(
                _issue(
                    "citation_definition_missing",
                    "正文引注没有对应尾注定义。",
                    path=context.chapter_dir / str(item.get("filename") or ""),
                    chapter_id=chapter_id,
                    values=inventory["missing_definitions"],
                )
            )
        if (
            inventory["unmatched_plain_numeric_markers"]
            and not inventory["square_note_style"]
            and not inventory["standard_footnote_style"]
        ):
            warnings.append(
                _issue(
                    "ambiguous_plain_numeric_markers",
                    "发现无法根据本章尾注体系判定的 [n] 标记，请人工确认。",
                    path=context.chapter_dir / str(item.get("filename") or ""),
                    chapter_id=chapter_id,
                    values=inventory["unmatched_plain_numeric_markers"],
                )
            )
        if inventory["unused_definitions"]:
            # In an explicitly reviewed chapter that uses body citations, the
            # relationship is bidirectional: deleting one reference must not
            # leave a release-passing orphan definition.  A notes-only source
            # list remains a visible warning because it can be intentional.
            target = (
                issues
                if strict and inventory["reference_count"] > 0
                else warnings
            )
            target.append(
                _issue(
                    "unused_note_definitions",
                    (
                        "审定章存在正文未引用的尾注定义。"
                        if target is issues
                        else "存在正文未引用的尾注定义，请确认这是否为原稿设计。"
                    ),
                    path=context.chapter_dir / str(item.get("filename") or ""),
                    chapter_id=chapter_id,
                    values=inventory["unused_definitions"],
                )
            )

    total_references = sum(value["reference_count"] for value in inventories.values())
    total_definitions = sum(value["definition_count"] for value in inventories.values())
    return _result(
        "人工审定章节的正文引注均有尾注对应。"
        if not issues
        else "正文引注与尾注定义未闭环。",
        metrics={
            "reference_count": total_references,
            "definition_count": total_definitions,
            "chapters": inventories,
        },
        issues=issues,
        warnings=warnings,
    )


def _check_content_hygiene(context: _VerificationContext) -> dict[str, Any]:
    issues: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []
    suspicious_english_count = 0
    for item in context.selected:
        chapter_id = str(item.get("id") or "")
        text = context.chapter_texts.get(chapter_id)
        if text is None:
            continue
        path = context.chapter_dir / str(item.get("filename") or "")
        issues.extend(_trace_issues(text, path=path, chapter_id=chapter_id))
        suspicious = _suspicious_english_paragraphs(text)
        suspicious_english_count += len(suspicious)
        if suspicious:
            warnings.append(
                _issue(
                    "long_english_passage",
                    "检测到长篇全英文段落；可能是合法引文或参考文献，建议抽查。",
                    path=path,
                    chapter_id=chapter_id,
                    count=len(suspicious),
                    snippets=suspicious[:4],
                )
            )
    return _result(
        "未发现乱码、模型前言、占位符或来源页码痕迹。"
        if not issues
        else "章节中发现发布污染。",
        metrics={
            "checked_chapter_count": len(context.chapter_texts),
            "suspicious_english_paragraph_count": suspicious_english_count,
        },
        issues=issues,
        warnings=warnings,
    )


def _check_epub(context: _VerificationContext) -> dict[str, Any]:
    issues: list[dict[str, Any]] = []
    chapter_text_match_count = 0
    package_title = ""
    package_language = ""
    path, candidates = _pick_artifact(
        context.output_dir, ".epub", book_title=context.book_title
    )
    if path is None:
        code = "epub_missing" if not candidates else "epub_ambiguous"
        return _result(
            "无法唯一确定 EPUB 成品。",
            issues=[
                _issue(
                    code,
                    "未找到 EPUB。" if not candidates else "存在多个 EPUB，且无法按书名唯一选择。",
                    path=context.output_dir,
                    candidates=[item.name for item in candidates],
                )
            ],
        )
    if len(candidates) != 1:
        issues.append(
            _issue(
                "epub_extra_artifacts",
                "输出目录含当前书名之外的旧 EPUB 成品。",
                path=context.output_dir,
                candidates=[item.name for item in candidates],
            )
        )
    try:
        with zipfile.ZipFile(path) as archive:
            corrupt = archive.testzip()
            if corrupt is not None:
                issues.append(
                    _issue("epub_corrupt_member", "EPUB ZIP 成员校验失败。", path=path, member=corrupt)
                )
            infos = archive.infolist()
            if not infos or infos[0].filename != "mimetype":
                issues.append(
                    _issue("epub_mimetype_not_first", "EPUB 的 mimetype 必须是首个 ZIP 成员。", path=path)
                )
            elif infos[0].compress_type != zipfile.ZIP_STORED:
                issues.append(
                    _issue("epub_mimetype_compressed", "EPUB 的 mimetype 不得压缩。", path=path)
                )
            try:
                mimetype = archive.read("mimetype")
            except KeyError:
                mimetype = b""
            if mimetype != b"application/epub+zip":
                issues.append(
                    _issue("epub_mimetype_invalid", "EPUB mimetype 内容无效。", path=path)
                )

            try:
                container_root = ET.fromstring(archive.read("META-INF/container.xml"))
                rootfile = next(
                    node for node in container_root.iter() if _local_name(node) == "rootfile"
                )
                opf_name = str(rootfile.attrib["full-path"])
                opf_root = ET.fromstring(archive.read(opf_name))
            except (KeyError, StopIteration, ET.ParseError) as exc:
                return _result(
                    "EPUB 包结构无法解析。",
                    issues=issues
                    + [_issue("epub_package_invalid", str(exc), path=path)],
                )

            package_titles = [
                _visible_text(node)
                for node in opf_root.iter()
                if _local_name(node) == "title"
            ]
            package_languages = [
                _visible_text(node)
                for node in opf_root.iter()
                if _local_name(node) == "language"
            ]
            if len(package_titles) != 1 or not package_titles[0]:
                issues.append(
                    _issue(
                        "epub_book_title_invalid",
                        "EPUB package 必须包含唯一且非空的 dc:title。",
                        path=path,
                        actual=package_titles,
                    )
                )
            else:
                package_title = package_titles[0]
                context.artifact_titles["epub"] = package_title
                if context.book_title and package_title != context.book_title:
                    issues.append(
                        _issue(
                            "epub_book_title_mismatch",
                            "EPUB 元数据书名与指定书名不一致。",
                            path=path,
                            expected=context.book_title,
                            actual=package_title,
                        )
                    )
            if len(package_languages) != 1 or not package_languages[0]:
                issues.append(
                    _issue(
                        "epub_language_invalid",
                        "EPUB package 必须包含唯一且非空的 dc:language。",
                        path=path,
                        actual=package_languages,
                    )
                )
            else:
                package_language = package_languages[0]
                if (
                    context.expected_language
                    and package_language != context.expected_language
                ):
                    issues.append(
                        _issue(
                            "epub_language_mismatch",
                            "EPUB 元数据语言与指定目标语言不一致。",
                            path=path,
                            expected=context.expected_language,
                            actual=package_language,
                        )
                    )

            expected_hrefs = [
                Path(str(item.get("filename") or "")).with_suffix(".xhtml").name
                for item in context.manifest
            ]
            manifest_by_id: dict[str, str] = {}
            manifest_ids: list[str] = []
            manifest_chapter_hrefs: list[str] = []
            nav_name: str | None = None
            for node in opf_root.iter():
                if _local_name(node) != "item":
                    continue
                item_id = node.attrib.get("id", "")
                href = node.attrib.get("href", "")
                manifest_ids.append(item_id)
                manifest_by_id[item_id] = href
                properties = node.attrib.get("properties", "").split()
                if "nav" in properties:
                    nav_name = href
                elif node.attrib.get("media-type") == "application/xhtml+xml":
                    manifest_chapter_hrefs.append(href)
            duplicate_manifest_ids = sorted(
                value
                for value, count in Counter(manifest_ids).items()
                if value and count > 1
            )
            if duplicate_manifest_ids:
                issues.append(
                    _issue(
                        "epub_manifest_duplicate_ids",
                        "EPUB manifest item id 不唯一。",
                        path=path,
                        values=duplicate_manifest_ids,
                    )
                )
            if manifest_chapter_hrefs != expected_hrefs:
                issues.append(
                    _issue(
                        "epub_manifest_chapters_mismatch",
                        "EPUB manifest 的章节 XHTML 没有与章节清单一一对应。",
                        path=path,
                        expected=expected_hrefs,
                        actual=manifest_chapter_hrefs,
                    )
                )
            spine_hrefs = [
                manifest_by_id.get(node.attrib.get("idref", ""), "")
                for node in opf_root.iter()
                if _local_name(node) == "itemref"
            ]
            if spine_hrefs != expected_hrefs:
                issues.append(
                    _issue(
                        "epub_spine_mismatch",
                        "EPUB spine 与章节清单不一致。",
                        path=path,
                        expected=expected_hrefs,
                        actual=spine_hrefs,
                    )
                )

            opf_dir = posixpath.dirname(opf_name)
            expected_xhtml_members = {
                posixpath.join(opf_dir, href) if opf_dir else href
                for href in expected_hrefs
            }
            actual_xhtml_members = {
                name
                for name in archive.namelist()
                if posixpath.splitext(name)[1].lower() == ".xhtml"
                and name
                != (
                    posixpath.join(opf_dir, nav_name)
                    if opf_dir and nav_name
                    else (nav_name or "")
                )
            }
            if actual_xhtml_members != expected_xhtml_members:
                issues.append(
                    _issue(
                        "epub_archive_chapters_mismatch",
                        "EPUB ZIP 中存在缺失或未登记的章节 XHTML。",
                        path=path,
                        expected=sorted(expected_xhtml_members),
                        actual=sorted(actual_xhtml_members),
                    )
                )
            nav_links: list[tuple[str, str]] = []
            if nav_name is None:
                issues.append(_issue("epub_nav_missing", "EPUB manifest 缺少导航文档。", path=path))
            else:
                nav_archive_name = (
                    posixpath.join(opf_dir, nav_name) if opf_dir else nav_name
                )
                try:
                    nav_root = ET.fromstring(archive.read(nav_archive_name))
                    nav_links = [
                        (node.attrib.get("href", ""), _visible_text(node))
                        for node in nav_root.iter()
                        if _local_name(node) == "a"
                    ]
                except (KeyError, ET.ParseError) as exc:
                    issues.append(_issue("epub_nav_invalid", str(exc), path=path))
            expected_nav = [
                (href, _display_title(item))
                for href, item in zip(expected_hrefs, context.manifest)
            ]
            if nav_links != expected_nav:
                issues.append(
                    _issue(
                        "epub_nav_mismatch",
                        "EPUB 导航标题或顺序与章节清单不一致。",
                        path=path,
                        expected=expected_nav,
                        actual=nav_links,
                    )
                )

            for href, item in zip(expected_hrefs, context.manifest):
                archive_name = posixpath.join(opf_dir, href) if opf_dir else href
                chapter_id = str(item.get("id") or "")
                try:
                    raw = archive.read(archive_name).decode("utf-8")
                    root = ET.fromstring(raw)
                except (KeyError, UnicodeError, ET.ParseError) as exc:
                    issues.append(
                        _issue(
                            "epub_chapter_invalid",
                            str(exc),
                            path=f"{path}!/{archive_name}",
                            chapter_id=chapter_id,
                        )
                    )
                    continue
                chapter_language = str(
                    root.attrib.get("lang")
                    or root.attrib.get("{http://www.w3.org/XML/1998/namespace}lang")
                    or ""
                )
                if package_language and chapter_language != package_language:
                    issues.append(
                        _issue(
                            "epub_chapter_language_mismatch",
                            "EPUB 章节 XHTML 的 lang 与 package 语言不一致。",
                            path=f"{path}!/{archive_name}",
                            chapter_id=chapter_id,
                            expected=package_language,
                            actual=chapter_language,
                        )
                    )
                h1_titles = [
                    _visible_text(node) for node in root.iter() if _local_name(node) == "h1"
                ]
                expected_title = _display_title(item)
                if h1_titles != [expected_title]:
                    issues.append(
                        _issue(
                            "epub_chapter_title_mismatch",
                            "EPUB 章节 H1 与清单标题不一致。",
                            path=f"{path}!/{archive_name}",
                            chapter_id=chapter_id,
                            expected=[expected_title],
                            actual=h1_titles,
                        )
                    )
                actual_heading_signature = [
                    (int(tag[1]), _visible_text(node))
                    for node in root.iter()
                    if re.fullmatch(
                        r"h[1-6]",
                        (tag := _local_name(node)),
                    )
                ]
                source_markdown = context.chapter_texts.get(chapter_id)
                expected_heading_signature = (
                    _markdown_heading_signature(source_markdown)
                    if source_markdown is not None
                    else []
                )
                if actual_heading_signature != expected_heading_signature:
                    issues.append(
                        _issue(
                            "epub_heading_structure_mismatch",
                            "EPUB 标题层级、文字或顺序与当前 Markdown 不一致。",
                            path=f"{path}!/{archive_name}",
                            chapter_id=chapter_id,
                            expected=expected_heading_signature,
                            actual=actual_heading_signature,
                        )
                    )
                body = next(
                    (node for node in root.iter() if _local_name(node) == "body"),
                    None,
                )
                if body is None:
                    issues.append(
                        _issue(
                            "epub_chapter_body_missing",
                            "EPUB 章节 XHTML 缺少 body。",
                            path=f"{path}!/{archive_name}",
                            chapter_id=chapter_id,
                        )
                    )
                elif source_markdown is not None:
                    expected_text = _canonical_visible_text(
                        _markdown_visible_text(source_markdown)
                    )
                    actual_text = _canonical_visible_text(
                        "".join(body.itertext())
                    )
                    if actual_text != expected_text:
                        issues.append(
                            _issue(
                                "epub_chapter_text_mismatch",
                                "EPUB 章节正文与当前 Markdown 不一致。",
                                path=f"{path}!/{archive_name}",
                                chapter_id=chapter_id,
                                expected_sha256=_sha256_text(expected_text),
                                actual_sha256=_sha256_text(actual_text),
                                expected_characters=len(expected_text),
                                actual_characters=len(actual_text),
                            )
                        )
                    else:
                        chapter_text_match_count += 1
                issues.extend(
                    _trace_issues(
                        raw,
                        path=f"{path}!/{archive_name}",
                        chapter_id=chapter_id,
                    )
                )
    except (OSError, zipfile.BadZipFile, RuntimeError) as exc:
        issues.append(_issue("epub_unreadable", str(exc), path=path))

    return _result(
        "EPUB 包、章节、spine 与导航完整。" if not issues else "EPUB 发布结构不完整。",
        metrics={
            "path": str(path),
            "book_title": package_title,
            "language": package_language,
            "chapter_count": len(context.manifest),
            "chapter_text_match_count": chapter_text_match_count,
        },
        issues=issues,
    )


def _iter_docx_paragraphs(document: Any) -> Iterable[Any]:
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def _docx_document_payload(
    document: Any,
) -> tuple[list[tuple[str, str]], list[dict[str, Any]]]:
    """Return pre-chapter content and chapter structure in document order."""

    from docx.table import Table  # type: ignore[import-not-found]
    from docx.text.paragraph import Paragraph  # type: ignore[import-not-found]

    preamble: list[tuple[str, str]] = []
    chapters: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None

    def new_chapter() -> dict[str, Any]:
        return {
            "text": [],
            "headings": [],
            "quotes": [],
            "inline_styles": {"bold": [], "italic": [], "underline": []},
        }

    def add_paragraph(payload: dict[str, Any], paragraph: Any) -> None:
        style_name = str(getattr(paragraph.style, "name", "") or "")
        payload["text"].append(paragraph.text)
        heading = re.fullmatch(r"Heading ([1-3])", style_name)
        if heading is not None:
            payload["headings"].append(
                (int(heading.group(1)), paragraph.text.strip())
            )
        if style_name == "Quote" and paragraph.text.strip():
            payload["quotes"].append(
                _canonical_visible_text(paragraph.text)
            )
        for run in paragraph.runs:
            value = _canonical_visible_text(run.text)
            if not value:
                continue
            if run.bold is True:
                payload["inline_styles"]["bold"].append(value)
            if run.italic is True:
                payload["inline_styles"]["italic"].append(value)
            if bool(run.underline):
                payload["inline_styles"]["underline"].append(value)

    for child in document.element.body.iterchildren():
        tag = _local_name(child)
        if tag == "p":
            paragraph = Paragraph(child, document)
            style_name = getattr(paragraph.style, "name", "")
            if style_name == "Heading 1":
                if current is not None:
                    current["text"] = "\n".join(current["text"])
                    chapters.append(current)
                current = new_chapter()
                add_paragraph(current, paragraph)
            elif current is not None:
                add_paragraph(current, paragraph)
            elif paragraph.text.strip():
                preamble.append((str(style_name or ""), paragraph.text.strip()))
        elif tag == "tbl":
            table = Table(child, document)
            if current is None:
                table_text = "\n".join(
                    paragraph.text
                    for row in table.rows
                    for cell in row.cells
                    for paragraph in cell.paragraphs
                    if paragraph.text.strip()
                ).strip()
                if table_text:
                    preamble.append(("Table", table_text))
            else:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            add_paragraph(current, paragraph)
    if current is not None:
        current["text"] = "\n".join(current["text"])
        chapters.append(current)
    return preamble, chapters


def _docx_paragraph_style_id(paragraph: ET.Element) -> str:
    for child in paragraph.iter():
        if _local_name(child) == "pstyle":
            return str(_xml_attribute(child, "val") or "")
    return ""


def _docx_active_on_off(element: ET.Element) -> bool:
    value = str(_xml_attribute(element, "val") or "").strip().casefold()
    return value not in {"0", "false", "off"}


def _check_docx_page_break_contract(
    document_root: ET.Element,
    path: Path,
) -> tuple[list[dict[str, Any]], dict[str, int]]:
    """Allow only deliberate book-layout page breaks.

    A book may start every chapter on a new page and may end its title block
    with one explicit break.  Other authored breaks remain suspicious because
    they can hide blank/source pages.  Render verification remains responsible
    for detecting the resulting visual blank-page and pagination anomalies.
    """

    issues: list[dict[str, Any]] = []
    body = next(
        (node for node in document_root.iter() if _local_name(node) == "body"),
        None,
    )
    if body is None:
        return issues, {
            "chapter_page_break_before_count": 0,
            "title_page_break_count": 0,
            "unexpected_page_break_count": 0,
            "last_rendered_page_break_count": 0,
        }

    body_children = list(body)
    first_heading_index: int | None = None
    for child_index, child in enumerate(body_children):
        if _local_name(child) != "p":
            continue
        style_id = re.sub(r"\s+", "", _docx_paragraph_style_id(child)).casefold()
        if style_id == "heading1":
            first_heading_index = child_index
            break

    chapter_breaks = 0
    title_breaks = 0
    unexpected_breaks = 0
    rendered_breaks = 0
    for child_index, paragraph in enumerate(body_children):
        if _local_name(paragraph) != "p":
            continue
        style_id = _docx_paragraph_style_id(paragraph)
        normalized_style = re.sub(r"\s+", "", style_id).casefold()
        paragraph_text = "".join(
            str(node.text or "")
            for node in paragraph.iter()
            if _local_name(node) == "t"
        ).strip()
        page_break_before = [
            node
            for node in paragraph.iter()
            if _local_name(node) == "pagebreakbefore"
            and _docx_active_on_off(node)
        ]
        explicit_breaks = [
            node
            for node in paragraph.iter()
            if _local_name(node) == "br"
            and str(_xml_attribute(node, "type") or "").casefold() == "page"
        ]
        last_rendered = [
            node
            for node in paragraph.iter()
            if _local_name(node) == "lastrenderedpagebreak"
        ]

        if page_break_before:
            if normalized_style == "heading1" and len(page_break_before) == 1:
                chapter_breaks += 1
            else:
                unexpected_breaks += len(page_break_before)
                issues.append(
                    _issue(
                        "docx_page_break_present",
                        "Word 成品含未登记的段前分页；仅一级章节标题可启用 pageBreakBefore。",
                        path=f"{path}!/word/document.xml",
                        paragraph_index=child_index,
                        paragraph_style=style_id,
                        text=paragraph_text[:160],
                    )
                )

        if explicit_breaks:
            title_page_break = bool(
                len(explicit_breaks) == 1
                and title_breaks == 0
                and first_heading_index is not None
                and child_index == first_heading_index - 1
                and not paragraph_text
            )
            if title_page_break:
                title_breaks += 1
            else:
                unexpected_breaks += len(explicit_breaks)
                issues.append(
                    _issue(
                        "docx_page_break_present",
                        "Word 成品含未登记的显式分页；仅扉页与首章之间允许一个空分页段落。",
                        path=f"{path}!/word/document.xml",
                        paragraph_index=child_index,
                        paragraph_style=style_id,
                        text=paragraph_text[:160],
                        break_count=len(explicit_breaks),
                    )
                )

        if last_rendered:
            rendered_breaks += len(last_rendered)
            unexpected_breaks += len(last_rendered)
            issues.append(
                _issue(
                    "docx_page_break_present",
                    "Word 成品含持久化的 lastRenderedPageBreak，无法视为受控章节分页。",
                    path=f"{path}!/word/document.xml",
                    paragraph_index=child_index,
                    text=paragraph_text[:160],
                    break_count=len(last_rendered),
                )
            )

    return issues, {
        "chapter_page_break_before_count": chapter_breaks,
        "title_page_break_count": title_breaks,
        "unexpected_page_break_count": unexpected_breaks,
        "last_rendered_page_break_count": rendered_breaks,
    }


def _docx_page_field_only_footer(
    root: ET.Element,
) -> bool:
    """Return true only for a footer containing one PAGE field and its cache."""

    if _local_name(root) != "ftr":
        return False
    instructions: list[str] = []
    for node in root.iter():
        local_name = _local_name(node)
        if local_name == "fldsimple":
            instruction = str(_xml_attribute(node, "instr") or "").strip()
            if instruction:
                instructions.append(instruction)
        elif local_name == "instrtext":
            instruction = str(node.text or "").strip()
            if instruction:
                instructions.append(instruction)
    if len(instructions) != 1:
        return False
    if re.fullmatch(
        r"PAGE(?:\s+\\\*\s+(?:MERGEFORMAT|CHARFORMAT|ARABIC|ROMAN))*",
        instructions[0],
        flags=re.I,
    ) is None:
        return False
    visible = _canonical_visible_text(
        "".join(
            str(node.text or "")
            for node in root.iter()
            if _local_name(node) == "t"
        )
    )
    return not visible or re.fullmatch(r"[0-9\uff10-\uff19]+", visible) is not None


def _xml_attribute(element: ET.Element, name: str) -> str | None:
    """Return an XML attribute by local name across strict/transitional OOXML."""

    wanted = name.lower()
    for key, value in element.attrib.items():
        local = key.rsplit("}", 1)[-1].lower() if isinstance(key, str) else ""
        if local == wanted:
            return str(value)
    return None


def _docx_relationship_target(target: str) -> str:
    value = target.replace("\\", "/")
    if value.startswith("/"):
        return posixpath.normpath(value.lstrip("/"))
    return posixpath.normpath(posixpath.join("word", value))


def _docx_positive_footnote_texts(
    archive: zipfile.ZipFile,
    document_root: ET.Element | None = None,
) -> list[str]:
    """Return positive footnote text in body-reference order.

    Word footnote IDs are relationship keys, not a promise that references
    appear in numeric order.  A paragraph can legitimately cite IDs ``2, 1``;
    the semantic Markdown contract follows that authored landing order.
    """

    member = "word/footnotes.xml"
    if member not in archive.namelist():
        return []
    root = ET.fromstring(archive.read(member))
    definitions: dict[int, str] = {}
    for note in root:
        if _local_name(note) != "footnote":
            continue
        raw_id = _xml_attribute(note, "id")
        try:
            note_id = int(raw_id) if raw_id is not None else 0
        except ValueError:
            continue
        if note_id <= 0:
            continue
        visible = "".join(
            str(node.text or "")
            for node in note.iter()
            if _local_name(node) == "t"
        )
        definitions[note_id] = _canonical_visible_text(visible)

    if document_root is None:
        try:
            document_root = ET.fromstring(archive.read("word/document.xml"))
        except (KeyError, ET.ParseError):
            return [definitions[note_id] for note_id in sorted(definitions)]
    reference_ids: list[int] = []
    for node in document_root.iter():
        if _local_name(node) != "footnotereference":
            continue
        raw_id = _xml_attribute(node, "id")
        try:
            note_id = int(raw_id) if raw_id is not None else 0
        except ValueError:
            continue
        if note_id > 0:
            reference_ids.append(note_id)
    return [definitions[note_id] for note_id in reference_ids if note_id in definitions]


def _check_docx_note_contract(
    archive: zipfile.ZipFile,
    path: Path,
    document_root: ET.Element | None,
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """Validate the package contract for true Word footnotes.

    True footnotes are supported; endnotes are deliberately outside the current
    publication contract.  The checks use local XML names so both strict and
    transitional OOXML namespaces remain parseable.
    """

    issues: list[dict[str, Any]] = []
    members = set(archive.namelist())
    footnotes_member = "word/footnotes.xml"
    footnotes_present = footnotes_member in members

    reference_ids: list[int] = []
    footnote_reference_node_count = 0
    endnote_reference_count = 0
    if document_root is not None:
        for node in document_root.iter():
            local_name = _local_name(node)
            if local_name == "endnotereference":
                endnote_reference_count += 1
                continue
            if local_name != "footnotereference":
                continue
            footnote_reference_node_count += 1
            raw_id = _xml_attribute(node, "id")
            try:
                note_id = int(raw_id) if raw_id is not None else None
            except ValueError:
                note_id = None
            if note_id is None or note_id < 1:
                issues.append(
                    _issue(
                        "docx_footnote_reference_id_invalid",
                        "Word 脚注引用 ID 必须是正整数。",
                        path=f"{path}!/word/document.xml",
                        actual=raw_id,
                    )
                )
                if note_id in {-1, 0}:
                    issues.append(
                        _issue(
                            "docx_footnote_reserved_reference",
                            "Word 正文不得引用脚注保留节点。",
                            path=f"{path}!/word/document.xml",
                            footnote_id=note_id,
                        )
                    )
                continue
            reference_ids.append(note_id)

    relationship_nodes: list[ET.Element] = []
    endnote_relationship_count = 0
    relationships_error: str | None = None
    relationships_member = "word/_rels/document.xml.rels"
    if relationships_member in members:
        try:
            relationships_root = ET.fromstring(archive.read(relationships_member))
            for node in relationships_root.iter():
                if _local_name(node) != "relationship":
                    continue
                relationship_type = str(_xml_attribute(node, "type") or "")
                relation_kind = (
                    relationship_type.rstrip("/").rsplit("/", 1)[-1].lower()
                )
                if relation_kind == "footnotes":
                    relationship_nodes.append(node)
                elif relation_kind == "endnotes":
                    endnote_relationship_count += 1
        except (KeyError, ET.ParseError) as exc:
            relationships_error = str(exc)

    footnote_overrides: list[ET.Element] = []
    endnote_content_type_count = 0
    content_types_error: str | None = None
    if "[Content_Types].xml" in members:
        try:
            content_types_root = ET.fromstring(archive.read("[Content_Types].xml"))
            for node in content_types_root.iter():
                if _local_name(node) != "override":
                    continue
                part_name = str(_xml_attribute(node, "partname") or "")
                content_type = str(_xml_attribute(node, "contenttype") or "")
                normalized_part = part_name.lstrip("/").lower()
                if normalized_part == footnotes_member:
                    footnote_overrides.append(node)
                if (
                    normalized_part == "word/endnotes.xml"
                    or content_type.lower().endswith(".endnotes+xml")
                ):
                    endnote_content_type_count += 1
        except (KeyError, ET.ParseError) as exc:
            content_types_error = str(exc)

    footnote_signal = bool(
        footnote_reference_node_count
        or footnotes_present
        or relationship_nodes
        or footnote_overrides
    )
    relationship_valid = False
    if footnote_signal:
        if relationships_error is not None:
            issues.append(
                _issue(
                    "docx_footnotes_relationship_invalid",
                    "Word 主文档关系文件无法解析。",
                    path=f"{path}!/{relationships_member}",
                    error=relationships_error,
                )
            )
        else:
            valid_relationships = [
                node
                for node in relationship_nodes
                if _docx_relationship_target(
                    str(_xml_attribute(node, "target") or "")
                )
                == footnotes_member
                and str(_xml_attribute(node, "targetmode") or "").lower()
                != "external"
            ]
            relationship_valid = (
                len(relationship_nodes) == 1 and len(valid_relationships) == 1
            )
            if not relationship_valid:
                issues.append(
                    _issue(
                        "docx_footnotes_relationship_invalid",
                        "Word 脚注必须由主文档通过唯一内部关系指向 footnotes.xml。",
                        path=f"{path}!/{relationships_member}",
                        relationship_count=len(relationship_nodes),
                        targets=[
                            _xml_attribute(node, "target")
                            for node in relationship_nodes
                        ],
                    )
                )

    content_type_valid = False
    if footnote_signal:
        if content_types_error is not None:
            issues.append(
                _issue(
                    "docx_footnotes_content_type_invalid",
                    "Word 内容类型清单无法解析。",
                    path=f"{path}!/[Content_Types].xml",
                    error=content_types_error,
                )
            )
        else:
            valid_overrides = [
                node
                for node in footnote_overrides
                if str(_xml_attribute(node, "contenttype") or "")
                .lower()
                .endswith(".footnotes+xml")
            ]
            content_type_valid = (
                len(footnote_overrides) == 1 and len(valid_overrides) == 1
            )
            if not content_type_valid:
                issues.append(
                    _issue(
                        "docx_footnotes_content_type_invalid",
                        "Word 脚注部件必须登记唯一且正确的内容类型。",
                        path=f"{path}!/[Content_Types].xml",
                        override_count=len(footnote_overrides),
                        content_types=[
                            _xml_attribute(node, "contenttype")
                            for node in footnote_overrides
                        ],
                    )
                )

    definition_ids: list[int] = []
    reserved_counts: Counter[int] = Counter()
    footnotes_root: ET.Element | None = None
    if footnote_signal and not footnotes_present:
        issues.append(
            _issue(
                "docx_footnotes_part_missing",
                "Word 脚注契约已启用，但包内缺少 footnotes.xml。",
                path=path,
            )
        )
    elif footnotes_present:
        try:
            footnotes_root = ET.fromstring(archive.read(footnotes_member))
            if _local_name(footnotes_root) != "footnotes":
                raise ET.ParseError("root element is not w:footnotes")
        except (KeyError, ET.ParseError) as exc:
            issues.append(
                _issue(
                    "docx_footnotes_part_invalid",
                    "Word 脚注部件无法解析。",
                    path=f"{path}!/{footnotes_member}",
                    error=str(exc),
                )
            )
            footnotes_root = None

    if footnotes_root is not None:
        for node in footnotes_root:
            if _local_name(node) != "footnote":
                continue
            raw_id = _xml_attribute(node, "id")
            try:
                note_id = int(raw_id) if raw_id is not None else None
            except ValueError:
                note_id = None
            if note_id is None or note_id < -1:
                issues.append(
                    _issue(
                        "docx_footnote_definition_id_invalid",
                        "Word 脚注定义 ID 必须为正整数或保留值 -1/0。",
                        path=f"{path}!/{footnotes_member}",
                        actual=raw_id,
                    )
                )
                continue
            if note_id in {-1, 0}:
                reserved_counts[note_id] += 1
                expected_type = (
                    "separator" if note_id == -1 else "continuationSeparator"
                )
                actual_type = _xml_attribute(node, "type")
                if actual_type != expected_type:
                    issues.append(
                        _issue(
                            "docx_footnote_reserved_type_invalid",
                            "Word 脚注保留节点缺少或使用了错误的 w:type。",
                            path=f"{path}!/{footnotes_member}",
                            footnote_id=note_id,
                            expected=expected_type,
                            actual=actual_type,
                        )
                    )
                continue
            definition_ids.append(note_id)

        for reserved_id in (-1, 0):
            if reserved_counts[reserved_id] == 0:
                issues.append(
                    _issue(
                        "docx_footnote_reserved_node_missing",
                        "Word 脚注部件缺少必需的保留节点。",
                        path=f"{path}!/{footnotes_member}",
                        footnote_id=reserved_id,
                    )
                )

        visible_notes = "".join(
            str(node.text or "")
            for node in footnotes_root.iter()
            if _local_name(node) == "t"
        )
        issues.extend(
            _trace_issues(visible_notes, path=f"{path}!/{footnotes_member}")
        )

    reference_counts = Counter(reference_ids)
    definition_counts = Counter(definition_ids)
    duplicate_references = sorted(
        note_id for note_id, count in reference_counts.items() if count != 1
    )
    duplicate_definitions = sorted(
        note_id for note_id, count in definition_counts.items() if count != 1
    )
    if duplicate_references:
        issues.append(
            _issue(
                "docx_footnote_reference_duplicate",
                "Word 脚注 ID 在正文中被引用多次。",
                path=f"{path}!/word/document.xml",
                footnote_ids=duplicate_references,
                counts={
                    str(note_id): reference_counts[note_id]
                    for note_id in duplicate_references
                },
            )
        )
    if duplicate_definitions:
        issues.append(
            _issue(
                "docx_footnote_definition_duplicate",
                "Word 脚注部件含重复的脚注定义 ID。",
                path=f"{path}!/{footnotes_member}",
                footnote_ids=duplicate_definitions,
                counts={
                    str(note_id): definition_counts[note_id]
                    for note_id in duplicate_definitions
                },
            )
        )
    duplicate_reserved = sorted(
        note_id for note_id, count in reserved_counts.items() if count != 1
    )
    if duplicate_reserved:
        issues.append(
            _issue(
                "docx_footnote_definition_duplicate",
                "Word 脚注部件含重复的保留节点 ID。",
                path=f"{path}!/{footnotes_member}",
                footnote_ids=duplicate_reserved,
                counts={
                    str(note_id): reserved_counts[note_id]
                    for note_id in duplicate_reserved
                },
            )
        )

    orphan_references = sorted(set(reference_ids) - set(definition_ids))
    orphan_definitions = sorted(set(definition_ids) - set(reference_ids))
    if orphan_references:
        issues.append(
            _issue(
                "docx_footnote_orphan_reference",
                "Word 正文含无对应定义的脚注引用。",
                path=f"{path}!/word/document.xml",
                footnote_ids=orphan_references,
            )
        )
    if orphan_definitions:
        issues.append(
            _issue(
                "docx_footnote_orphan_definition",
                "Word 脚注部件含未被正文引用的定义。",
                path=f"{path}!/{footnotes_member}",
                footnote_ids=orphan_definitions,
            )
        )

    endnotes_present = "word/endnotes.xml" in members
    if (
        endnote_reference_count
        or endnotes_present
        or endnote_relationship_count
        or endnote_content_type_count
    ):
        issues.append(
            _issue(
                "docx_unexpected_endnotes",
                "当前 Word 发布契约不支持尾注。",
                path=path,
                reference_count=endnote_reference_count,
                part_present=endnotes_present,
                relationship_count=endnote_relationship_count,
                content_type_count=endnote_content_type_count,
            )
        )

    metrics = {
        "footnotes_part_present": footnotes_present,
        "footnotes_relationship_valid": relationship_valid,
        "footnotes_content_type_valid": content_type_valid,
        "footnote_reference_count": footnote_reference_node_count,
        "footnote_unique_reference_count": len(reference_counts),
        "footnote_definition_count": len(definition_ids),
        "footnote_unique_definition_count": len(definition_counts),
        "footnote_reserved_node_count": sum(reserved_counts.values()),
        "footnote_orphan_reference_count": len(orphan_references),
        "footnote_orphan_definition_count": len(orphan_definitions),
        "endnotes_present": endnotes_present,
    }
    return issues, metrics


def _check_docx(context: _VerificationContext) -> dict[str, Any]:
    issues: list[dict[str, Any]] = []
    chapter_text_match_count = 0
    footnote_text_match_count = 0
    expected_footnote_text_count = 0
    note_metrics: dict[str, Any] = {
        "footnotes_part_present": False,
        "footnotes_relationship_valid": False,
        "footnotes_content_type_valid": False,
        "footnote_reference_count": 0,
        "footnote_unique_reference_count": 0,
        "footnote_definition_count": 0,
        "footnote_unique_definition_count": 0,
        "footnote_reserved_node_count": 0,
        "footnote_orphan_reference_count": 0,
        "footnote_orphan_definition_count": 0,
        "endnotes_present": False,
    }
    page_break_metrics: dict[str, int] = {
        "chapter_page_break_before_count": 0,
        "title_page_break_count": 0,
        "unexpected_page_break_count": 0,
        "last_rendered_page_break_count": 0,
    }
    page_number_footer_count = 0
    path, candidates = _pick_artifact(
        context.output_dir, ".docx", book_title=context.book_title
    )
    if path is None:
        code = "docx_missing" if not candidates else "docx_ambiguous"
        return _result(
            "无法唯一确定 Word 成品。",
            issues=[
                _issue(
                    code,
                    "未找到 DOCX。" if not candidates else "存在多个 DOCX，且无法按书名唯一选择。",
                    path=context.output_dir,
                    candidates=[item.name for item in candidates],
                )
            ],
        )
    if len(candidates) != 1:
        issues.append(
            _issue(
                "docx_extra_artifacts",
                "输出目录含当前书名之外的旧 Word 成品。",
                path=context.output_dir,
                candidates=[item.name for item in candidates],
            )
        )

    try:
        with zipfile.ZipFile(path) as archive:
            corrupt = archive.testzip()
            if corrupt is not None:
                issues.append(
                    _issue("docx_corrupt_member", "DOCX ZIP 成员校验失败。", path=path, member=corrupt)
                )
            document_xml = archive.read("word/document.xml").decode("utf-8")
            try:
                document_root = ET.fromstring(document_xml)
            except ET.ParseError as exc:
                issues.append(_issue("docx_xml_invalid", str(exc), path=path))
                document_root = None
            if document_root is not None:
                page_break_issues, page_break_metrics = (
                    _check_docx_page_break_contract(document_root, path)
                )
                issues.extend(page_break_issues)
                for node in document_root.iter():
                    if _local_name(node) != "txbxcontent":
                        continue
                    textbox_text = "".join(
                        str(child.text or "")
                        for child in node.iter()
                        if _local_name(child) == "t"
                    ).strip()
                    if textbox_text:
                        issues.append(
                            _issue(
                                "docx_unexpected_textbox",
                                "Word 成品含生成器未登记的文本框文字。",
                                path=path,
                                text=textbox_text[:300],
                            )
                        )
                        issues.extend(
                            _trace_issues(textbox_text, path=f"{path}!/word/document.xml")
                        )
            note_issues, note_metrics = _check_docx_note_contract(
                archive, path, document_root
            )
            issues.extend(note_issues)
            expected_footnotes: list[str] = []
            for item in context.manifest:
                chapter_id = str(item.get("id") or "")
                source_markdown = context.chapter_texts.get(chapter_id)
                if source_markdown is not None:
                    expected_footnotes.extend(
                        _expected_docx_footnote_texts(source_markdown)
                    )
            expected_footnote_text_count = len(expected_footnotes)
            try:
                actual_footnotes = _docx_positive_footnote_texts(
                    archive,
                    document_root,
                )
            except (KeyError, ET.ParseError, UnicodeError, ValueError):
                actual_footnotes = []
            if actual_footnotes != expected_footnotes:
                issues.append(
                    _issue(
                        "docx_footnote_text_mismatch",
                        "Word 脚注正文与当前 Markdown 语义脚注不一致。",
                        path=f"{path}!/word/footnotes.xml",
                        expected_count=len(expected_footnotes),
                        actual_count=len(actual_footnotes),
                        expected_sha256=_sha256_text("\n".join(expected_footnotes)),
                        actual_sha256=_sha256_text("\n".join(actual_footnotes)),
                    )
                )
            else:
                footnote_text_match_count = len(expected_footnotes)
            auxiliary_pattern = re.compile(
                r"word/(?:header\d*|footer\d*)\.xml$",
                flags=re.I,
            )
            for member in archive.namelist():
                if auxiliary_pattern.fullmatch(member) is None:
                    continue
                try:
                    raw_part = archive.read(member).decode("utf-8")
                    part_root = ET.fromstring(raw_part)
                except (KeyError, UnicodeError, ET.ParseError) as exc:
                    issues.append(
                        _issue(
                            "docx_auxiliary_part_invalid",
                            str(exc),
                            path=f"{path}!/{member}",
                        )
                    )
                    continue
                visible_part = "".join(
                    str(node.text or "")
                    for node in part_root.iter()
                    if _local_name(node) == "t"
                ).strip()
                if _docx_page_field_only_footer(part_root):
                    page_number_footer_count += 1
                    continue
                if visible_part or re.search(r"\b(?:PAGE|NUMPAGES)\b", raw_part, re.I):
                    issues.append(
                        _issue(
                            "docx_unexpected_header_footer_or_notes",
                            "Word 成品含生成器未登记的页眉或页脚内容。",
                            path=f"{path}!/{member}",
                            text=visible_part[:300],
                        )
                    )
                    issues.extend(
                        _trace_issues(visible_part, path=f"{path}!/{member}")
                    )
    except (OSError, KeyError, UnicodeError, zipfile.BadZipFile) as exc:
        return _result(
            "Word 包无法读取。",
            issues=issues + [_issue("docx_unreadable", str(exc), path=path)],
        )

    try:
        from docx import Document  # type: ignore[import-not-found]

        document = Document(path)
        paragraphs = list(_iter_docx_paragraphs(document))
        chapter_headings = [
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if getattr(paragraph.style, "name", "") == "Heading 1"
        ]
        expected_headings = [_display_title(item) for item in context.manifest]
        if chapter_headings != expected_headings:
            issues.append(
                _issue(
                    "docx_chapter_headings_mismatch",
                    "Word 一级标题与章节清单不一致。",
                    path=path,
                    expected=expected_headings,
                    actual=chapter_headings,
                )
            )
        title_paragraphs = [
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if getattr(paragraph.style, "name", "") in DOCX_BOOK_TITLE_STYLES
        ]
        core_title = str(document.core_properties.title or "").strip()
        inferred_title = (
            context.book_title
            or context.artifact_titles.get("epub")
            or (title_paragraphs[0] if len(title_paragraphs) == 1 else "")
            or core_title
        )
        if len(title_paragraphs) != 1 or not title_paragraphs[0]:
            issues.append(
                _issue(
                    "docx_book_title_invalid",
                    "Word 必须包含唯一且非空的 Title 或 Codex Book Title 段落。",
                    path=path,
                    actual=title_paragraphs,
                )
            )
        elif inferred_title and title_paragraphs != [inferred_title]:
            issues.append(
                _issue(
                    "docx_book_title_mismatch",
                    "Word 文档标题与期望书名不一致。",
                    path=path,
                    expected=[inferred_title],
                    actual=title_paragraphs,
                )
            )
        if not core_title:
            issues.append(
                _issue(
                    "docx_core_title_invalid",
                    "Word 核心属性中的书名为空。",
                    path=path,
                )
            )
        elif inferred_title and core_title != inferred_title:
            issues.append(
                _issue(
                    "docx_core_title_mismatch",
                    "Word 核心属性中的书名与期望书名不一致。",
                    path=path,
                    expected=inferred_title,
                    actual=core_title,
                )
            )
        if core_title:
            context.artifact_titles["docx"] = core_title

        preamble, actual_chapters = _docx_document_payload(document)
        front_paragraphs: list[Any] = []
        for paragraph in document.paragraphs:
            if getattr(paragraph.style, "name", "") == "Heading 1":
                break
            if paragraph.text.strip():
                front_paragraphs.append(paragraph)
        title_style = (
            str(getattr(front_paragraphs[0].style, "name", "") or "")
            if front_paragraphs
            and getattr(front_paragraphs[0].style, "name", "")
            in DOCX_BOOK_TITLE_STYLES
            else "Title"
        )
        expected_preamble = (
            [(title_style, inferred_title)] if inferred_title else []
        )
        core_author = str(document.core_properties.author or "").strip()
        if len(front_paragraphs) >= 2:
            from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore[import-not-found]

            author_paragraph = front_paragraphs[1]
            if (
                core_author
                and author_paragraph.text.strip() == core_author
                and getattr(author_paragraph.style, "name", "") == "Normal"
                and author_paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER
            ):
                expected_preamble.append(("Normal", core_author))
        if preamble != expected_preamble:
            issues.append(
                _issue(
                    "docx_front_matter_mismatch",
                    "Word 书名与首章之间含有未登记的正文或表格。",
                    path=path,
                    expected=expected_preamble,
                    actual=preamble,
                )
            )

        for item, payload in zip(context.manifest, actual_chapters):
            actual_text = str(payload["text"])
            chapter_id = str(item.get("id") or "")
            source_markdown = context.chapter_texts.get(chapter_id)
            if source_markdown is None:
                continue
            # Mirror the Word publisher's soft-wrap merge so the expected
            # text is derived from exactly what build_docx renders.
            from book_pipeline import _normalize_wrapped_markdown_for_docx

            expected_text = _canonical_visible_text(
                _markdown_visible_text(
                    _normalize_wrapped_markdown_for_docx(
                        _docx_markdown_body(source_markdown)
                    )
                )
            )
            canonical_actual = _canonical_visible_text(actual_text)
            if canonical_actual != expected_text:
                issues.append(
                    _issue(
                        "docx_chapter_text_mismatch",
                        "Word 章节正文与当前 Markdown 不一致。",
                        path=path,
                        chapter_id=chapter_id,
                        expected_sha256=_sha256_text(expected_text),
                        actual_sha256=_sha256_text(canonical_actual),
                        expected_characters=len(expected_text),
                        actual_characters=len(canonical_actual),
                    )
                )
            else:
                chapter_text_match_count += 1

        expectations = _docx_expectations(context.manifest, context.chapter_texts)
        actual_table_shapes = [
            [len(row.cells) for row in table.rows] for table in document.tables
        ]
        if actual_table_shapes != expectations["table_shapes"]:
            issues.append(
                _issue(
                    "docx_tables_mismatch",
                    "Word 表格数量或行列结构与审定 Markdown 不一致。",
                    path=path,
                    expected=expectations["table_shapes"],
                    actual=actual_table_shapes,
                )
            )
        actual_quotes = sum(
            1
            for paragraph in document.paragraphs
            if getattr(paragraph.style, "name", "") == "Quote"
        )
        if actual_quotes != expectations["quote_paragraphs"]:
            issues.append(
                _issue(
                    "docx_quotes_mismatch",
                    "Word 引文段落数量与审定 Markdown 不一致。",
                    path=path,
                    expected=expectations["quote_paragraphs"],
                    actual=actual_quotes,
                )
            )
        for item, payload in zip(context.manifest, actual_chapters):
            chapter_id = str(item.get("id") or "")
            expected_quotes = expectations["quotes"].get(chapter_id, [])
            if payload["quotes"] != expected_quotes:
                issues.append(
                    _issue(
                        "docx_quote_structure_mismatch",
                        "Word 引文段落的文字、归属或顺序与 Markdown 不一致。",
                        path=path,
                        chapter_id=chapter_id,
                        expected=expected_quotes[:20],
                        actual=payload["quotes"][:20],
                    )
                )
            expected_headings = expectations["headings"].get(chapter_id, [])
            if payload["headings"] != expected_headings:
                issues.append(
                    _issue(
                        "docx_heading_structure_mismatch",
                        "Word 标题层级、文字或顺序与当前 Markdown 不一致。",
                        path=path,
                        chapter_id=chapter_id,
                        expected=expected_headings,
                        actual=payload["headings"],
                    )
                )
            expected_styles = expectations["inline_styles"].get(chapter_id, {})
            for style_name in ("bold", "italic", "underline"):
                expected_fragments = expected_styles.get(style_name, [])
                actual_fragments = payload["inline_styles"].get(style_name, [])
                if actual_fragments != expected_fragments:
                    issues.append(
                        _issue(
                            "docx_inline_style_mismatch",
                            "Word 未精确保留 Markdown 的行内样式。",
                            path=path,
                            chapter_id=chapter_id,
                            style=style_name,
                            expected=expected_fragments[:20],
                            actual=actual_fragments[:20],
                        )
                    )

        visible = "\n".join(paragraph.text for paragraph in paragraphs)
        issues.extend(_trace_issues(visible, path=path))
    except ImportError as exc:
        issues.append(_issue("docx_dependency_missing", str(exc), path=path))
        expectations = {
            "table_shapes": [],
            "quote_paragraphs": 0,
            "headings": {},
            "inline_styles": {},
            "quotes": {},
        }
    except Exception as exc:  # python-docx emits several package/XML exceptions.
        issues.append(_issue("docx_parse_failed", f"{type(exc).__name__}: {exc}", path=path))
        expectations = {
            "table_shapes": [],
            "quote_paragraphs": 0,
            "headings": {},
            "inline_styles": {},
            "quotes": {},
        }

    return _result(
        "Word 书名、标题、表格、引文、行内样式及脚注包契约完整。"
        if not issues
        else "Word 发布结构不完整。",
        metrics={
            "path": str(path),
            "chapter_count": len(context.manifest),
            "chapter_text_match_count": chapter_text_match_count,
            "expected_footnote_text_count": expected_footnote_text_count,
            "footnote_text_match_count": footnote_text_match_count,
            "expected_table_count": len(expectations["table_shapes"]),
            "expected_quote_paragraph_count": expectations["quote_paragraphs"],
            "expected_inline_style_fragment_count": sum(
                len(values)
                for chapter in expectations["inline_styles"].values()
                for values in chapter.values()
            ),
            **page_break_metrics,
            "page_number_footer_count": page_number_footer_count,
            **note_metrics,
        },
        issues=issues,
    )


def _check_docx_render(context: _VerificationContext) -> dict[str, Any]:
    path, candidates = _pick_artifact(
        context.output_dir,
        ".docx",
        book_title=context.book_title,
    )
    if path is None:
        return _result(
            "无法唯一确定待渲染的 Word 成品。",
            issues=[
                _issue(
                    "docx_render_artifact_missing_or_ambiguous",
                    "渲染门要求唯一 Word 成品。",
                    path=context.output_dir,
                    candidates=[item.name for item in candidates],
                )
            ],
        )
    from docx_render_gate import verify_docx_render

    result = verify_docx_render(path)
    issues = list(result.get("issues") or [])
    warnings = list(result.get("warnings") or [])
    for issue in (*issues, *warnings):
        if isinstance(issue, dict):
            issue.setdefault("path", str(path))
    if result.get("status") != "passed" and not issues:
        issues.append(
            _issue(
                "docx_render_gate_failed",
                "Word 渲染门返回失败但未提供诊断。",
                path=path,
            )
        )
    return _result(
        "Word 已在固定渲染环境中通过逐页版面检查。"
        if not issues
        else "Word 渲染后发现版面异常。",
        metrics=dict(result.get("metrics") or {}),
        issues=issues,
        warnings=warnings,
    )


def _check_knowledge_base(context: _VerificationContext) -> dict[str, Any]:
    path = context.output_dir / "knowledge_base.jsonl"
    issues: list[dict[str, Any]] = []
    if not path.is_file():
        return _result(
            "知识库文件不存在。",
            issues=[_issue("knowledge_base_missing", "缺少 knowledge_base.jsonl。", path=path)],
        )
    rows: list[dict[str, Any]] = []
    try:
        for line_number, line in enumerate(path.read_text(encoding="utf-8").splitlines(), start=1):
            if not line.strip():
                continue
            try:
                row = json.loads(line)
            except json.JSONDecodeError as exc:
                issues.append(
                    _issue(
                        "knowledge_base_invalid_jsonl",
                        str(exc),
                        path=path,
                        line=line_number,
                    )
                )
                continue
            if not isinstance(row, dict):
                issues.append(
                    _issue(
                        "knowledge_base_row_invalid",
                        "知识库行不是 JSON 对象。",
                        path=path,
                        line=line_number,
                    )
                )
                continue
            rows.append(row)
    except (OSError, UnicodeError) as exc:
        return _result(
            "知识库文件无法读取。",
            issues=[_issue("knowledge_base_unreadable", str(exc), path=path)],
        )

    rag_embedding_status: str | None = None
    rag_manifest_path = rag_knowledge_base.manifest_path_for(path)
    try:
        rag_manifest = rag_knowledge_base.read_rag_manifest(path)
        rag_embedding_status = str(
            rag_manifest["retrieval"]["embedding"]["status"]
        )
    except rag_knowledge_base.RagError as exc:
        issues.append(
            _issue(
                "knowledge_base_rag_manifest_invalid",
                "RAG 清单缺失、过期或与知识库不一致。",
                path=rag_manifest_path,
                detail=str(exc),
            )
        )

    manifest_by_id = {str(item.get("id") or ""): item for item in context.manifest}
    rows_by_chapter: dict[str, list[dict[str, Any]]] = defaultdict(list)
    row_ids: list[str] = []
    encountered_orders: list[int] = []
    for index, row in enumerate(rows, start=1):
        keys = set(row)
        if keys != KB_FIELDS:
            issues.append(
                _issue(
                    "knowledge_base_fields_invalid",
                    "知识库字段必须严格符合白名单。",
                    path=path,
                    line=index,
                    expected=sorted(KB_FIELDS),
                    actual=sorted(keys),
                )
            )
        chapter_id_value = row.get("chapter_id")
        row_id_value = row.get("id")
        chapter_id = chapter_id_value if isinstance(chapter_id_value, str) else ""
        row_id = row_id_value if isinstance(row_id_value, str) else ""
        content = row.get("content")
        row_ids.append(row_id)
        for field_name, valid in (
            ("id", isinstance(row_id_value, str)),
            ("title", isinstance(row.get("title"), str)),
            ("chapter_id", isinstance(chapter_id_value, str)),
            (
                "chapter_order",
                isinstance(row.get("chapter_order"), int)
                and not isinstance(row.get("chapter_order"), bool),
            ),
            ("content", isinstance(content, str)),
        ):
            if not valid:
                issues.append(
                    _issue(
                        "knowledge_base_field_type_invalid",
                        "知识库字段类型不符合 schema。",
                        path=path,
                        line=index,
                        field=field_name,
                    )
                )
        item = manifest_by_id.get(chapter_id)
        if item is None:
            issues.append(
                _issue(
                    "knowledge_base_unknown_chapter",
                    "知识库行指向未知章节。",
                    path=path,
                    line=index,
                    chapter_id=chapter_id or None,
                )
            )
        else:
            if row.get("title") != _display_title(item):
                issues.append(
                    _issue(
                        "knowledge_base_title_mismatch",
                        "知识库章节标题与清单不一致。",
                        path=path,
                        line=index,
                        chapter_id=chapter_id,
                    )
                )
            if row.get("chapter_order") != item.get("sequence"):
                issues.append(
                    _issue(
                        "knowledge_base_order_mismatch",
                        "知识库章节顺序与清单不一致。",
                        path=path,
                        line=index,
                        chapter_id=chapter_id,
                    )
                )
            if isinstance(row.get("chapter_order"), int) and not isinstance(
                row.get("chapter_order"), bool
            ):
                encountered_orders.append(int(row["chapter_order"]))
        if not row_id:
            issues.append(_issue("knowledge_base_id_missing", "知识库行缺少 id。", path=path, line=index))
        elif re.fullmatch(r"[0-9a-f]{40}", row_id) is None:
            issues.append(
                _issue(
                    "knowledge_base_id_invalid",
                    "知识库 id 必须是稳定的 40 位小写 SHA-1。",
                    path=path,
                    line=index,
                    value=row_id,
                )
            )
        if not isinstance(content, str) or not content.strip():
            issues.append(
                _issue(
                    "knowledge_base_content_empty",
                    "知识库文本块为空。",
                    path=path,
                    line=index,
                    chapter_id=chapter_id or None,
                )
            )
        else:
            issues.extend(_trace_issues(content, path=path, chapter_id=chapter_id or None))
        rows_by_chapter[chapter_id].append(row)

    duplicates = sorted(value for value, count in Counter(row_ids).items() if value and count > 1)
    if duplicates:
        issues.append(
            _issue(
                "knowledge_base_duplicate_ids",
                "知识库文本块 id 不唯一。",
                path=path,
                values=duplicates,
            )
        )
    if encountered_orders != sorted(encountered_orders):
        issues.append(
            _issue(
                "knowledge_base_row_order_invalid",
                "知识库文本块没有按章节顺序排列。",
                path=path,
                actual=encountered_orders,
            )
        )
    missing_chapters = [chapter_id for chapter_id in manifest_by_id if not rows_by_chapter[chapter_id]]
    if missing_chapters:
        issues.append(
            _issue(
                "knowledge_base_chapter_missing",
                "知识库未覆盖全部章节。",
                path=path,
                values=missing_chapters,
            )
        )

    content_match_count = 0
    for item in context.manifest:
        chapter_id = str(item.get("id") or "")
        chapter_text = context.chapter_texts.get(chapter_id)
        if chapter_text is None:
            continue
        expected = _chapter_body(chapter_text)
        chunks = [
            str(row.get("content") or "").strip()
            for row in rows_by_chapter[chapter_id]
        ]
        if context.source_pdf is not None:
            row_source = "reviewed" if item.get("reviewed_override") else "compiled"
            expected_ids = [
                hashlib.sha1(
                    (
                        f"{context.source_pdf.name}:{chapter_id}:"
                        f"{row_source}:{chunk_index}"
                    ).encode("utf-8")
                ).hexdigest()
                for chunk_index in range(1, len(chunks) + 1)
            ]
            actual_ids = [
                str(row.get("id") or "")
                for row in rows_by_chapter[chapter_id]
            ]
            if actual_ids != expected_ids:
                issues.append(
                    _issue(
                        "knowledge_base_stable_ids_mismatch",
                        "知识库文本块 id 与框架的稳定 SHA-1 派生规则不一致。",
                        path=path,
                        chapter_id=chapter_id,
                        expected=expected_ids,
                        actual=actual_ids,
                    )
                )
        # ``split_text`` normally separates chunks on paragraph boundaries,
        # but a single paragraph longer than the chunk budget is cut inside
        # the paragraph.  Validate ordered, lossless coverage without
        # inventing one fixed join delimiter for both cases.
        cursor = 0
        coverage_ok = True
        for chunk in chunks:
            position = expected.find(chunk, cursor)
            if position < 0 or expected[cursor:position].strip():
                coverage_ok = False
                break
            cursor = position + len(chunk)
        if expected[cursor:].strip():
            coverage_ok = False
        actual = "\n\n".join(chunks)
        if not coverage_ok:
            issues.append(
                _issue(
                    "knowledge_base_chapter_content_mismatch",
                    "知识库文本块未完整重组当前 Markdown 章节正文。",
                    path=path,
                    chapter_id=chapter_id,
                    expected_sha256=_sha256_text(expected),
                    actual_sha256=_sha256_text(actual),
                )
            )
        else:
            content_match_count += 1

    return _result(
        "知识库字段合规并覆盖全部章节。" if not issues else "知识库结构或覆盖不完整。",
        metrics={
            "path": str(path),
            "chunk_count": len(rows),
            "covered_chapter_count": sum(bool(rows_by_chapter[key]) for key in manifest_by_id),
            "content_match_chapter_count": content_match_count,
            "chapter_count": len(manifest_by_id),
            "rag_manifest_path": str(rag_manifest_path),
            "rag_lexical_status": (
                "ready" if rag_embedding_status is not None else "invalid"
            ),
            "rag_embedding_status": rag_embedding_status or "invalid",
        },
        issues=issues,
    )


def _expected_bookmarks(toc_payload: dict[str, Any]) -> list[list[Any]]:
    entries: list[dict[str, Any]] = []
    raw_entries = toc_payload.get("entries")
    if not isinstance(raw_entries, list):
        raise ValueError("toc.json 的 entries 必须是数组。")
    for raw in raw_entries:
        if not isinstance(raw, dict) or not raw.get("pdf_page"):
            continue
        if (
            raw.get("kind") == "part"
            and _normalize_cover_title(str(raw.get("title") or ""))
            in {"封面", "封底", "frontcover", "backcover"}
        ):
            continue
        entries.append(raw)
    if not entries:
        raise ValueError("toc.json 没有可用于 PDF 书签的映射条目。")
    min_level = min(int(item.get("level") or 1) for item in entries)
    result: list[list[Any]] = []
    previous_level = 0
    for item in entries:
        level = max(1, int(item.get("level") or 1) - min_level + 1)
        level = min(level, previous_level + 1)
        result.append([level, _display_title(item), int(item["pdf_page"])])
        previous_level = level
    return result


def _check_pdf(context: _VerificationContext) -> dict[str, Any]:
    issues: list[dict[str, Any]] = []
    source_page_count = 0
    geometry_match_count = 0
    text_layer_match_count = 0
    render_match_count = 0
    path, candidates = _pick_artifact(
        context.output_dir,
        ".pdf",
        book_title=context.book_title,
        bookmarked_pdf=True,
    )
    if path is None:
        code = "bookmarked_pdf_missing" if not candidates else "bookmarked_pdf_ambiguous"
        return _result(
            "无法唯一确定带书签 PDF。",
            issues=[
                _issue(
                    code,
                    "未找到带目录 PDF。" if not candidates else "存在多个带目录 PDF，且无法按书名唯一选择。",
                    path=context.output_dir,
                    candidates=[item.name for item in candidates],
                )
            ],
        )
    if len(candidates) != 1:
        issues.append(
            _issue(
                "bookmarked_pdf_extra_artifacts",
                "输出目录含当前书名之外的旧带书签 PDF。",
                path=context.output_dir,
                candidates=[item.name for item in candidates],
            )
        )
    toc_path = context.output_dir / "toc.json"
    try:
        toc_payload = json.loads(toc_path.read_text(encoding="utf-8"))
        if not isinstance(toc_payload, dict):
            raise ValueError("toc.json 顶层必须是对象。")
        expected_toc = _expected_bookmarks(toc_payload)
    except (OSError, UnicodeError, json.JSONDecodeError, ValueError) as exc:
        return _result(
            "目录信息无法用于书签验证。",
            issues=[_issue("toc_invalid", str(exc), path=toc_path)],
        )
    output_page_count = 0
    if context.source_pdf is None:
        issues.append(
            _issue(
                "source_pdf_required",
                "完整 PDF 验收必须提供源 PDF，才能证明页数未被改变。",
                path=path,
            )
        )
    try:
        import fitz  # type: ignore[import-not-found]

        with fitz.open(path) as document:
            output_page_count = document.page_count
            actual_toc = [entry[:3] for entry in document.get_toc(simple=True)]
        if actual_toc != expected_toc:
            issues.append(
                _issue(
                    "pdf_bookmarks_mismatch",
                    "PDF 书签标题、层级或目标页与 toc.json 不一致。",
                    path=path,
                    expected=expected_toc,
                    actual=actual_toc,
                )
            )
        invalid_targets = [entry for entry in actual_toc if not 1 <= int(entry[2]) <= output_page_count]
        if invalid_targets:
            issues.append(
                _issue(
                    "pdf_bookmark_target_invalid",
                    "PDF 书签目标页超出文档范围。",
                    path=path,
                    values=invalid_targets,
                )
            )
        if context.source_pdf is not None:
            if not context.source_pdf.is_file():
                issues.append(
                    _issue(
                        "source_pdf_missing",
                        "指定的源 PDF 不存在。",
                        path=context.source_pdf,
                    )
                )
            else:
                with (
                    fitz.open(context.source_pdf) as source_document,
                    fitz.open(path) as output_document,
                ):
                    source_page_count = source_document.page_count
                    if output_page_count == source_page_count:
                        geometry_mismatches: list[int] = []
                        text_mismatches: list[int] = []
                        render_mismatches: list[int] = []
                        # RGB catches colour substitutions that grayscale can
                        # hide; quarter scale keeps the all-page gate cheap.
                        matrix = fitz.Matrix(0.25, 0.25)
                        for page_index in range(source_page_count):
                            source_page = source_document[page_index]
                            output_page = output_document[page_index]
                            source_geometry = (
                                round(float(source_page.rect.width), 4),
                                round(float(source_page.rect.height), 4),
                                int(source_page.rotation),
                            )
                            output_geometry = (
                                round(float(output_page.rect.width), 4),
                                round(float(output_page.rect.height), 4),
                                int(output_page.rotation),
                            )
                            if source_geometry == output_geometry:
                                geometry_match_count += 1
                            else:
                                geometry_mismatches.append(page_index + 1)

                            if source_page.get_text("text") == output_page.get_text("text"):
                                text_layer_match_count += 1
                            else:
                                text_mismatches.append(page_index + 1)

                            source_pixmap = source_page.get_pixmap(
                                matrix=matrix,
                                colorspace=fitz.csRGB,
                                alpha=False,
                            )
                            output_pixmap = output_page.get_pixmap(
                                matrix=matrix,
                                colorspace=fitz.csRGB,
                                alpha=False,
                            )
                            source_render = (
                                source_pixmap.width,
                                source_pixmap.height,
                                hashlib.sha256(source_pixmap.samples).digest(),
                            )
                            output_render = (
                                output_pixmap.width,
                                output_pixmap.height,
                                hashlib.sha256(output_pixmap.samples).digest(),
                            )
                            if source_render == output_render:
                                render_match_count += 1
                            else:
                                render_mismatches.append(page_index + 1)

                        for code, message, mismatches in (
                            (
                                "pdf_page_geometry_mismatch",
                                "带书签 PDF 的页面尺寸或旋转与源 PDF 不一致。",
                                geometry_mismatches,
                            ),
                            (
                                "pdf_text_layer_mismatch",
                                "带书签 PDF 的可复制文字层与源 PDF 不一致。",
                                text_mismatches,
                            ),
                            (
                                "pdf_page_render_mismatch",
                                "带书签 PDF 的页面外观与源 PDF 不一致。",
                                render_mismatches,
                            ),
                        ):
                            if mismatches:
                                issues.append(
                                    _issue(
                                        code,
                                        message,
                                        path=path,
                                        pages=mismatches[:50],
                                        count=len(mismatches),
                                    )
                                )
                if output_page_count != source_page_count:
                    issues.append(
                        _issue(
                            "pdf_page_count_mismatch",
                            "带书签 PDF 与源 PDF 页数不一致。",
                            path=path,
                            expected=source_page_count,
                            actual=output_page_count,
                        )
                    )
    except ImportError as exc:
        issues.append(_issue("pdf_dependency_missing", str(exc), path=path))
    except Exception as exc:  # PyMuPDF exposes several format-specific errors.
        issues.append(_issue("pdf_unreadable", f"{type(exc).__name__}: {exc}", path=path))

    return _result(
        "带书签 PDF 的书签、页数、页面外观和文字层均与源文件一致。"
        if not issues
        else "带书签 PDF 验证失败。",
        metrics={
            "path": str(path),
            "page_count": output_page_count,
            "source_page_count": source_page_count,
            "geometry_match_page_count": geometry_match_count,
            "text_layer_match_page_count": text_layer_match_count,
            "render_match_page_count": render_match_count,
            "bookmark_count": len(expected_toc),
        },
        issues=issues,
    )


def _check_runtime_hygiene(context: _VerificationContext) -> dict[str, Any]:
    issues: list[dict[str, Any]] = []
    temporary_paths: list[str] = []
    if context.output_dir.exists():
        for path in context.output_dir.rglob("*"):
            name = path.name.lower()
            if (
                name in {"tmp", "temp", ".tmp"}
                or name.endswith((".tmp", ".part", ".partial"))
                or (name.startswith(".") and ".tmp" in name)
                or re.fullmatch(r"_page_images_\d+_[0-9a-f-]+", name) is not None
            ):
                temporary_paths.append(str(path))
    if temporary_paths:
        issues.append(
            _issue(
                "temporary_artifacts_present",
                "输出目录仍有临时文件或目录。",
                path=context.output_dir,
                values=temporary_paths[:50],
                count=len(temporary_paths),
            )
        )

    active_stage_locks: list[str] = []
    stage_lock_dir = context.output_dir / ".stage_locks"
    if fcntl is not None and stage_lock_dir.is_dir():
        for lock_path in sorted(stage_lock_dir.glob("*.lock")):
            try:
                with lock_path.open("a+", encoding="utf-8") as handle:
                    try:
                        fcntl.flock(
                            handle.fileno(),
                            fcntl.LOCK_EX | fcntl.LOCK_NB,
                        )
                    except BlockingIOError:
                        active_stage_locks.append(str(lock_path))
                    else:
                        fcntl.flock(handle.fileno(), fcntl.LOCK_UN)
            except OSError as exc:
                issues.append(
                    _issue(
                        "stage_lock_unreadable",
                        f"无法检查阶段锁：{exc}",
                        path=lock_path,
                    )
                )
    if active_stage_locks:
        issues.append(
            _issue(
                "active_stage_locks",
                "仍有流水线阶段锁被占用。",
                path=stage_lock_dir,
                values=active_stage_locks,
            )
        )

    active_processes: list[dict[str, Any]] = []
    proc_root = Path("/proc")
    output_token = str(context.output_dir.resolve())
    if proc_root.is_dir():
        # The verifier itself, its launcher shell, and a supervising pipeline
        # naturally contain the output path in their command line.  They are
        # the current execution chain, not stale sibling workers.
        ancestor_pids: set[int] = {os.getpid()}
        ancestor = os.getpid()
        while ancestor > 1:
            try:
                stat = (proc_root / str(ancestor) / "stat").read_text()
                ancestor = int(stat.rsplit(")", 1)[1].split()[1])
            except (OSError, IndexError, ValueError):
                break
            ancestor_pids.add(ancestor)
        for proc_dir in proc_root.iterdir():
            if not proc_dir.name.isdigit() or int(proc_dir.name) in ancestor_pids:
                continue
            try:
                raw = (proc_dir / "cmdline").read_bytes()
                command = raw.replace(b"\x00", b" ").decode("utf-8", errors="replace").strip()
            except (OSError, PermissionError):
                continue
            if output_token not in command:
                continue
            if not re.search(
                r"(?:book_pipeline(?:\.py)?|publication_verifier|regenerate|ocr|translate)",
                command,
                flags=re.I,
            ):
                continue
            # Never copy argv into the report: deprecated raw-key flags may
            # still be present in another process even though this pipeline
            # no longer recommends them.
            executable = raw.split(b"\x00", 1)[0].decode(
                "utf-8", errors="replace"
            )
            active_processes.append(
                {
                    "pid": int(proc_dir.name),
                    "executable": Path(executable).name[:200],
                }
            )
    if active_processes:
        issues.append(
            _issue(
                "active_pipeline_processes",
                "仍有指向该输出目录的流水线进程。",
                path=context.output_dir,
                values=active_processes,
            )
        )

    return _result(
        "未发现临时产物或残留流水线进程。"
        if not issues
        else "发现临时产物或残留流水线进程。",
        metrics={
            "temporary_path_count": len(temporary_paths),
            "active_stage_lock_count": len(active_stage_locks),
            "active_process_count": len(active_processes),
        },
        issues=issues,
    )


def _normalise_chapter_ids(chapter_ids: Iterable[str] | str | None) -> list[str] | None:
    if chapter_ids is None:
        return None
    if isinstance(chapter_ids, str):
        values = chapter_ids.split(",")
    else:
        values = list(chapter_ids)
    return [str(value).strip() for value in values if str(value).strip()]


def _write_report(path: Path, report: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = json.dumps(report, ensure_ascii=False, indent=2) + "\n"
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{path.name}.", suffix=".tmp", dir=path.parent
    )
    temporary = Path(temporary_name)
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8") as handle:
            handle.write(payload)
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)


def verify_publication(
    output_dir: str | os.PathLike[str],
    source_pdf: str | os.PathLike[str] | None = None,
    book_title: str | None = None,
    expected_language: str | None = None,
    expected_translation_fingerprint: str | None = None,
    require_translation: bool = False,
    require_epub: bool = True,
    require_docx: bool = True,
    require_docx_render: bool = True,
    require_knowledge_base: bool = True,
    require_bookmarked_pdf: bool = True,
    require_all_reviewed: bool = False,
    publication_profile: str = "full",
    chapter_ids: Iterable[str] | str | None = None,
    report_path: str | os.PathLike[str] | None = None,
) -> dict[str, Any]:
    """Verify one compiled publication and return a machine-readable report.

    ``chapter_ids`` selects the fast, incremental chapter gate.  In that mode
    only the manifest, Markdown files, reviewed-source equality, citations,
    and content hygiene checks run.  Full-publication artifacts are reported
    as ``skipped``.  Every validation failure is captured in the returned
    report; normal bad input never escapes as an exception.
    """

    if publication_profile not in {"full", "word"}:
        raise ValueError(
            "publication_profile must be either 'full' or 'word'"
        )
    output_path = Path(output_dir).expanduser().resolve()
    source_path = Path(source_pdf).expanduser().resolve() if source_pdf is not None else None
    selected_ids = _normalise_chapter_ids(chapter_ids)
    destination = (
        Path(report_path).expanduser().resolve()
        if report_path is not None
        else output_path
        / "audit"
        / ("chapter-report.json" if selected_ids is not None else "release-report.json")
    )
    context = _VerificationContext(
        output_path,
        source_pdf=source_path,
        book_title=book_title,
        expected_language=expected_language,
        expected_translation_fingerprint=expected_translation_fingerprint,
        require_translation=require_translation,
        require_all_reviewed=require_all_reviewed,
        chapter_ids=selected_ids,
    )

    checks: list[dict[str, Any]] = []

    def run(check_id: str, operation: Callable[[_VerificationContext], dict[str, Any]]) -> None:
        try:
            value = operation(context)
        except Exception as exc:  # A verifier bug must become an actionable report.
            value = _result(
                "验证器执行该检查时发生内部错误。",
                issues=[
                    _issue(
                        "verifier_internal_error",
                        f"{type(exc).__name__}: {exc}",
                    )
                ],
            )
        issues = list(value.get("issues") or [])
        checks.append(
            {
                "id": check_id,
                "status": "failed" if issues else "passed",
                "summary": str(value.get("summary") or ""),
                "metrics": value.get("metrics") or {},
                "issues": issues,
                "warnings": list(value.get("warnings") or []),
            }
        )

    run("manifest.valid", _check_manifest)
    run("chapters.files", _check_chapter_files)
    run("reviewed.exact", _check_reviewed_exact)
    run("semantics.integrity", _check_semantics)
    run("citations.integrity", _check_citations)
    run("content.hygiene", _check_content_hygiene)

    full_checks: tuple[tuple[str, bool, Callable[[_VerificationContext], dict[str, Any]]], ...] = (
        ("checkpoints.complete", True, _check_checkpoints),
        ("epub.structure", require_epub, _check_epub),
        ("docx.structure", require_docx, _check_docx),
        (
            "docx.render",
            require_docx and require_docx_render,
            _check_docx_render,
        ),
        ("knowledge_base.structure", require_knowledge_base, _check_knowledge_base),
        ("pdf.bookmarks", require_bookmarked_pdf, _check_pdf),
        ("runtime.hygiene", True, _check_runtime_hygiene),
    )
    for check_id, required, operation in full_checks:
        if context.incremental:
            reason = "chapter_ids 启用了增量章节 gate。"
        elif not required:
            reason = "调用方未要求验证此发布产物。"
        else:
            run(check_id, operation)
            continue
        checks.append(
            {
                "id": check_id,
                "status": "skipped",
                "summary": reason,
                "metrics": {},
                "issues": [],
                "warnings": [],
            }
        )

    errors = [
        {"check_id": check["id"], **issue}
        for check in checks
        for issue in check["issues"]
    ]
    warnings = [
        {"check_id": check["id"], **warning}
        for check in checks
        for warning in check["warnings"]
    ]
    profile_allowed_skips = (
        {
            "epub.structure",
            "knowledge_base.structure",
            "pdf.bookmarks",
        }
        if publication_profile == "word"
        else set()
    )
    skipped_check_ids = {
        str(check["id"])
        for check in checks
        if check["status"] == "skipped"
    }
    unexpected_skips = bool(
        not context.incremental
        and (skipped_check_ids - profile_allowed_skips)
    )
    status = "failed" if errors else ("partial" if unexpected_skips else "passed")
    release_ready = bool(
        not context.incremental
        and not errors
        and not unexpected_skips
    )
    checks_by_id = {check["id"]: check for check in checks}

    def metric(check_id: str, name: str, default: int = 0) -> int:
        value = checks_by_id.get(check_id, {}).get("metrics", {}).get(name, default)
        return int(value) if isinstance(value, int) and not isinstance(value, bool) else default

    summary = {
        "passed": sum(check["status"] == "passed" for check in checks),
        "failed": sum(check["status"] == "failed" for check in checks),
        "skipped": sum(check["status"] == "skipped" for check in checks),
        "chapter_count": metric("manifest.valid", "chapter_count"),
        "selected_chapter_count": metric("manifest.valid", "selected_chapter_count"),
        "reviewed_chapter_count": metric("manifest.valid", "reviewed_chapter_count"),
        "reviewed_exact": metric("reviewed.exact", "exact_match_count"),
        "reference_count": metric("citations.integrity", "reference_count"),
        "definition_count": metric("citations.integrity", "definition_count"),
        "checkpoint_pages": metric(
            "checkpoints.complete", "checkpoint_page_count"
        ),
        "translated_pages": metric(
            "checkpoints.complete", "translated_page_count"
        ),
        "epub_chapters": metric("epub.structure", "chapter_count"),
        "epub_text_matches": metric(
            "epub.structure", "chapter_text_match_count"
        ),
        "docx_chapters": metric("docx.structure", "chapter_count"),
        "docx_text_matches": metric(
            "docx.structure", "chapter_text_match_count"
        ),
        "docx_footnotes": metric(
            "docx.structure", "footnote_definition_count"
        ),
        "docx_render_pages": metric("docx.render", "page_count"),
        "kb_chunks": metric("knowledge_base.structure", "chunk_count"),
        "kb_covered_chapters": metric(
            "knowledge_base.structure", "covered_chapter_count"
        ),
        "kb_content_matches": metric(
            "knowledge_base.structure", "content_match_chapter_count"
        ),
        "pdf_pages": metric("pdf.bookmarks", "page_count"),
        "pdf_render_matches": metric(
            "pdf.bookmarks", "render_match_page_count"
        ),
        "pdf_text_layer_matches": metric(
            "pdf.bookmarks", "text_layer_match_page_count"
        ),
        "pdf_bookmarks": metric("pdf.bookmarks", "bookmark_count"),
    }
    report: dict[str, Any] = {
        "schema_version": SCHEMA_VERSION,
        "status": status,
        "ok": (not errors) if context.incremental else release_ready,
        "release_ready": release_ready,
        "mode": "chapters" if context.incremental else "full",
        "generated_at": dt.datetime.now(dt.timezone.utc).isoformat(),
        "output_dir": str(output_path),
        "source_pdf": str(source_path) if source_path is not None else None,
        "book_title": (
            book_title
            or context.artifact_titles.get("epub")
            or context.artifact_titles.get("docx")
        ),
        "publication_profile": publication_profile,
        "expected_language": expected_language,
        "translation_required": require_translation,
        "docx_render_required": bool(require_docx and require_docx_render),
        "chapter_ids": selected_ids,
        "report_path": str(destination),
        "summary": summary,
        "metrics": {
            "check_count": len(checks),
            "passed_check_count": summary["passed"],
            "failed_check_count": summary["failed"],
            "skipped_check_count": summary["skipped"],
            "error_count": len(errors),
            "warning_count": len(warnings),
        },
        "checks": checks,
        "errors": errors,
        "warnings": warnings,
    }
    try:
        _write_report(destination, report)
    except Exception as exc:
        issue = _issue(
            "report_write_failed",
            f"{type(exc).__name__}: {exc}",
            path=destination,
        )
        write_check = {
            "id": "report.write",
            "status": "failed",
            "summary": "发布验收报告写入失败。",
            "metrics": {},
            "issues": [issue],
            "warnings": [],
        }
        write_error = {"check_id": "report.write", **issue}
        report["checks"].append(write_check)
        report["status"] = "failed"
        report["ok"] = False
        report["release_ready"] = False
        report["errors"].append(write_error)
        report["summary"]["failed"] += 1
        report["metrics"]["check_count"] = len(report["checks"])
        report["metrics"]["failed_check_count"] = report["summary"]["failed"]
        report["metrics"]["error_count"] = len(report["errors"])
    return report


__all__ = ["CHECK_IDS", "SCHEMA_VERSION", "verify_publication"]
